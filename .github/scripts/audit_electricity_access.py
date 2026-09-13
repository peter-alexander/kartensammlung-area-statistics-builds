#!/usr/bin/env python3
from __future__ import annotations

import math
import statistics
import sys
import tempfile
import time
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from urllib.error import HTTPError, URLError
from urllib.request import Request, urlopen

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "scripts"))

from build_world_bank import fetch_world_bank_records, load_registry  # noqa: E402

REGISTRY = "https://tiles.radlobby.at/AreaStatistics/area-registry-countries.json"
XLSX_URL = "https://trackingsdg7.esmap.org/sites/default/files/download-documents/sdg7.1.1_-_access_to_electricity.xlsx"
WDI_API = "https://api.worldbank.org/v2"
WDI_SOURCE_ID = 2
WDI_INDICATOR = "EG.ELC.ACCS.ZS"
USER_AGENT = "kartensammlung-area-statistics-builds/1"


def download(url: str, suffix: str) -> Path:
	last_error: Exception | None = None
	for attempt, delay in enumerate((0, 5, 15, 30), start=1):
		if delay:
			time.sleep(delay)
		handle = tempfile.NamedTemporaryFile(prefix="electricity-audit-", suffix=suffix, delete=False)
		path = Path(handle.name)
		handle.close()
		try:
			request = Request(url, headers={"User-Agent": USER_AGENT, "Accept": "*/*"})
			with urlopen(request, timeout=90) as response, path.open("wb") as output:
				while True:
					chunk = response.read(1024 * 1024)
					if not chunk:
						break
					output.write(chunk)
			if path.stat().st_size < 1000000:
				raise RuntimeError(f"Workbook unexpectedly small: {path.stat().st_size} bytes")
			return path
		except (HTTPError, URLError, TimeoutError, OSError, RuntimeError) as error:
			last_error = error
			path.unlink(missing_ok=True)
			print(f"Download failed ({attempt}/4): {error}")
	raise RuntimeError("Tracking SDG7 workbook download failed") from last_error


def numeric(value: object) -> float | None:
	if value is None or isinstance(value, bool):
		return None
	if isinstance(value, (int, float)):
		result = float(value)
	else:
		text = str(value).strip()
		if not text or text.upper() in {"NA", "N/A", "..", "..."}:
			return None
		result = float(text)
	if not math.isfinite(result):
		raise ValueError(f"Non-finite value: {value!r}")
	return result


def read_direct(path: Path, area_by_iso3: dict[str, str]) -> tuple[dict[tuple[str, int], float], Counter[str], set[str], set[int]]:
	wb = load_workbook(path, read_only=True, data_only=True)
	if "UN reporting" not in wb.sheetnames:
		raise RuntimeError("Tracking SDG7 workbook is missing UN reporting sheet")
	ws = wb["UN reporting"]
	rows = ws.iter_rows(values_only=True)
	header = next(rows)
	index = {str(value).strip(): position for position, value in enumerate(header) if value is not None}
	required = {
		"SeriesCode",
		"Indicator",
		"GeoAreaName/Reference Area Name",
		"TimePeriod",
		"Value",
		"Units",
		"Nature",
		"Location",
		"Reporting Type",
		"Source",
		"ISOalpha3",
		"Type",
	}
	missing = sorted(required - set(index))
	if missing:
		raise RuntimeError(f"UN reporting sheet missing columns: {missing}")

	values: dict[tuple[str, int], float] = {}
	natures: Counter[str] = Counter()
	ignored_iso3: set[str] = set()
	years: set[int] = set()
	invariant_values: dict[str, set[str]] = {
		"SeriesCode": set(),
		"Indicator": set(),
		"Units": set(),
		"Location": set(),
		"Reporting Type": set(),
		"Source": set(),
		"Type": set(),
	}

	for row in rows:
		series = str(row[index["SeriesCode"]] or "").strip()
		indicator = str(row[index["Indicator"]] or "").strip()
		location = str(row[index["Location"]] or "").strip()
		type_name = str(row[index["Type"]] or "").strip()
		if series != "EG_ACS_ELEC" or indicator != "7.1.1" or location != "ALLAREA" or type_name != "Country":
			continue

		for field in invariant_values:
			invariant_values[field].add(str(row[index[field]] or "").strip())

		iso3 = str(row[index["ISOalpha3"]] or "").strip().upper()
		if iso3 not in area_by_iso3:
			if iso3:
				ignored_iso3.add(iso3)
			continue
		year_raw = row[index["TimePeriod"]]
		if isinstance(year_raw, bool) or not isinstance(year_raw, (int, float)):
			continue
		year = int(year_raw)
		value = numeric(row[index["Value"]])
		if value is None:
			continue
		if not 0 <= value <= 100:
			raise RuntimeError(f"Direct value out of range: {iso3} {year} {value}")
		key = (area_by_iso3[iso3], year)
		if key in values:
			raise RuntimeError(f"Duplicate direct value: {key}")
		values[key] = value
		years.add(year)
		natures[str(row[index["Nature"]] or "").strip()] += 1

	print("Direct source invariants:")
	for field, found in invariant_values.items():
		print(f"  {field}={sorted(found)}")
	return values, natures, ignored_iso3, years


def read_wdi(area_by_iso3: dict[str, str]) -> tuple[dict[tuple[str, int], float], set[str], set[int]]:
	records = fetch_world_bank_records(
		WDI_API,
		f"country/all/indicator/{WDI_INDICATOR}",
		WDI_SOURCE_ID,
		90,
	)
	values: dict[tuple[str, int], float] = {}
	ignored_iso3: set[str] = set()
	years: set[int] = set()
	current_year = datetime.now(timezone.utc).year
	for record in records:
		value = numeric(record.get("value"))
		if value is None:
			continue
		if str(record.get("obs_status", "")).strip().upper() == "F":
			continue
		iso3 = str(record.get("countryiso3code", "")).strip().upper()
		if iso3 not in area_by_iso3:
			if iso3:
				ignored_iso3.add(iso3)
			continue
		year_text = str(record.get("date", "")).strip()
		if len(year_text) != 4 or not year_text.isdigit():
			continue
		year = int(year_text)
		if year > current_year:
			continue
		if not 0 <= value <= 100:
			raise RuntimeError(f"WDI value out of range: {iso3} {year} {value}")
		key = (area_by_iso3[iso3], year)
		if key in values:
			raise RuntimeError(f"Duplicate WDI value: {key}")
		values[key] = value
		years.add(year)
	return values, ignored_iso3, years


def year_counts(values: dict[tuple[str, int], float]) -> Counter[int]:
	return Counter(year for _, year in values)


def area_counts(values: dict[tuple[str, int], float]) -> Counter[str]:
	return Counter(area for area, _ in values)


def main() -> None:
	area_by_iso3, _ = load_registry(REGISTRY, 90)
	iso3_by_area = {area: iso3 for iso3, area in area_by_iso3.items()}
	path = download(XLSX_URL, ".xlsx")
	try:
		direct, natures, direct_ignored, direct_years = read_direct(path, area_by_iso3)
	finally:
		path.unlink(missing_ok=True)
	wdi, wdi_ignored, wdi_years = read_wdi(area_by_iso3)

	common = sorted(set(direct) & set(wdi))
	direct_only = sorted(set(direct) - set(wdi))
	wdi_only = sorted(set(wdi) - set(direct))
	diffs = [abs(direct[key] - wdi[key]) for key in common]

	print("\nCoverage:")
	print(f"direct observations={len(direct)} areas={len(area_counts(direct))} years={min(direct_years)}-{max(direct_years)}")
	print(f"wdi observations={len(wdi)} areas={len(area_counts(wdi))} years={min(wdi_years)}-{max(wdi_years)}")
	print(f"common={len(common)} directOnly={len(direct_only)} wdiOnly={len(wdi_only)}")
	print(f"directIgnoredIso3={sorted(direct_ignored)}")
	print(f"wdiIgnoredIso3Count={len(wdi_ignored)}")
	print(f"natureCounts={dict(sorted(natures.items()))}")

	print("\nNumeric overlap:")
	for tolerance in (0.0, 1e-12, 1e-9, 1e-6, 1e-4, 0.05):
		count = sum(diff <= tolerance for diff in diffs)
		print(f"within {tolerance:g}: {count}/{len(diffs)}")
	print(f"medianAbsDiff={statistics.median(diffs) if diffs else None}")
	print(f"meanAbsDiff={statistics.fmean(diffs) if diffs else None}")
	print(f"maxAbsDiff={max(diffs) if diffs else None}")
	if diffs:
		largest = sorted(((abs(direct[key] - wdi[key]), key, direct[key], wdi[key]) for key in common), reverse=True)[:20]
		print("largest differences:")
		for diff, (area, year), direct_value, wdi_value in largest:
			print(f"  {iso3_by_area[area]} {year}: direct={direct_value!r} wdi={wdi_value!r} diff={diff!r}")

	print("\nExclusive observations by year:")
	direct_only_years = year_counts({key: direct[key] for key in direct_only})
	wdi_only_years = year_counts({key: wdi[key] for key in wdi_only})
	print(f"directOnlyYears={dict(sorted(direct_only_years.items()))}")
	print(f"wdiOnlyYears={dict(sorted(wdi_only_years.items()))}")
	print(f"wdiOnlyBefore2000={sum(count for year, count in wdi_only_years.items() if year < 2000)}")
	print(f"wdiOnlyFrom2000={sum(count for year, count in wdi_only_years.items() if year >= 2000)}")

	for year in (2000, 2020, 2023, 2024):
		direct_year = {area: value for (area, y), value in direct.items() if y == year}
		wdi_year = {area: value for (area, y), value in wdi.items() if y == year}
		common_areas = set(direct_year) & set(wdi_year)
		year_diffs = [abs(direct_year[area] - wdi_year[area]) for area in common_areas]
		print(
			f"year {year}: direct={len(direct_year)} wdi={len(wdi_year)} common={len(common_areas)} "
			f"directOnly={len(set(direct_year)-set(wdi_year))} wdiOnly={len(set(wdi_year)-set(direct_year))} "
			f"exact={sum(diff == 0 for diff in year_diffs)}/{len(year_diffs)} "
			f"maxDiff={max(year_diffs) if year_diffs else None}"
		)

	print("\nDirect-only areas:")
	for area, count in area_counts({key: direct[key] for key in direct_only}).most_common():
		print(f"  {iso3_by_area[area]}: {count}")
	print("WDI-only areas (top 40):")
	for area, count in area_counts({key: wdi[key] for key in wdi_only}).most_common(40):
		print(f"  {iso3_by_area[area]}: {count}")


if __name__ == "__main__":
	main()
