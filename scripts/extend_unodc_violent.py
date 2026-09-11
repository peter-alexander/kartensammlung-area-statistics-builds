#!/usr/bin/env python3
from __future__ import annotations

import argparse
import hashlib
import json
import math
import shutil
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from urllib.parse import urljoin

import openpyxl
import pycountry

import build_unodc as base

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_CONFIG = ROOT / "config" / "unodc-violent-indicators.json"
DEFAULT_OUTPUT_DIR = ROOT / "dist" / "statistics"
DEFAULT_REGISTRY = "https://tiles.radlobby.at/AreaStatistics/area-registry-countries.json"


def parse_args() -> argparse.Namespace:
	parser = argparse.ArgumentParser(description="Extend the UNODC provider with violent and sexual crime statistics.")
	parser.add_argument("--config", type=Path, default=DEFAULT_CONFIG)
	parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
	parser.add_argument("--registry", default=DEFAULT_REGISTRY)
	parser.add_argument("--timeout", type=int, default=90)
	return parser.parse_args()


def validate_config(payload: Any) -> dict[str, Any]:
	if not isinstance(payload, dict) or payload.get("schema") != "kartensammlung.unodc-violent-statistics/v1":
		raise ValueError("Invalid UNODC violent-crime config.")
	urls = payload.get("downloadUrls")
	if not isinstance(urls, list) or not urls or not all(str(url).startswith("https://") for url in urls):
		raise ValueError("UNODC violent-crime downloadUrls must be a non-empty HTTPS URL list.")
	for key in ("sourcePage", "metadataUrl", "populationIndexUrl"):
		if not str(payload.get(key, "")).startswith("https://"):
			raise ValueError(f"UNODC violent-crime config requires HTTPS {key}.")
	fraction = payload.get("broadCoverageFraction")
	if not isinstance(fraction, (int, float)) or not 0 < float(fraction) <= 1:
		raise ValueError("Invalid broadCoverageFraction.")
	if not isinstance(payload.get("maxDefaultYearAge"), int) or int(payload["maxDefaultYearAge"]) < 0:
		raise ValueError("Invalid maxDefaultYearAge.")
	for field in ("countryAliases", "excludedCountryNames", "countryCompositions"):
		if not isinstance(payload.get(field), dict):
			raise ValueError(f"UNODC violent-crime config requires {field} object.")
	categories = payload.get("categories")
	if not isinstance(categories, list) or not categories:
		raise ValueError("UNODC violent-crime config requires categories.")
	keys: set[str] = set()
	source_categories: set[str] = set()
	for category in categories:
		if not isinstance(category, dict):
			raise ValueError("UNODC violent-crime category entries must be objects.")
		for key in ("key", "sourceCategory", "title", "description"):
			if not str(category.get(key, "")).strip():
				raise ValueError(f"UNODC violent-crime category entry is missing {key}.")
		category_key = str(category["key"])
		source_category = str(category["sourceCategory"])
		if category_key in keys or source_category.casefold() in source_categories:
			raise ValueError(f"Duplicate UNODC violent-crime category: {category_key}.")
		keys.add(category_key)
		source_categories.add(source_category.casefold())
		if not isinstance(category.get("minAreasWithAnyValue"), int) or int(category["minAreasWithAnyValue"]) < 1:
			raise ValueError(f"Category {category_key} has invalid minAreasWithAnyValue.")
		for break_field in ("countBreaks", "rateBreaks"):
			breaks = category.get(break_field)
			if not isinstance(breaks, list) or len(breaks) != 6:
				raise ValueError(f"Category {category_key} requires six {break_field} values.")
			numbers = [float(value) for value in breaks]
			if any(not math.isfinite(value) for value in numbers) or any(numbers[index] <= numbers[index - 1] for index in range(1, 6)):
				raise ValueError(f"Category {category_key} has invalid {break_field}.")
	return payload


def read_rows(path: Path) -> list[dict[str, Any]]:
	workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
	try:
		if "data" not in workbook.sheetnames:
			raise RuntimeError(f"UNODC violent-crime workbook has unexpected sheets: {workbook.sheetnames}")
		worksheet = workbook["data"]
		rows = worksheet.iter_rows(values_only=True)
		next(rows, None)
		next(rows, None)
		header = [base.text(value) for value in next(rows, ())]
		expected = ["Region", "Subregion", "Country", "Category", "variable", "value"]
		if header[:6] != expected:
			raise RuntimeError(f"UNODC violent-crime header changed: {header}")
		records: list[dict[str, Any]] = []
		for row in rows:
			if len(row) < 6:
				continue
			country = base.text(row[2])
			category = base.text(row[3])
			year = base.parse_year(row[4])
			value = base.parse_number(row[5])
			if not country or not category or year is None or value is None:
				continue
			if float(value) < 0:
				raise RuntimeError(f"Negative UNODC violent-crime count: {country} {category} {year} = {value}")
			records.append({
				"region": base.text(row[0]),
				"subregion": base.text(row[1]),
				"country": country,
				"category": category,
				"year": year,
				"value": value,
			})
		return records
	finally:
		workbook.close()


def resolve_iso3(country_name: str, config: dict[str, Any]) -> str | None:
	aliases = config["countryAliases"]
	if country_name in aliases:
		return str(aliases[country_name]).upper()
	if country_name in config["excludedCountryNames"]:
		return None
	try:
		return str(pycountry.countries.lookup(country_name).alpha_3).upper()
	except LookupError:
		return None


def set_count(target: dict[str, dict[int, dict[str, int | float]]], category: str, year: int, area_id: str, value: int | float) -> None:
	year_values = target.setdefault(category, {}).setdefault(year, {})
	if area_id in year_values and year_values[area_id] != value:
		raise RuntimeError(f"Conflicting UNODC violent-crime values for {category} {year} {area_id}.")
	year_values[area_id] = value


def map_counts(records: list[dict[str, Any]], config: dict[str, Any], area_by_iso3: dict[str, str]) -> tuple[dict[str, dict[int, dict[str, int | float]]], set[str], set[str], list[str]]:
	mapped: dict[str, dict[int, dict[str, int | float]]] = {}
	unresolved: set[str] = set()
	excluded: set[str] = set()
	composition_names = {
		component
		for components in config["countryCompositions"].values()
		for component in components
	}
	composition_values: dict[str, dict[int, dict[str, int | float]]] = {}

	for record in records:
		country = str(record["country"])
		category = str(record["category"])
		year = int(record["year"])
		value = record["value"]
		if country in composition_names:
			year_values = composition_values.setdefault(category, {}).setdefault(year, {})
			if country in year_values and year_values[country] != value:
				raise RuntimeError(f"Conflicting UNODC component values for {country} {category} {year}.")
			year_values[country] = value
			continue
		if country in config["excludedCountryNames"]:
			excluded.add(country)
			continue
		iso3 = resolve_iso3(country, config)
		if not iso3 or iso3 not in area_by_iso3:
			unresolved.add(country)
			continue
		set_count(mapped, category, year, area_by_iso3[iso3], value)

	incomplete_compositions: list[str] = []
	for target_iso3, components in config["countryCompositions"].items():
		if target_iso3 not in area_by_iso3:
			raise RuntimeError(f"Composition target {target_iso3} is missing from area registry.")
		target_area_id = area_by_iso3[target_iso3]
		for category, years in composition_values.items():
			for year, values in years.items():
				present = [component for component in components if component in values]
				if not present:
					continue
				if len(present) != len(components):
					incomplete_compositions.append(f"{target_iso3}:{category}:{year}:{','.join(present)}")
					continue
				total = sum(float(values[component]) for component in components)
				value: int | float = int(total) if total.is_integer() else total
				set_count(mapped, category, year, target_area_id, value)
	return mapped, unresolved, excluded, incomplete_compositions


def load_population(index_url: str, timeout: int) -> tuple[dict[str, Any], dict[str, Any], str]:
	index = base.fetch_json(index_url, timeout)
	if not isinstance(index, dict) or index.get("schema") != "kartensammlung.statistics-provider-index/v1":
		raise RuntimeError("Invalid UN WPP provider index for UNODC rate denominator.")
	entry = next((item for item in index.get("indicators", []) if isinstance(item, dict) and item.get("id") == "population.total"), None)
	if not isinstance(entry, dict) or not str(entry.get("path", "")):
		raise RuntimeError("UN WPP population.total indicator is missing.")
	payload_url = urljoin(index_url, str(entry["path"]))
	payload = base.fetch_json(payload_url, timeout)
	if not isinstance(payload, dict) or payload.get("schema") != "kartensammlung.statistics-indicator/v1":
		raise RuntimeError("Invalid UN WPP population.total payload.")
	return index, payload, payload_url


def calculate_rates(counts_by_year: dict[int, dict[str, int | float]], population_payload: dict[str, Any]) -> dict[int, dict[str, int | float]]:
	rates: dict[int, dict[str, int | float]] = {}
	population_values = population_payload.get("values")
	if not isinstance(population_values, dict):
		raise RuntimeError("UN WPP population.total payload contains no values.")
	for year, counts in counts_by_year.items():
		population_for_year = population_values.get(str(year))
		if not isinstance(population_for_year, dict):
			continue
		for area_id, count in counts.items():
			population = population_for_year.get(area_id)
			if not isinstance(population, (int, float)) or float(population) <= 0:
				continue
			rate = 100000.0 * float(count) / float(population)
			rates.setdefault(year, {})[area_id] = round(rate, 4)
	return rates


def build_payload(
	config: dict[str, Any],
	category: dict[str, Any],
	mode: str,
	values_by_year: dict[int, dict[str, int | float]],
	area_count: int,
	now: datetime,
	download_url: str,
	provider: dict[str, Any],
	population_index: dict[str, Any],
	population_url: str,
) -> tuple[dict[str, Any], dict[str, Any]]:
	key = str(category["key"])
	if mode == "count":
		indicator_id = f"crime.{key}-count"
		slug = f"crime-{key}-count"
		title = f"{category['title']} – Anzahl"
		unit = {"id": "recorded-offences", "label": "registrierte Fälle"}
		classification = {"type": "fixed", "scale": "logarithmic", "breaks": category["countBreaks"]}
	else:
		indicator_id = f"crime.{key}-rate"
		slug = f"crime-{key}-rate"
		title = f"{category['title']} – je 100.000"
		unit = {"id": "per-100000-population", "label": "je 100.000 Einwohner", "symbol": "je 100.000"}
		classification = {"type": "fixed", "breaks": category["rateBreaks"]}
	areas = {area_id for year_values in values_by_year.values() for area_id in year_values}
	minimum = int(category["minAreasWithAnyValue"])
	if len(areas) < minimum:
		raise RuntimeError(f"Coverage too small for {indicator_id}: {len(areas)} areas, expected at least {minimum}.")
	available_years = sorted(values_by_year)
	if not available_years:
		raise RuntimeError(f"No values for {indicator_id}.")
	coverage_fraction = float(config["broadCoverageFraction"])
	broad_threshold = max(1, math.ceil(len(areas) * coverage_fraction))
	broad_years = [year for year in available_years if year <= now.year and len(values_by_year[year]) >= broad_threshold]
	if not broad_years:
		raise RuntimeError(f"No year reaches broad coverage for {indicator_id}.")
	default_year = broad_years[-1]
	if now.year - default_year > int(config["maxDefaultYearAge"]):
		raise RuntimeError(f"Latest broadly covered year for {indicator_id} is unexpectedly old: {default_year}.")
	latest_year = available_years[-1]
	warning = "Registrierte Polizei-/Verwaltungsdaten; Unterschiede in Anzeige-, Rechts- und Erfassungspraxis schränken Ländervergleiche ein. Die Daten messen nicht die tatsächliche Prävalenz."
	description = f"{category['description']} {warning}"
	if mode == "rate":
		description += " Die Rate wurde aus der UNODC-Fallzahl und der UN-WPP-Bevölkerung berechnet."
	source: dict[str, Any] = {
		"providerId": provider["id"],
		"providerName": provider["name"],
		"dataset": "UNODC Violent Crime & Sexual Violence",
		"datasetId": "UNODC-CTS-Violent-Sexual-Crime",
		"indicator": {
			"category": category["sourceCategory"],
			"measure": "recorded offences",
		},
		"license": provider["license"],
		"licenseUrl": provider["licenseUrl"],
		"attribution": provider["attribution"],
		"url": config["sourcePage"],
		"downloadUrl": download_url,
		"metadataUrl": config["metadataUrl"],
	}
	if mode == "rate":
		source["derivedFrom"] = {
			"formula": "100000 * UNODC recorded offences / UN WPP total population",
			"numerator": {
				"providerId": "unodc",
				"datasetId": "UNODC-CTS-Violent-Sexual-Crime",
				"category": category["sourceCategory"],
			},
			"denominator": {
				"providerId": "un-wpp",
				"indicatorId": "population.total",
				"activeSnapshot": population_index.get("activeSnapshot"),
				"url": population_url,
				"projectionStartYear": 2024,
				"projectionVariant": "Medium",
			},
		}
	payload = {
		"schema": "kartensammlung.statistics-indicator/v1",
		"indicator": {
			"id": indicator_id,
			"title": title,
			"description": description,
			"areaLevel": "country",
			"frequency": "annual",
			"unit": unit,
			"classification": classification,
		},
		"source": source,
		"availableYears": available_years,
		"defaultYear": default_year,
		"coverage": {
			"registryAreas": area_count,
			"areasWithAnyValue": len(areas),
			"latestYear": latest_year,
			"areasInLatestYear": len(values_by_year[latest_year]),
			"broadCoverageFraction": coverage_fraction,
			"broadCoverageThreshold": broad_threshold,
			"areasInDefaultYear": len(values_by_year[default_year]),
		},
		"values": {
			str(year): {area_id: values_by_year[year][area_id] for area_id in sorted(values_by_year[year])}
			for year in available_years
		},
	}
	entry = {
		"id": indicator_id,
		"title": title,
		"description": description,
		"areaLevel": "country",
		"frequency": "annual",
		"unit": unit,
		"classification": classification,
		"sourceIndicator": source["indicator"],
		"slug": slug,
		"availableYears": available_years,
		"defaultYear": default_year,
		"coverage": payload["coverage"],
	}
	return payload, entry


def load_existing_payloads(provider_dir: Path, index: dict[str, Any]) -> list[tuple[dict[str, Any], dict[str, Any], str]]:
	result: list[tuple[dict[str, Any], dict[str, Any], str]] = []
	for entry in index.get("indicators", []):
		if not isinstance(entry, dict) or not str(entry.get("path", "")):
			raise RuntimeError("Existing UNODC provider index contains invalid indicator entry.")
		path = provider_dir / str(entry["path"])
		payload = base.read_json(path)
		filename = Path(str(entry["path"])).name
		result.append((dict(entry), payload, filename))
	return result


def main() -> None:
	args = parse_args()
	config = validate_config(base.read_json(args.config))
	provider_dir = args.output_dir / "unodc"
	index_path = provider_dir / "index.json"
	if not index_path.exists():
		raise RuntimeError("Base UNODC provider index is missing; run build_unodc.py first.")
	base_index = base.read_json(index_path)
	if not isinstance(base_index, dict) or base_index.get("schema") != "kartensammlung.statistics-provider-index/v1":
		raise RuntimeError("Invalid base UNODC provider index.")
	provider = base_index.get("provider")
	if not isinstance(provider, dict) or provider.get("id") != "unodc":
		raise RuntimeError("Base provider index is not UNODC.")
	for key in ("id", "name", "license", "licenseUrl", "attribution"):
		if not str(provider.get(key, "")).strip():
			raise RuntimeError(f"Base UNODC provider metadata is missing {key}.")

	area_by_iso3, registry = base.load_registry(args.registry, args.timeout)
	population_index, population_payload, population_url = load_population(config["populationIndexUrl"], args.timeout)
	now = datetime.now(timezone.utc)
	path, download_url = base.download_xlsx([str(url) for url in config["downloadUrls"]], args.timeout)
	try:
		print(f"Downloaded UNODC violent-crime workbook: {path.stat().st_size} bytes")
		records = read_rows(path)
	finally:
		path.unlink(missing_ok=True)
	if not records:
		raise RuntimeError("UNODC violent-crime workbook contained no parseable rows.")
	print(f"UNODC violent-crime parsed rows: {len(records)}")
	counts_by_category, unresolved, excluded, incomplete = map_counts(records, config, area_by_iso3)
	if unresolved:
		raise RuntimeError(f"Unresolved UNODC country names: {sorted(unresolved)}")
	print(f"Excluded partial/non-country rows: {sorted(excluded) or '(none)'}")
	print(f"Incomplete country-composition rows skipped: {len(incomplete)}")
	for item in incomplete[:30]:
		print(f"Incomplete composition: {item}")

	existing = load_existing_payloads(provider_dir, base_index)
	new_items: list[tuple[dict[str, Any], dict[str, Any], str]] = []
	for category in config["categories"]:
		source_category = str(category["sourceCategory"])
		counts = counts_by_category.get(source_category, {})
		if not counts:
			raise RuntimeError(f"UNODC source category is missing: {source_category}")
		rates = calculate_rates(counts, population_payload)
		for mode, values in (("count", counts), ("rate", rates)):
			payload, entry = build_payload(
				config,
				category,
				mode,
				values,
				len(area_by_iso3),
				now,
				download_url,
				provider,
				population_index,
				population_url,
			)
			filename = f"{entry.pop('slug')}.json"
			new_items.append((entry, payload, filename))
			coverage = payload["coverage"]
			print(
				f"{entry['id']}: mapped={coverage['areasWithAnyValue']} "
				f"years={payload['availableYears'][0]}-{payload['availableYears'][-1]} "
				f"defaultYear={payload['defaultYear']} defaultCoverage={coverage['areasInDefaultYear']}"
			)

	all_ids = [str(entry["id"]) for entry, _, _ in existing] + [str(entry["id"]) for entry, _, _ in new_items]
	if len(all_ids) != len(set(all_ids)):
		raise RuntimeError("Duplicate indicator IDs while combining UNODC datasets.")
	hasher = hashlib.sha256()
	for entry, payload, _ in sorted(existing + new_items, key=lambda item: str(item[0]["id"])):
		hasher.update(str(entry["id"]).encode("utf-8") + b"\0" + base.canonical_bytes(payload) + b"\0")
	snapshot = hasher.hexdigest()[:16]
	old_snapshot = str(base_index.get("activeSnapshot", ""))
	old_release_dir = provider_dir / "releases" / old_snapshot
	new_release_dir = provider_dir / "releases" / snapshot
	if new_release_dir.exists():
		shutil.rmtree(new_release_dir)
	new_release_dir.mkdir(parents=True, exist_ok=True)

	combined_entries: list[dict[str, Any]] = []
	for entry, payload, filename in existing + new_items:
		base.write_json(new_release_dir / filename, payload)
		new_entry = dict(entry)
		new_entry["path"] = f"releases/{snapshot}/{filename}"
		combined_entries.append(new_entry)
	combined_entries.sort(key=lambda item: (0 if str(item["id"]).startswith("homicide.") else 1, str(item["id"])))

	registry_source = base_index.get("areaRegistry")
	if not isinstance(registry_source, dict):
		registry_source = {
			"url": args.registry if args.registry.startswith(("https://", "http://")) else None,
			"schema": registry.get("schema"),
			"generatedAt": registry.get("generatedAt"),
			"areaCount": len(area_by_iso3),
		}
	index_payload = dict(base_index)
	index_payload["retrievedAt"] = now.isoformat().replace("+00:00", "Z")
	index_payload["activeSnapshot"] = snapshot
	index_payload["areaRegistry"] = registry_source
	index_payload["indicators"] = combined_entries
	index_payload["additionalDatasets"] = [
		{
			"id": "UNODC-CTS-Violent-Sexual-Crime",
			"name": "UNODC Violent Crime & Sexual Violence",
			"downloadUrl": download_url,
			"sourcePage": config["sourcePage"],
			"metadataUrl": config["metadataUrl"],
			"measures": ["recorded offences", "derived rate per 100,000 population"],
			"populationDenominator": {
				"providerId": "un-wpp",
				"indicatorId": "population.total",
				"activeSnapshot": population_index.get("activeSnapshot"),
				"url": population_url,
			},
		}
	]
	index_payload["excludedCountryNames"] = config["excludedCountryNames"]
	index_payload["countryCompositions"] = config["countryCompositions"]
	base.write_json(index_path, index_payload)

	if old_snapshot and old_snapshot != snapshot and old_release_dir.exists():
		shutil.rmtree(old_release_dir)
	print(f"Extended UNODC snapshot {snapshot}")
	print(f"Indicators: {len(combined_entries)} ({len(existing)} existing + {len(new_items)} violent-crime)")
	print(f"UN WPP denominator snapshot: {population_index.get('activeSnapshot')}")


if __name__ == "__main__":
	main()
