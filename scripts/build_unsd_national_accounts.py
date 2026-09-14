#!/usr/bin/env python3
from __future__ import annotations

import argparse
from collections import defaultdict
from datetime import datetime, timezone
import hashlib
import html
import io
import json
import math
from pathlib import Path
import re
import time
from typing import Any
from urllib.error import HTTPError, URLError
from urllib.request import Request, urlopen

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_CONFIG = ROOT / "config" / "unsd-national-accounts-indicators.json"
DEFAULT_PROVIDERS = ROOT / "config" / "statistics-providers.json"
DEFAULT_OUTPUT_DIR = ROOT / "dist" / "statistics"
DEFAULT_REGISTRY = "https://tiles.radlobby.at/AreaStatistics/area-registry-countries.json"
USER_AGENT = "kartensammlung-area-statistics-builds/1"


def parse_args() -> argparse.Namespace:
	parser = argparse.ArgumentParser(
		description="Build normalized country statistics from UNSD National Accounts Main Aggregates."
	)
	parser.add_argument("--config", type=Path, default=DEFAULT_CONFIG)
	parser.add_argument("--providers", type=Path, default=DEFAULT_PROVIDERS)
	parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
	parser.add_argument("--registry", default=DEFAULT_REGISTRY)
	parser.add_argument("--timeout", type=int, default=120)
	return parser.parse_args()


def read_json(path: Path) -> Any:
	return json.loads(path.read_text(encoding="utf-8"))


def fetch_bytes(url: str, timeout: int, accept: str = "*/*") -> tuple[bytes, dict[str, str]]:
	delays = (0, 5, 15, 30)
	last_error: Exception | None = None
	for attempt, delay in enumerate(delays, start=1):
		if delay:
			time.sleep(delay)
		request = Request(url, headers={"Accept": accept, "User-Agent": USER_AGENT})
		try:
			with urlopen(request, timeout=timeout) as response:
				return response.read(), {key.lower(): value for key, value in response.headers.items()}
		except (HTTPError, URLError, TimeoutError) as error:
			last_error = error
			print(f"Request failed ({attempt}/{len(delays)}): {url}: {error}")
	if last_error is None:
		raise RuntimeError(f"Request failed without an exception: {url}")
	raise RuntimeError(f"Request failed after {len(delays)} attempts: {url}") from last_error


def fetch_json(source: str, timeout: int) -> Any:
	if not source.startswith(("https://", "http://")):
		return read_json(Path(source))
	data, _ = fetch_bytes(source, timeout, "application/json")
	return json.loads(data.decode("utf-8"))


def validate_classification(indicator_id: str, classification: Any) -> None:
	if not isinstance(classification, dict) or classification.get("type") != "fixed":
		raise ValueError(f"Indicator {indicator_id} requires fixed classification metadata.")
	scale = str(classification.get("scale", "linear"))
	if scale not in ("linear", "logarithmic"):
		raise ValueError(f"Indicator {indicator_id} has invalid classification scale: {scale}.")
	breaks = classification.get("breaks")
	if not isinstance(breaks, list) or len(breaks) != 6:
		raise ValueError(f"Indicator {indicator_id} requires exactly 6 classification breaks.")
	numbers = [float(value) for value in breaks]
	if not all(math.isfinite(value) for value in numbers):
		raise ValueError(f"Indicator {indicator_id} has non-finite classification breaks.")
	if any(numbers[index] <= numbers[index - 1] for index in range(1, len(numbers))):
		raise ValueError(f"Indicator {indicator_id} classification breaks must be strictly ascending.")
	if scale == "logarithmic" and numbers[0] <= 0:
		raise ValueError(f"Indicator {indicator_id} logarithmic classification requires positive breaks.")


def validate_config(payload: Any) -> dict[str, Any]:
	if not isinstance(payload, dict) or payload.get("schema") != "kartensammlung.unsd-national-accounts-statistics/v1":
		raise ValueError("Invalid UNSD national accounts statistics config.")
	for key in ("sourcePage", "termsUrl"):
		if not str(payload.get(key, "")).startswith("https://"):
			raise ValueError(f"UNSD {key} must use HTTPS.")
	start_year = payload.get("expectedStartYear")
	latest_year = payload.get("minimumLatestYear")
	if not isinstance(start_year, int) or not 1900 <= start_year <= 2200:
		raise ValueError("UNSD expectedStartYear is invalid.")
	if not isinstance(latest_year, int) or latest_year < start_year:
		raise ValueError("UNSD minimumLatestYear is invalid.")

	provider = payload.get("provider")
	if not isinstance(provider, dict) or provider.get("id") != "unsd-national-accounts":
		raise ValueError("UNSD provider metadata is invalid.")
	for key in ("name", "dataset", "license", "licenseUrl", "attribution"):
		if not str(provider.get(key, "")).strip():
			raise ValueError(f"UNSD provider metadata is missing {key}.")

	ignored = payload.get("ignoredSourceCodes")
	if not isinstance(ignored, list) or not ignored or any(not re.fullmatch(r"\d{3}", str(code)) for code in ignored):
		raise ValueError("UNSD ignoredSourceCodes must be a non-empty list of three-digit codes.")
	if len(ignored) != len(set(ignored)):
		raise ValueError("UNSD ignoredSourceCodes contains duplicates.")

	aliases = payload.get("countryIdAliases")
	if not isinstance(aliases, dict):
		raise ValueError("UNSD countryIdAliases must be an object.")
	for code, area_id in aliases.items():
		if not re.fullmatch(r"\d{3}", str(code)) or not str(area_id).startswith("country:"):
			raise ValueError(f"Invalid UNSD country id alias: {code!r} -> {area_id!r}")

	compositions = payload.get("compositions")
	if not isinstance(compositions, dict) or set(compositions) != {"country:TZA"}:
		raise ValueError("UNSD config requires the reviewed country:TZA composition.")
	for area_id, composition in compositions.items():
		if not isinstance(composition, dict):
			raise ValueError(f"Invalid UNSD composition for {area_id}.")
		codes = composition.get("sourceCodes")
		if not isinstance(codes, list) or len(codes) != 2 or any(not re.fullmatch(r"\d{3}", str(code)) for code in codes):
			raise ValueError(f"Invalid UNSD source codes for {area_id}.")
		if not isinstance(composition.get("minimumYear"), int):
			raise ValueError(f"Invalid UNSD minimumYear for {area_id}.")
		if not str(composition.get("label", "")).strip():
			raise ValueError(f"UNSD composition {area_id} requires a label.")

	resources = payload.get("resources")
	if not isinstance(resources, dict) or set(resources) != {"gdp", "gdpPerCapita", "gni", "population"}:
		raise ValueError("UNSD resource configuration is incomplete.")
	file_ids: set[int] = set()
	for key, resource in resources.items():
		if not isinstance(resource, dict):
			raise ValueError(f"UNSD resource {key} must be an object.")
		file_id = resource.get("fileId")
		if not isinstance(file_id, int) or file_id <= 0 or file_id in file_ids:
			raise ValueError(f"UNSD resource {key} has an invalid or duplicate fileId.")
		file_ids.add(file_id)
		if not str(resource.get("downloadUrl", "")).startswith("https://"):
			raise ValueError(f"UNSD resource {key} downloadUrl must use HTTPS.")
		if not isinstance(resource.get("minimumBytes"), int) or resource["minimumBytes"] <= 0:
			raise ValueError(f"UNSD resource {key} minimumBytes is invalid.")
		selector = resource.get("rowSelector")
		if selector is not None:
			if not isinstance(selector, dict) or not str(selector.get("column", "")).strip() or not str(selector.get("value", "")).strip():
				raise ValueError(f"UNSD resource {key} rowSelector is invalid.")

	indicators = payload.get("indicators")
	if not isinstance(indicators, list) or len(indicators) != 3:
		raise ValueError("UNSD config requires exactly three migrated indicators.")
	ids: set[str] = set()
	slugs: set[str] = set()
	for indicator in indicators:
		if not isinstance(indicator, dict):
			raise ValueError("UNSD indicator entries must be objects.")
		for key in ("id", "slug", "resource", "sourceIndicator", "title", "description"):
			if not str(indicator.get(key, "")).strip():
				raise ValueError(f"UNSD indicator entry is missing {key}.")
		indicator_id = str(indicator["id"])
		slug = str(indicator["slug"])
		if indicator_id in ids or slug in slugs:
			raise ValueError(f"Duplicate UNSD indicator id or slug: {indicator_id}.")
		ids.add(indicator_id)
		slugs.add(slug)
		if indicator["resource"] not in resources:
			raise ValueError(f"Indicator {indicator_id} references an unknown resource.")
		unit = indicator.get("unit")
		if not isinstance(unit, dict) or not str(unit.get("id", "")).strip() or not str(unit.get("label", "")).strip():
			raise ValueError(f"Indicator {indicator_id} has invalid unit metadata.")
		validate_classification(indicator_id, indicator.get("classification"))
		for key in ("minAreasWithAnyValue", "minAreasInDefaultYear"):
			value = indicator.get(key)
			if not isinstance(value, int) or value <= 0:
				raise ValueError(f"Indicator {indicator_id} has invalid {key}.")
		required_areas = indicator.get("requiredAreas")
		if not isinstance(required_areas, list) or not required_areas or any(not str(area_id).startswith("country:") for area_id in required_areas):
			raise ValueError(f"Indicator {indicator_id} requires reviewed requiredAreas.")

	expected_ids = {"gdp.current-usd", "gdp-per-capita.current-usd", "gni.current-usd"}
	if ids != expected_ids:
		raise ValueError(f"Unexpected UNSD indicator set: {sorted(ids)}")
	return payload


def validate_provider_catalog(payload: Any, provider_id: str) -> dict[str, Any]:
	if not isinstance(payload, dict) or payload.get("schema") != "kartensammlung.statistics-providers/v1":
		raise ValueError("Invalid statistics provider catalog.")
	providers = payload.get("providers")
	if not isinstance(providers, list) or not providers:
		raise ValueError("Statistics provider catalog is empty.")
	ids = [str(item.get("id", "")) for item in providers if isinstance(item, dict)]
	if len(ids) != len(set(ids)):
		raise ValueError("Duplicate statistics provider id.")
	if provider_id not in ids:
		raise RuntimeError(f"Provider {provider_id} is missing from statistics-providers.json.")
	return payload


def load_registry(source: str, timeout: int) -> tuple[dict[str, str], set[str], dict[str, Any]]:
	payload = fetch_json(source, timeout)
	if not isinstance(payload, dict) or payload.get("schema") != "kartensammlung.area-registry/v1":
		raise ValueError("Invalid area registry.")
	areas = payload.get("areas")
	if not isinstance(areas, list):
		raise ValueError("Area registry has no areas.")
	area_by_m49: dict[str, str] = {}
	area_ids: set[str] = set()
	for area in areas:
		if not isinstance(area, dict) or area.get("level") != "country":
			continue
		area_id = str(area.get("area_id", "")).strip()
		if not area_id:
			continue
		if area_id in area_ids:
			raise ValueError(f"Duplicate area id in registry: {area_id}")
		area_ids.add(area_id)
		codes = area.get("codes") or {}
		m49 = str(codes.get("m49", "")).strip()
		if m49:
			m49 = m49.zfill(3)
			if not re.fullmatch(r"\d{3}", m49):
				raise ValueError(f"Invalid M49 code in area registry: {m49!r}")
			if m49 in area_by_m49:
				raise ValueError(f"Duplicate M49 code in area registry: {m49}")
			area_by_m49[m49] = area_id
	if len(area_ids) < 240 or len(area_by_m49) < 240:
		raise RuntimeError(
			f"Area registry unexpectedly small: {len(area_ids)} country areas, {len(area_by_m49)} M49 codes."
		)
	for area_id in ("country:AUT", "country:DEU", "country:USA", "country:IND", "country:XKX", "country:TZA"):
		if area_id not in area_ids:
			raise RuntimeError(f"Area registry sanity check failed: {area_id} is missing.")
	return area_by_m49, area_ids, payload


def verify_terms(config: dict[str, Any], timeout: int) -> dict[str, Any]:
	data, _ = fetch_bytes(str(config["termsUrl"]), timeout, "text/html,*/*")
	text = data.decode("utf-8", errors="replace")
	text = html.unescape(re.sub(r"<[^>]+>", " ", text))
	text = " ".join(text.lower().split())
	required_phrases = ("may be copied freely", "further distributed")
	missing = [phrase for phrase in required_phrases if phrase not in text]
	if missing:
		raise RuntimeError("UNdata terms no longer contain expected reuse language: " + ", ".join(missing))
	return {
		"url": config["termsUrl"],
		"verifiedReuseLanguage": list(required_phrases),
		"bytes": len(data),
		"sha256": hashlib.sha256(data).hexdigest(),
	}


def parse_country_id(value: Any) -> str | None:
	if value is None or isinstance(value, bool):
		return None
	if isinstance(value, int):
		integer = value
	elif isinstance(value, float) and value.is_integer():
		integer = int(value)
	else:
		text = str(value).strip()
		try:
			integer = int(float(text))
		except (TypeError, ValueError):
			return None
	if not 0 <= integer <= 999:
		return None
	return f"{integer:03d}"


def parse_year(value: Any) -> int | None:
	if value is None or isinstance(value, bool):
		return None
	if isinstance(value, int):
		return value if 1900 <= value <= 2200 else None
	if isinstance(value, float) and value.is_integer():
		integer = int(value)
		return integer if 1900 <= integer <= 2200 else None
	text = str(value).strip()
	if len(text) == 4 and text.isdigit():
		year = int(text)
		return year if 1900 <= year <= 2200 else None
	return None


def parse_number(value: Any, context: str) -> int | float | None:
	if value is None or value == "" or isinstance(value, bool):
		return None
	if isinstance(value, str) and not value.strip():
		return None
	try:
		number = float(value)
	except (TypeError, ValueError) as error:
		raise RuntimeError(f"Non-numeric UNSD value in {context}: {value!r}") from error
	if not math.isfinite(number):
		raise RuntimeError(f"Non-finite UNSD value in {context}: {value!r}")
	rounded = round(number, 6)
	if rounded == 0:
		rounded = 0.0
	return int(rounded) if rounded.is_integer() else rounded


def content_filename(headers: dict[str, str], fallback: str) -> str:
	disposition = str(headers.get("content-disposition", ""))
	match = re.search(r"filename\*?=(?:UTF-8''|\")?([^\";]+)", disposition, flags=re.IGNORECASE)
	if not match:
		return fallback
	return match.group(1).strip().strip('"') or fallback


def parse_workbook(
	resource_key: str,
	resource: dict[str, Any],
	data: bytes,
) -> tuple[dict[str, dict[str, Any]], dict[str, Any]]:
	if len(data) < int(resource["minimumBytes"]):
		raise RuntimeError(
			f"UNSD resource {resource_key} is unexpectedly small: {len(data)} < {resource['minimumBytes']} bytes."
		)
	if not data.startswith(b"PK"):
		raise RuntimeError(f"UNSD resource {resource_key} is no longer an XLSX/ZIP document.")
	workbook = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
	try:
		if not workbook.sheetnames:
			raise RuntimeError(f"UNSD resource {resource_key} contains no sheets.")
		sheet_names = list(workbook.sheetnames)
		sheet = workbook[sheet_names[0]]
		rows = sheet.iter_rows(values_only=True)
		header: tuple[Any, ...] | None = None
		for index, row in enumerate(rows, start=1):
			if index > 12:
				break
			first = str(row[0] or "").strip() if len(row) > 0 else ""
			second = str(row[1] or "").strip() if len(row) > 1 else ""
			if first == "CountryID" and second == "Country":
				header = row
				break
		if header is None:
			raise RuntimeError(f"UNSD resource {resource_key} has no CountryID/Country header row.")
		columns: dict[str, int] = {}
		year_columns: dict[int, int] = {}
		for index, raw in enumerate(header):
			name = str(raw or "").strip()
			if name:
				columns[name] = index
			year = parse_year(raw)
			if year is not None:
				year_columns[year] = index
		for required in ("CountryID", "Country"):
			if required not in columns:
				raise RuntimeError(f"UNSD resource {resource_key} is missing column {required}.")
		selector = resource.get("rowSelector")
		if selector is not None and str(selector["column"]) not in columns:
			raise RuntimeError(
				f"UNSD resource {resource_key} is missing selector column {selector['column']!r}."
			)
		if not year_columns:
			raise RuntimeError(f"UNSD resource {resource_key} contains no annual columns.")

		result: dict[str, dict[str, Any]] = {}
		for row in rows:
			if selector is not None:
				selector_index = columns[str(selector["column"])]
				actual = str(row[selector_index] or "").strip() if selector_index < len(row) else ""
				if actual != str(selector["value"]):
					continue
			code_index = columns["CountryID"]
			code = parse_country_id(row[code_index] if code_index < len(row) else None)
			if code is None:
				continue
			name_index = columns["Country"]
			country = str(row[name_index] or "").strip() if name_index < len(row) else ""
			values: dict[int, int | float] = {}
			for year, column_index in year_columns.items():
				raw = row[column_index] if column_index < len(row) else None
				value = parse_number(raw, f"{resource_key} {code} {country} {year}")
				if value is not None:
					values[year] = value
			if not values:
				continue
			if code in result:
				raise RuntimeError(f"Duplicate UNSD source row for {resource_key} CountryID {code}.")
			result[code] = {"country": country, "values": values}
	finally:
		workbook.close()
	return result, {
		"sheetNames": sheet_names,
		"sourceStartYear": min(year_columns),
		"sourceLatestYear": max(year_columns),
		"sourceRows": len(result),
	}


def map_source_rows(
	resource_key: str,
	rows: dict[str, dict[str, Any]],
	config: dict[str, Any],
	area_by_m49: dict[str, str],
	area_ids: set[str],
) -> tuple[dict[int, dict[str, int | float]], dict[str, dict[int, int | float]], list[dict[str, Any]]]:
	ignored = {str(code) for code in config["ignoredSourceCodes"]}
	aliases = {str(code): str(area_id) for code, area_id in config["countryIdAliases"].items()}
	component_codes = {
		str(code)
		for composition in config["compositions"].values()
		for code in composition["sourceCodes"]
	}
	for area_id in aliases.values():
		if area_id not in area_ids:
			raise RuntimeError(f"UNSD alias target is missing from registry: {area_id}")

	missing_ignored = sorted(ignored - set(rows))
	if missing_ignored:
		raise RuntimeError(
			f"Reviewed historical UNSD source codes disappeared from {resource_key}; review before continuing: "
			+ ", ".join(missing_ignored)
		)

	values: dict[int, dict[str, int | float]] = defaultdict(dict)
	components: dict[str, dict[int, int | float]] = {}
	unexpected: list[dict[str, Any]] = []
	for code, row in rows.items():
		if code in ignored:
			continue
		if code in component_codes:
			components[code] = dict(row["values"])
			continue
		area_id = aliases.get(code) or area_by_m49.get(code)
		if area_id is None:
			unexpected.append(
				{
					"code": code,
					"country": row["country"],
					"startYear": min(row["values"]),
					"latestYear": max(row["values"]),
				}
			)
			continue
		for year, value in row["values"].items():
			if area_id in values[year]:
				raise RuntimeError(f"Duplicate mapped UNSD value for {resource_key} {year} {area_id}.")
			values[year][area_id] = value
	if unexpected:
		details = ", ".join(
			f"{entry['code']} {entry['country']} ({entry['startYear']}-{entry['latestYear']})"
			for entry in unexpected
		)
		raise RuntimeError(f"Unreviewed UNSD CountryID row(s) in {resource_key}: {details}")
	missing_components = sorted(component_codes - set(components))
	if missing_components:
		raise RuntimeError(f"UNSD resource {resource_key} is missing composition components: {', '.join(missing_components)}")
	return dict(values), components, unexpected


def add_sum_composition(
	values_by_year: dict[int, dict[str, int | float]],
	components: dict[str, dict[int, int | float]],
	composition: dict[str, Any],
	target_area: str,
) -> list[int]:
	codes = [str(code) for code in composition["sourceCodes"]]
	minimum_year = int(composition["minimumYear"])
	common_years = set(components[codes[0]])
	for code in codes[1:]:
		common_years &= set(components[code])
	used_years: list[int] = []
	for year in sorted(common_years):
		if year < minimum_year:
			continue
		if target_area in values_by_year.setdefault(year, {}):
			raise RuntimeError(f"UNSD composition would overwrite {target_area} in {year}.")
		total = sum(float(components[code][year]) for code in codes)
		rounded = round(total, 6)
		values_by_year[year][target_area] = int(rounded) if rounded.is_integer() else rounded
		used_years.append(year)
	if not used_years or used_years[0] != minimum_year:
		raise RuntimeError(
			f"UNSD {target_area} sum composition does not begin in reviewed year {minimum_year}: "
			f"{used_years[0] if used_years else 'none'}"
		)
	return used_years


def add_tanzania_gdp_per_capita(
	values_by_year: dict[int, dict[str, int | float]],
	gdp_components: dict[str, dict[int, int | float]],
	population_components: dict[str, dict[int, int | float]],
	composition: dict[str, Any],
	target_area: str,
) -> list[int]:
	codes = [str(code) for code in composition["sourceCodes"]]
	minimum_year = int(composition["minimumYear"])
	common_years = set(gdp_components[codes[0]]) & set(population_components[codes[0]])
	for code in codes[1:]:
		common_years &= set(gdp_components[code]) & set(population_components[code])
	used_years: list[int] = []
	for year in sorted(common_years):
		if year < minimum_year:
			continue
		population = sum(float(population_components[code][year]) for code in codes)
		if population <= 0:
			raise RuntimeError(f"UNSD Tanzania population is not positive in {year}: {population}")
		gdp = sum(float(gdp_components[code][year]) for code in codes)
		value = round(gdp / population, 6)
		if target_area in values_by_year.setdefault(year, {}):
			raise RuntimeError(f"UNSD GDP-per-capita composition would overwrite {target_area} in {year}.")
		values_by_year[year][target_area] = int(value) if value.is_integer() else value
		used_years.append(year)
	if not used_years or used_years[0] != minimum_year:
		raise RuntimeError(
			f"UNSD {target_area} GDP-per-capita composition does not begin in reviewed year {minimum_year}: "
			f"{used_years[0] if used_years else 'none'}"
		)
	return used_years


def normalize_year_values(values_by_year: dict[int, dict[str, int | float]]) -> dict[int, dict[str, int | float]]:
	return {
		year: {area_id: year_values[area_id] for area_id in sorted(year_values)}
		for year, year_values in sorted(values_by_year.items())
		if year_values
	}


def build_indicator(
	config: dict[str, Any],
	indicator: dict[str, Any],
	values_by_year: dict[int, dict[str, int | float]],
	registry_area_count: int,
	resource_metadata: dict[str, dict[str, Any]],
	composition_metadata: dict[str, Any],
) -> dict[str, Any]:
	indicator_id = str(indicator["id"])
	values_by_year = normalize_year_values(values_by_year)
	if not values_by_year:
		raise RuntimeError(f"UNSD returned no mapped values for {indicator_id}.")
	available_years = sorted(values_by_year)
	if available_years[0] != int(config["expectedStartYear"]):
		raise RuntimeError(
			f"UNSD {indicator_id} starts in {available_years[0]}, expected {config['expectedStartYear']}."
		)
	latest_year = available_years[-1]
	if latest_year < int(config["minimumLatestYear"]):
		raise RuntimeError(
			f"UNSD {indicator_id} latest year is unexpectedly old: {latest_year} < {config['minimumLatestYear']}."
		)
	areas = {area_id for year_values in values_by_year.values() for area_id in year_values}
	if len(areas) < int(indicator["minAreasWithAnyValue"]):
		raise RuntimeError(f"UNSD coverage too small for {indicator_id}: {len(areas)} areas.")
	broad_years = [
		year
		for year in available_years
		if len(values_by_year[year]) >= int(indicator["minAreasInDefaultYear"])
	]
	if not broad_years:
		raise RuntimeError(f"UNSD has no broadly covered year for {indicator_id}.")
	default_year = broad_years[-1]
	if default_year < int(config["minimumLatestYear"]):
		raise RuntimeError(
			f"UNSD latest broadly covered year is unexpectedly old for {indicator_id}: {default_year}."
		)
	for area_id in indicator["requiredAreas"]:
		if area_id not in values_by_year[default_year]:
			raise RuntimeError(f"Required area {area_id} is missing in default year {default_year} for {indicator_id}.")

	if indicator_id == "gdp-per-capita.current-usd":
		resource_keys = ["gdpPerCapita", "gdp", "population"]
	else:
		resource_keys = [str(indicator["resource"])]
	resources = [resource_metadata[key] for key in resource_keys]
	provider = config["provider"]
	return {
		"schema": "kartensammlung.statistics-indicator/v1",
		"indicator": {
			"id": indicator_id,
			"title": indicator["title"],
			"description": indicator["description"],
			"areaLevel": "country",
			"frequency": "annual",
			"unit": indicator["unit"],
			"classification": indicator["classification"],
		},
		"source": {
			"providerId": provider["id"],
			"providerName": provider["name"],
			"dataset": provider["dataset"],
			"indicator": indicator["sourceIndicator"],
			"license": provider["license"],
			"licenseUrl": provider["licenseUrl"],
			"attribution": provider["attribution"],
			"url": config["sourcePage"],
			"resources": [
				{
					"fileId": item["fileId"],
					"downloadUrl": item["downloadUrl"],
					"sha256": item["sha256"],
				}
				for item in resources
			],
			"countryIdentifier": "UNSD AMA CountryID / M49",
			"ignoredHistoricalCountryIds": list(config["ignoredSourceCodes"]),
			"explicitCountryIdAliases": config["countryIdAliases"],
			"compositions": composition_metadata,
			"fallback": None,
			"interpolation": None,
		},
		"availableYears": available_years,
		"defaultYear": default_year,
		"coverage": {
			"registryAreas": registry_area_count,
			"areasWithAnyValue": len(areas),
			"latestYear": latest_year,
			"areasInLatestYear": len(values_by_year[latest_year]),
			"broadCoverageThreshold": int(indicator["minAreasInDefaultYear"]),
			"areasInDefaultYear": len(values_by_year[default_year]),
		},
		"values": {str(year): values_by_year[year] for year in available_years},
	}


def canonical_bytes(payload: Any) -> bytes:
	return json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")


def write_json(path: Path, payload: Any) -> None:
	path.parent.mkdir(parents=True, exist_ok=True)
	path.write_text(json.dumps(payload, ensure_ascii=False, indent="\t") + "\n", encoding="utf-8")


def main() -> None:
	args = parse_args()
	config = validate_config(read_json(args.config))
	provider = config["provider"]
	provider_catalog = validate_provider_catalog(read_json(args.providers), str(provider["id"]))
	area_by_m49, area_ids, registry = load_registry(args.registry, args.timeout)
	terms_metadata = verify_terms(config, args.timeout)
	now = datetime.now(timezone.utc)

	aliases = {str(code): str(area_id) for code, area_id in config["countryIdAliases"].items()}
	if set(aliases) != {"412"} or aliases["412"] != "country:XKX":
		raise RuntimeError("UNSD Kosovo mapping changed from reviewed CountryID 412 -> country:XKX.")

	resource_rows: dict[str, dict[str, dict[str, Any]]] = {}
	resource_metadata: dict[str, dict[str, Any]] = {}
	for resource_key, resource in config["resources"].items():
		print(f"Downloading UNSD resource {resource_key} (file {resource['fileId']})")
		data, headers = fetch_bytes(
			str(resource["downloadUrl"]),
			args.timeout,
			"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet,application/octet-stream,*/*",
		)
		rows, workbook_metadata = parse_workbook(resource_key, resource, data)
		if workbook_metadata["sourceStartYear"] != int(config["expectedStartYear"]):
			raise RuntimeError(
				f"UNSD resource {resource_key} starts in {workbook_metadata['sourceStartYear']}, "
				f"expected {config['expectedStartYear']}."
			)
		if workbook_metadata["sourceLatestYear"] < int(config["minimumLatestYear"]):
			raise RuntimeError(
				f"UNSD resource {resource_key} latest year is unexpectedly old: "
				f"{workbook_metadata['sourceLatestYear']}."
			)
		resource_rows[resource_key] = rows
		resource_metadata[resource_key] = {
			"key": resource_key,
			"fileId": resource["fileId"],
			"downloadUrl": resource["downloadUrl"],
			"filename": content_filename(headers, f"unsd-ama-file-{resource['fileId']}.xlsx"),
			"bytes": len(data),
			"sha256": hashlib.sha256(data).hexdigest(),
			**workbook_metadata,
		}
		print(
			f"{resource_key}: rows={workbook_metadata['sourceRows']} "
			f"years={workbook_metadata['sourceStartYear']}-{workbook_metadata['sourceLatestYear']} "
			f"bytes={len(data)}"
		)

	mapped: dict[str, dict[int, dict[str, int | float]]] = {}
	components: dict[str, dict[str, dict[int, int | float]]] = {}
	for resource_key in ("gdp", "gdpPerCapita", "gni"):
		mapped_values, component_values, _ = map_source_rows(
			resource_key,
			resource_rows[resource_key],
			config,
			area_by_m49,
			area_ids,
		)
		mapped[resource_key] = mapped_values
		components[resource_key] = component_values

	population_rows = resource_rows["population"]
	population_components: dict[str, dict[int, int | float]] = {}
	tanzania_codes = [str(code) for code in config["compositions"]["country:TZA"]["sourceCodes"]]
	for code in tanzania_codes:
		row = population_rows.get(code)
		if row is None:
			raise RuntimeError(f"UNSD population resource is missing Tanzania component {code}.")
		population_components[code] = dict(row["values"])

	composition = config["compositions"]["country:TZA"]
	gdp_years = add_sum_composition(mapped["gdp"], components["gdp"], composition, "country:TZA")
	gni_years = add_sum_composition(mapped["gni"], components["gni"], composition, "country:TZA")
	gdp_pc_years = add_tanzania_gdp_per_capita(
		mapped["gdpPerCapita"],
		components["gdp"],
		population_components,
		composition,
		"country:TZA",
	)
	if gdp_years != gni_years or gdp_years != gdp_pc_years:
		raise RuntimeError("UNSD Tanzania composition year coverage differs between GDP, GDP per capita and GNI.")

	composition_metadata = {
		"country:TZA": {
			"label": composition["label"],
			"sourceCountryIds": tanzania_codes,
			"startYear": gdp_years[0],
			"endYear": gdp_years[-1],
			"gdpAndGni": "sum of Mainland and Zanzibar",
			"gdpPerCapita": "(Mainland GDP + Zanzibar GDP) / (Mainland population + Zanzibar population)",
		},
	}

	indicator_values = {
		"gdp.current-usd": mapped["gdp"],
		"gdp-per-capita.current-usd": mapped["gdpPerCapita"],
		"gni.current-usd": mapped["gni"],
	}
	payloads: list[tuple[dict[str, Any], dict[str, Any]]] = []
	for indicator in config["indicators"]:
		indicator_id = str(indicator["id"])
		payload = build_indicator(
			config,
			indicator,
			indicator_values[indicator_id],
			len(area_ids),
			resource_metadata,
			composition_metadata,
		)
		payloads.append((indicator, payload))
		coverage = payload["coverage"]
		print(
			f"{indicator_id}: mapped={coverage['areasWithAnyValue']} "
			f"years={payload['availableYears'][0]}-{payload['availableYears'][-1]} "
			f"defaultYear={payload['defaultYear']} defaultCoverage={coverage['areasInDefaultYear']}"
		)

	hasher = hashlib.sha256()
	for indicator, payload in sorted(payloads, key=lambda item: item[0]["id"]):
		hasher.update(str(indicator["id"]).encode("utf-8") + b"\0" + canonical_bytes(payload) + b"\0")
	snapshot = hasher.hexdigest()[:16]
	provider_dir = args.output_dir / str(provider["id"])
	release_dir = provider_dir / "releases" / snapshot
	index_indicators: list[dict[str, Any]] = []
	for indicator, payload in payloads:
		filename = f"{indicator['slug']}.json"
		write_json(release_dir / filename, payload)
		index_indicators.append(
			{
				"id": indicator["id"],
				"title": indicator["title"],
				"description": indicator["description"],
				"areaLevel": "country",
				"frequency": "annual",
				"unit": indicator["unit"],
				"classification": indicator["classification"],
				"sourceIndicator": indicator["sourceIndicator"],
				"path": f"releases/{snapshot}/{filename}",
				"availableYears": payload["availableYears"],
				"defaultYear": payload["defaultYear"],
				"coverage": payload["coverage"],
			}
		)

	registry_source = {
		"url": args.registry if args.registry.startswith(("https://", "http://")) else None,
		"schema": registry.get("schema"),
		"generatedAt": registry.get("generatedAt"),
		"areaCount": len(area_ids),
		"m49AreaCount": len(area_by_m49),
	}
	if registry_source["url"] is None:
		registry_source.pop("url")
	provider_index = {
		"schema": "kartensammlung.statistics-provider-index/v1",
		"provider": provider,
		"retrievedAt": now.isoformat().replace("+00:00", "Z"),
		"activeSnapshot": snapshot,
		"areaRegistry": registry_source,
		"dataset": {
			"sourcePage": config["sourcePage"],
			"expectedStartYear": config["expectedStartYear"],
			"minimumLatestYear": config["minimumLatestYear"],
			"terms": terms_metadata,
			"resources": [resource_metadata[key] for key in sorted(resource_metadata)],
			"ignoredHistoricalCountryIds": list(config["ignoredSourceCodes"]),
			"explicitCountryIdAliases": config["countryIdAliases"],
			"compositions": composition_metadata,
		},
		"indicators": index_indicators,
	}
	write_json(provider_dir / "index.json", provider_index)
	write_json(
		args.output_dir / "index.json",
		{
			"schema": "kartensammlung.statistics-index/v1",
			"areaRegistry": "../area-registry-countries.json",
			"providers": provider_catalog["providers"],
		},
	)
	print(f"Built UNSD national accounts snapshot {snapshot}")
	print(f"Indicators: {len(index_indicators)}")
	print(f"Tanzania composition years: {gdp_years[0]}-{gdp_years[-1]}")


if __name__ == "__main__":
	main()
