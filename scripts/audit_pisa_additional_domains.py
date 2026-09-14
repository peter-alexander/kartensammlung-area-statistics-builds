#!/usr/bin/env python3
from __future__ import annotations

import hashlib
import io
import zipfile
from urllib.request import Request, urlopen

import xlrd
from openpyxl import load_workbook

# Temporary source audit. Remove this file before the production PR.
SOURCES = {
	"pisa-2025-core-and-digital": "https://stat.link/mrq53f",
	"pisa-2022-financial-literacy": "https://stat.link/4nx1lb",
	"pisa-2022-creative-thinking": "https://stat.link/opbe7a",
	"pisa-2018-global-competence": "https://doi.org/10.1787/888934171229",
	"pisa-2015-collaborative-problem-solving": "https://doi.org/10.1787/888933616769",
	"pisa-2012-computer-based-mathematics": "https://doi.org/10.1787/888932935781",
	"pisa-2012-digital-reading": "https://doi.org/10.1787/888932935781",
	"pisa-2012-creative-problem-solving": "https://doi.org/10.1787/888933003668",
	"pisa-2009-digital-reading": "https://doi.org/10.1787/888932436556",
	"pisa-2003-cross-curricular-problem-solving": "https://doi.org/10.1787/402381481733",
}

TITLE_PATTERNS = {
	"pisa-2025-core-and-digital": ("mean score and variation in computational problem-solving performance",),
	"pisa-2022-financial-literacy": ("mean financial literacy scores in 2012, 2015, 2018 and 2022",),
	"pisa-2022-creative-thinking": ("mean score and variation in creative thinking performance",),
	"pisa-2018-global-competence": ("performance on the global competence test",),
	"pisa-2015-collaborative-problem-solving": ("mean score and variation in collaborative problem-solving performance",),
	"pisa-2012-computer-based-mathematics": (
		"mean score, variation and gender differences in student performance on the computer-based mathematics scale",
	),
	"pisa-2012-digital-reading": (
		"mean score, variation and gender differences in student performance on the digital reading scale",
	),
	"pisa-2012-creative-problem-solving": ("mean score and variation in student performance in problem solving",),
	"pisa-2009-digital-reading": ("mean score", "digital", "reading"),
	"pisa-2003-cross-curricular-problem-solving": ("problem-solving",),
}


def fetch(url: str) -> tuple[bytes, str, str]:
	request = Request(url, headers={"User-Agent": "kartensammlung-pisa-source-audit/1"})
	with urlopen(request, timeout=120) as response:
		return response.read(), response.geturl(), response.headers.get("Content-Type", "")


def normalized_text(value: object) -> str:
	return " ".join(str(value).replace("\n", " ").split()).lower()


def xlsx_rows(raw: bytes) -> dict[str, list[list[object]]]:
	workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
	return {
		sheet_name: [list(row) for row in workbook[sheet_name].iter_rows(values_only=True)]
		for sheet_name in workbook.sheetnames
	}


def xls_rows(raw: bytes) -> dict[str, list[list[object]]]:
	workbook = xlrd.open_workbook(file_contents=raw, on_demand=True)
	result: dict[str, list[list[object]]] = {}
	for sheet_name in workbook.sheet_names():
		sheet = workbook.sheet_by_name(sheet_name)
		result[sheet_name] = [sheet.row_values(row_index) for row_index in range(sheet.nrows)]
	return result


def find_target_sheets(name: str, sheets: dict[str, list[list[object]]]) -> list[str]:
	patterns = TITLE_PATTERNS[name]
	matches: list[str] = []
	for sheet_name, rows in sheets.items():
		search_text = " ".join(
			normalized_text(value)
			for row in rows[:25]
			for value in row
			if value not in (None, "")
		)
		if all(pattern in search_text for pattern in patterns):
			matches.append(sheet_name)
	return matches


def print_target_table(name: str, sheets: dict[str, list[list[object]]]) -> None:
	print(f"sheet names: {list(sheets)}")
	matches = find_target_sheets(name, sheets)
	print(f"target sheets: {matches}")
	for sheet_name in matches:
		print(f"--- {sheet_name} ---")
		rows = sheets[sheet_name]
		for row_number, row in enumerate(rows, start=1):
			values = [value for value in row if value not in (None, "")]
			if not values:
				continue
			if row_number <= 20 or isinstance(values[0], str):
				print(f"row {row_number}: {values[:32]!r}")


def main() -> None:
	for name, url in SOURCES.items():
		print(f"\n===== {name} =====")
		print(f"requested: {url}")
		try:
			raw, effective_url, content_type = fetch(url)
		except Exception as error:
			print(f"ERROR: {type(error).__name__}: {error}")
			continue
		print(f"effective: {effective_url}")
		print(f"content-type: {content_type}")
		print(f"bytes: {len(raw)}")
		print(f"sha256: {hashlib.sha256(raw).hexdigest()}")
		print(f"magic: {raw[:16]!r}")
		try:
			if raw.startswith(b"PK") and zipfile.is_zipfile(io.BytesIO(raw)):
				sheets = xlsx_rows(raw)
			elif raw.startswith(b"\xd0\xcf\x11\xe0"):
				sheets = xls_rows(raw)
			else:
				print(f"unsupported source preview: {raw[:500]!r}")
				continue
			print_target_table(name, sheets)
		except Exception as error:
			print(f"PARSE ERROR: {type(error).__name__}: {error}")


if __name__ == "__main__":
	main()
