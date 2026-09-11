#!/usr/bin/env python3
from __future__ import annotations

import argparse
import hashlib
import json
import math
import shutil
import tempfile
import time
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from urllib.error import HTTPError, URLError
from urllib.request import Request, urlopen

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_CONFIG = ROOT / "config" / "unodc-indicators.json"
DEFAULT_PROVIDERS = ROOT / "config" / "statistics-providers.json"
DEFAULT_OUTPUT_DIR = ROOT / "dist" / "statistics"
DEFAULT_REGISTRY = "https://tiles.radlobby.at/AreaStatistics/area-registry-countries.json"
USER_AGENT = "kartensammlung-area-statistics-builds/1"


def parse_args() -> argparse.Namespace:
	parser = argparse.ArgumentParser(description="Build normalized country statistics from UNODC UN-CTS data.")
	parser.add_argument("--config", type=Path, default=DEFAULT_CONFIG)
	parser.add_argument("--providers", type=Path, default=DEFAULT_PROVIDERS)
	parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
	parser.add_argument("--registry", default=DEFAULT_REGISTRY)
	parser.add_argument("--timeout", type=int, default=90)
	return parser.parse_args()


def read_json(path: Path) -> Any:
	return json.loads(path.read_text(encoding="utf-8"))


def fetch_json(source: str, timeout: int) -> Any:
	if not source.startswith(("https://", "http://")):
		return read_json(Path(source))
	last_error: Exception | None = None
	for attempt, delay in enumerate((0, 5, 15, 30, 60), start=1):
		if delay:
			time.sleep(delay)
		try:
			request = Request(source, headers={"Accept": "application/json", "User-Agent": USER_AGENT})
			with urlopen(request, timeout=timeout) as response:
				return json.load(response)
		except (HTTPError, URLError, TimeoutError, json.JSONDecodeError) as error:
			last_error = error
			print(f"Request failed ({attempt}/5): {source}: {error}")
	raise RuntimeError(f"Request failed after 5 attempts: {source}") from last_error


def download_xlsx(urls: list[str], timeout: int) -> tuple[Path, str]:
	last_error: Exception | None = None
	for url in urls:
		for attempt, delay in enumerate((0, 5, 15), start=1):
			if delay:
				time.sleep(delay)
			handle = tempfile.NamedTemporaryFile(prefix="unodc-", suffix=".xlsx", delete=False)
			path = Path(handle.name)
			handle.close()
			try:
				request = Request(url, headers={
					"Accept": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet, application/octet-stream, */*",
					"User-Agent": USER_AGENT,
				})
				with urlopen(request, timeout=timeout) as response, path.open("wb") as output:
					shutil.copyfileobj(response, output, length=1024 * 1024)
				if path.stat().st_size < 10_000:
					raise RuntimeError(f"Downloaded file is unexpectedly small: {path.stat().st_size} bytes")
				if not zipfile.is_zipfile(path):
					raise RuntimeError("Downloaded UNODC file is not a valid XLSX/ZIP archive")
				workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
				try:
					if not workbook.sheetnames:
						raise RuntimeError("Downloaded UNODC workbook contains no worksheets")
				finally:
					workbook.close()
				print(f"Using UNODC source: {url}")
				return path, url
			except (HTTPError, URLError, TimeoutError, OSError, RuntimeError, zipfile.BadZipFile) as error:
				last_error = error
				path.unlink(missing_ok=True)
				print(f"UNODC download/validation failed ({attempt}/3): {url}: {error}")
	raise RuntimeError("Unable to download a valid UNODC workbook") from last_error


def validate_config(payload: Any) -> dict[str, Any]:
	if not isinstance(payload, dict) or payload.get("schema") != "kartensammlung.unodc-statistics/v1":
		raise ValueError("Invalid UNODC statistics config.")
	urls = payload.get("downloadUrls")
	if not isinstance(urls, list) or not urls or not all(str(url).startswith("https://") for url in urls):
		raise ValueError("UNODC downloadUrls must be a non-empty HTTPS URL list.")
	provider = payload.get("provider")
	if not isinstance(provider, dict):
		raise ValueError("UNODC provider metadata is required.")
	for key in ("id", "name", "dataset", "license", "licenseUrl", "attribution"):
		if not str(provider.get(key, "")).strip():
			raise ValueError(f"UNODC provider metadata is missing {key}.")
	fraction = payload.get("broadCoverageFraction")
	if not isinstance(fraction, (int, float)) or not 0 < float(fraction) <= 1:
		raise ValueError("Invalid broadCoverageFraction.")
	indicators = payload.get("indicators")
	if not isinstance(indicators, list) or not indicators:
		raise ValueError("UNODC config requires indicators.")
	ids: set[str] = set()
	slugs: set[str] = set()
	for indicator in indicators:
		if not isinstance(indicator, dict):
			raise ValueError("UNODC indicator entries must be objects.")
		for key in ("id", "slug", "title", "description"):
			if not str(indicator.get(key, "")).strip():
				raise ValueError(f"UNODC indicator entry is missing {key}.")
		indicator_id = str(indicator["id"])
		slug = str(indicator["slug"])
		if indicator_id in ids or slug in slugs:
			raise ValueError(f"Duplicate UNODC indicator id or slug: {indicator_id}.")
		ids.add(indicator_id)
		slugs.add(slug)
		if not isinstance(indicator.get("filters"), dict):
			raise ValueError(f"Indicator {indicator_id} requires filters.")
		if not isinstance(indicator.get("unit"), dict):
			raise ValueError(f"Indicator {indicator_id} requires unit metadata.")
		classification = indicator.get("classification")
		if not isinstance(classification, dict) or classification.get("type") != "fixed":
			raise ValueError(f"Indicator {indicator_id} requires fixed classification.")
		breaks = classification.get("breaks")
		if not isinstance(breaks, list) or len(breaks) != 6:
			raise ValueError(f"Indicator {indicator_id} requires exactly six breaks.")
		numbers = [float(value) for value in breaks]
		if any(not math.isfinite(value) for value in numbers) or any(numbers[index] <= numbers[index - 1] for index in range(1, 6)):
			raise ValueError(f"Indicator {indicator_id} has invalid breaks.")
		if classification.get("scale") == "logarithmic" and numbers[0] <= 0:
			raise ValueError(f"Indicator {indicator_id} logarithmic breaks must be positive.")
	return payload


def validate_provider_catalog(payload: Any) -> dict[str, Any]:
	if not isinstance(payload, dict) or payload.get("schema") != "kartensammlung.statistics-providers/v1":
		raise ValueError("Invalid statistics provider catalog.")
	providers = payload.get("providers")
	if not isinstance(providers, list) or not providers:
		raise ValueError("Statistics provider catalog is empty.")
	ids = [str(provider.get("id", "")) for provider in providers if isinstance(provider, dict)]
	if len(ids) != len(set(ids)):
		raise ValueError("Duplicate statistics provider id.")
	return payload


def load_registry(source: str, timeout: int) -> tuple[dict[str, str], dict[str, Any]]:
	payload = fetch_json(source, timeout)
	if not isinstance(payload, dict) or payload.get("schema") != "kartensammlung.area-registry/v1":
		raise ValueError("Invalid area registry.")
	area_by_iso3: dict[str, str] = {}
	for area in payload.get("areas", []):
		if not isinstance(area, dict) or area.get("level") != "country":
			continue
		codes = area.get("codes")
		if not isinstance(codes, dict):
			continue
		iso3 = str(codes.get("iso3", "")).strip().upper()
		area_id = str(area.get("area_id", "")).strip()
		if len(iso3) == 3 and area_id:
			if iso3 in area_by_iso3:
				raise ValueError(f"Duplicate ISO3 in area registry: {iso3}")
			area_by_iso3[iso3] = area_id
	if len(area_by_iso3) < 240:
		raise RuntimeError(f"Area registry unexpectedly small: {len(area_by_iso3)} country areas.")
	for iso3 in ("AUT", "DEU", "USA", "IND", "XKX"):
		if iso3 not in area_by_iso3:
			raise RuntimeError(f"Area registry sanity check failed: {iso3} is missing.")
	return area_by_iso3, payload


def text(value: Any) -> str:
	return "" if value is None else " ".join(str(value).strip().split())


def parse_year(value: Any) -> int | None:
	try:
		year = int(float(value))
	except (TypeError, ValueError):
		return None
	return year if 1900 <= year <= 2200 else None


def parse_number(value: Any) -> int | float | None:
	if value is None or isinstance(value, bool):
		return None
	try:
		number = float(value)
	except (TypeError, ValueError):
		return None
	if not math.isfinite(number):
		return None
	if number == 0:
		number = 0.0
	return int(number) if number.is_integer() else number


def find_data_sheet(workbook: Any) -> str:
	candidates = [name for name in workbook.sheetnames if "intentional_homicide" in name.lower() and "reg_est" not in name.lower()]
	if len(candidates) != 1:
		raise RuntimeError(f"Expected one intentional-homicide sheet, found: {workbook.sheetnames}")
	return candidates[0]


def read_records(path: Path) -> list[dict[str, Any]]:
	workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
	try:
		worksheet = workbook[find_data_sheet(workbook)]
		records: list[dict[str, Any]] = []
		for row in worksheet.iter_rows(values_only=True):
			if len(row) < 12:
				continue
			iso3 = text(row[0]).upper()
			year = parse_year(row[9])
			value = parse_number(row[11])
			indicator = text(row[4])
			if len(iso3) != 3 or not iso3.isalpha() or year is None or value is None or not indicator:
				continue
			records.append({
				"iso3": iso3,
				"country": text(row[1]),
				"region": text(row[2]),
				"subregion": text(row[3]),
				"indicator": indicator,
				"dimension": text(row[5]),
				"category": text(row[6]),
				"sex": text(row[7]),
				"age": text(row[8]),
				"year": year,
				"unit": text(row[10]),
				"value": value,
			})
		return records
	finally:
		workbook.close()


def equal_text(actual: Any, expected: Any) -> bool:
	return str(actual).casefold() == str(expected).casefold()


def matches(record: dict[str, Any], filters: dict[str, Any]) -> bool:
	for key in ("indicator", "dimension", "category", "sex", "age", "unit"):
		if key in filters and not equal_text(record.get(key, ""), filters[key]):
			return False
	if "unitContains" in filters and str(filters["unitContains"]).casefold() not in str(record.get("unit", "")).casefold():
		return False
	return True


def print_diagnostics(records: list[dict[str, Any]]) -> None:
	print(f"UNODC parsed country-year rows: {len(records)}")
	for field in ("indicator", "dimension", "sex", "age", "unit"):
		values = sorted({str(record[field]) for record in records if record[field]})
		print(f"UNODC {field} values ({len(values)}): {values[:100]}")
	victims = [record for record in records if equal_text(record["indicator"], "Victims of intentional homicide")]
	categories = sorted({record["category"] for record in victims if record["category"]})
	print(f"UNODC homicide victim categories ({len(categories)}): {categories[:150]}")
	for dimension in sorted({record["dimension"] for record in victims if record["dimension"]}):
		dimension_categories = sorted({record["category"] for record in victims if record["dimension"] == dimension and record["category"]})
		print(f"UNODC categories for dimension {dimension!r}: {dimension_categories[:100]}")


def select_values(records: list[dict[str, Any]], indicator: dict[str, Any], area_by_iso3: dict[str, str]) -> tuple[dict[int, dict[str, int | float]], set[str]]:
	values: dict[int, dict[str, int | float]] = {}
	ignored: set[str] = set()
	matched = 0
	for record in records:
		if not matches(record, indicator["filters"]):
			continue
		matched += 1
		iso3 = record["iso3"]
		if iso3 not in area_by_iso3:
			ignored.add(iso3)
			continue
		year = int(record["year"])
		area_id = area_by_iso3[iso3]
		year_values = values.setdefault(year, {})
		value = record["value"]
		if area_id in year_values and year_values[area_id] != value:
			raise RuntimeError(f"Conflicting UNODC values for {indicator['id']} {year} {area_id}.")
		year_values[area_id] = value
	print(f"{indicator['id']}: matchedRows={matched}")
	return values, ignored


def build_indicator(config: dict[str, Any], indicator: dict[str, Any], values_by_year: dict[int, dict[str, int | float]], area_count: int, now: datetime, download_url: str) -> dict[str, Any]:
	indicator_id = str(indicator["id"])
	areas = {area_id for year_values in values_by_year.values() for area_id in year_values}
	if len(areas) < int(indicator["minAreasWithAnyValue"]):
		raise RuntimeError(f"Coverage too small for {indicator_id}: {len(areas)} areas.")
	for area_id in indicator["requiredAreas"]:
		if area_id not in areas:
			raise RuntimeError(f"Required area {area_id} has no values for {indicator_id}.")
	available_years = sorted(values_by_year)
	if not available_years:
		raise RuntimeError(f"UNODC returned no mapped values for {indicator_id}.")
	broad_threshold = max(1, math.ceil(len(areas) * float(config["broadCoverageFraction"])))
	broad_years = [year for year in available_years if year <= now.year and len(values_by_year[year]) >= broad_threshold]
	if not broad_years:
		raise RuntimeError(f"No year reaches broad coverage for {indicator_id}.")
	default_year = broad_years[-1]
	if now.year - default_year > int(indicator["maxDefaultYearAge"]):
		raise RuntimeError(f"Latest broadly covered year for {indicator_id} is unexpectedly old: {default_year}.")
	latest_year = available_years[-1]
	provider = config["provider"]
	return {
		"schema": "kartensammlung.statistics-indicator/v1",
		"indicator": {
			"id": indicator_id,
			"title": indicator["title"],
			"description": indicator["description"],
			"areaLevel": "country",
			"frequency": "annual",
			"unit": indicator["unit"],
			"classification": indicator["classification"],
		},
		"source": {
			"providerId": provider["id"],
			"providerName": provider["name"],
			"dataset": provider["dataset"],
			"datasetId": "UNODC-CTS-Intentional-Homicide",
			"indicator": indicator["filters"],
			"license": provider["license"],
			"licenseUrl": provider["licenseUrl"],
			"attribution": provider["attribution"],
			"url": config["sourcePage"],
			"downloadUrl": download_url,
		},
		"availableYears": available_years,
		"defaultYear": default_year,
		"coverage": {
			"registryAreas": area_count,
			"areasWithAnyValue": len(areas),
			"latestYear": latest_year,
			"areasInLatestYear": len(values_by_year[latest_year]),
			"broadCoverageThreshold": broad_threshold,
			"areasInDefaultYear": len(values_by_year[default_year]),
		},
		"values": {str(year): {area_id: values_by_year[year][area_id] for area_id in sorted(values_by_year[year])} for year in available_years},
	}


def canonical_bytes(payload: Any) -> bytes:
	return json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")


def write_json(path: Path, payload: Any) -> None:
	path.parent.mkdir(parents=True, exist_ok=True)
	path.write_text(json.dumps(payload, ensure_ascii=False, indent="\t") + "\n", encoding="utf-8")


def main() -> None:
	args = parse_args()
	config = validate_config(read_json(args.config))
	provider_catalog = validate_provider_catalog(read_json(args.providers))
	provider = config["provider"]
	if not any(item.get("id") == provider["id"] for item in provider_catalog["providers"]):
		raise RuntimeError(f"Provider {provider['id']} is missing from statistics-providers.json.")
	area_by_iso3, registry = load_registry(args.registry, args.timeout)
	now = datetime.now(timezone.utc)
	path, download_url = download_xlsx([str(url) for url in config["downloadUrls"]], args.timeout)
	try:
		print(f"Downloaded UNODC workbook: {path.stat().st_size} bytes")
		records = read_records(path)
	finally:
		path.unlink(missing_ok=True)
	if not records:
		raise RuntimeError("UNODC workbook contained no parseable country-year rows.")
	print_diagnostics(records)
	payloads: list[tuple[dict[str, Any], dict[str, Any]]] = []
	ignored_codes: set[str] = set()
	for indicator in config["indicators"]:
		values, ignored = select_values(records, indicator, area_by_iso3)
		ignored_codes.update(ignored)
		payload = build_indicator(config, indicator, values, len(area_by_iso3), now, download_url)
		payloads.append((indicator, payload))
		coverage = payload["coverage"]
		print(
			f"{indicator['id']}: mapped={coverage['areasWithAnyValue']} "
			f"years={payload['availableYears'][0]}-{payload['availableYears'][-1]} "
			f"defaultYear={payload['defaultYear']} defaultCoverage={coverage['areasInDefaultYear']}"
		)
	hasher = hashlib.sha256()
	for indicator, payload in sorted(payloads, key=lambda item: item[0]["id"]):
		hasher.update(str(indicator["id"]).encode("utf-8") + b"\0" + canonical_bytes(payload) + b"\0")
	snapshot = hasher.hexdigest()[:16]
	provider_dir = args.output_dir / provider["id"]
	release_dir = provider_dir / "releases" / snapshot
	index_indicators = []
	for indicator, payload in payloads:
		filename = f"{indicator['slug']}.json"
		write_json(release_dir / filename, payload)
		index_indicators.append({
			"id": indicator["id"],
			"title": indicator["title"],
			"description": indicator["description"],
			"areaLevel": "country",
			"frequency": "annual",
			"unit": indicator["unit"],
			"classification": indicator["classification"],
			"sourceIndicator": indicator["filters"],
			"path": f"releases/{snapshot}/{filename}",
			"availableYears": payload["availableYears"],
			"defaultYear": payload["defaultYear"],
			"coverage": payload["coverage"],
		})
	registry_source = {
		"url": args.registry if args.registry.startswith(("https://", "http://")) else None,
		"schema": registry.get("schema"),
		"generatedAt": registry.get("generatedAt"),
		"areaCount": len(area_by_iso3),
	}
	if registry_source["url"] is None:
		registry_source.pop("url")
	write_json(provider_dir / "index.json", {
		"schema": "kartensammlung.statistics-provider-index/v1",
		"provider": provider,
		"retrievedAt": now.isoformat().replace("+00:00", "Z"),
		"activeSnapshot": snapshot,
		"areaRegistry": registry_source,
		"dataset": {"downloadUrl": download_url, "sourcePage": config["sourcePage"]},
		"ignoredIso3Codes": sorted(ignored_codes),
		"indicators": index_indicators,
	})
	write_json(args.output_dir / "index.json", {
		"schema": "kartensammlung.statistics-index/v1",
		"areaRegistry": "../area-registry-countries.json",
		"providers": provider_catalog["providers"],
	})
	print(f"Built UNODC snapshot {snapshot}")
	print(f"Indicators: {len(index_indicators)}")
	print(f"Ignored ISO3 codes: {', '.join(sorted(ignored_codes)) or '(none)'}")


if __name__ == "__main__":
	main()
