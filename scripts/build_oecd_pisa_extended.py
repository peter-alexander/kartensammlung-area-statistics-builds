#!/usr/bin/env python3
from __future__ import annotations

import argparse
import hashlib
import io
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

import build_oecd_pisa as core
import build_oecd_pisa_additional as additional
import build_world_bank as common

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_CONFIG = ROOT / "config" / "oecd-pisa-indicators.json"
DEFAULT_ADDITIONAL_CONFIG = ROOT / "config" / "oecd-pisa-additional-indicators.json"
DEFAULT_PROVIDERS = ROOT / "config" / "statistics-providers.json"
DEFAULT_OUTPUT_DIR = ROOT / "dist" / "statistics"
DEFAULT_REGISTRY = "https://tiles.radlobby.at/AreaStatistics/area-registry-countries.json"


def parse_args() -> argparse.Namespace:
	parser = argparse.ArgumentParser(
		description="Build core and additional country-level OECD PISA statistics from official result tables."
	)
	parser.add_argument("--config", type=Path, default=DEFAULT_CONFIG)
	parser.add_argument("--additional-config", type=Path, default=DEFAULT_ADDITIONAL_CONFIG)
	parser.add_argument("--providers", type=Path, default=DEFAULT_PROVIDERS)
	parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
	parser.add_argument("--registry", default=DEFAULT_REGISTRY)
	parser.add_argument("--timeout", type=int, default=120)
	return parser.parse_args()


def index_entry(indicator: dict[str, Any], payload: dict[str, Any], snapshot: str) -> dict[str, Any]:
	filename = f"{indicator['slug']}.json"
	return {
		"id": indicator["id"],
		"title": indicator["title"],
		"description": indicator["description"],
		"areaLevel": "country",
		"frequency": "pisa-cycle",
		"unit": indicator["unit"],
		"classification": indicator["classification"],
		"sourceIndicator": payload["source"]["indicator"],
		"path": f"releases/{snapshot}/{filename}",
		"availableYears": payload["availableYears"],
		"defaultYear": payload["defaultYear"],
		"coverage": payload["coverage"],
	}


def main() -> None:
	args = parse_args()
	config = core.validate_config(common.read_json_path(args.config))
	additional_config = additional.validate_config(common.read_json_path(args.additional_config))
	provider_catalog = common.validate_provider_catalog(common.read_json_path(args.providers))
	provider = config["provider"]
	if not any(item["id"] == provider["id"] for item in provider_catalog["providers"]):
		raise RuntimeError("OECD PISA is missing from statistics-providers.json.")

	area_by_iso3, registry_payload = common.load_registry(args.registry, args.timeout)
	print(f"Fetching OECD PISA Annex B1A from {config['dataUrl']}")
	raw, headers, effective_url = core.fetch_xlsx(str(config["dataUrl"]), args.timeout)
	source_sha256 = hashlib.sha256(raw).hexdigest()
	workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)

	expected_sheets = {
		str(indicator["sourceIndicator"])
		for indicator in config["indicators"]
	} | {
		str(indicator["currentTable"])
		for indicator in config["indicators"]
	}
	missing_sheets = sorted(expected_sheets - set(workbook.sheetnames))
	if missing_sheets:
		workbook.close()
		raise RuntimeError(f"OECD PISA workbook is missing target sheets: {missing_sheets}")

	trend_rows_by_id: dict[str, dict[str, dict[int, float]]] = {}
	starred_by_id: dict[str, set[str]] = {}
	all_source_names: set[str] = set()
	for indicator in config["indicators"]:
		indicator_id = str(indicator["id"])
		trend_rows, starred = core.extract_trend_rows(
			workbook[str(indicator["sourceIndicator"])],
			list(indicator["expectedYears"]),
			list(indicator["valueRange"]),
		)
		current_rows = core.extract_current_2025(
			workbook[str(indicator["currentTable"])],
			list(indicator["valueRange"]),
		)
		core.validate_current_matches_trend(indicator, trend_rows, current_rows)
		trend_rows_by_id[indicator_id] = trend_rows
		starred_by_id[indicator_id] = starred
		all_source_names.update(trend_rows)
		print(
			f"{indicator_id}: sourceEntities={len(trend_rows)} current2025={len(current_rows)} "
			f"exactCurrentTrendMatches={len(current_rows)}"
		)
	workbook.close()

	mapping = core.build_country_mapping(all_source_names, config, area_by_iso3)
	print(
		f"PISA core entity mapping: source={len(all_source_names)} mapped={len(mapping)} "
		f"excludedPartial={len(config['excludedPartialEntities'])} registry={len(area_by_iso3)}"
	)

	excluded = set(config["excludedPartialEntities"])
	core_payloads: list[tuple[dict[str, Any], dict[str, Any]]] = []
	source_observations_by_id: dict[str, int] = {}
	retained_observations_by_id: dict[str, int] = {}
	for indicator in config["indicators"]:
		indicator_id = str(indicator["id"])
		values, source_observations, retained_observations = core.normalize_values(
			indicator,
			trend_rows_by_id[indicator_id],
			mapping,
			excluded,
		)
		payload = core.build_payload(
			config,
			indicator,
			area_by_iso3,
			values,
			mapping,
			starred_by_id[indicator_id],
			source_sha256,
			source_observations,
			retained_observations,
		)
		core_payloads.append((indicator, payload))
		source_observations_by_id[indicator_id] = source_observations
		retained_observations_by_id[indicator_id] = retained_observations

	cached_sources = {
		str(config["dataUrl"]): additional.SourceFile(raw=raw, headers=headers, effective_url=effective_url)
	}
	additional_payloads, additional_sources = additional.build_additional_payloads(
		additional_config,
		provider,
		area_by_iso3,
		args.timeout,
		cached_sources=cached_sources,
	)
	payloads = core_payloads + additional_payloads

	all_ids = [str(indicator["id"]) for indicator, _ in payloads]
	if len(all_ids) != 12 or len(set(all_ids)) != 12:
		raise RuntimeError(f"OECD PISA extended build must contain exactly 12 unique indicators, got {all_ids}.")

	hasher = hashlib.sha256()
	for indicator, payload in payloads:
		hasher.update(str(indicator["id"]).encode("utf-8"))
		hasher.update(b"\0")
		hasher.update(common.canonical_bytes(payload))
		hasher.update(b"\0")
	snapshot = hasher.hexdigest()[:16]

	now = datetime.now(timezone.utc)
	provider_dir = args.output_dir / provider["id"]
	release_dir = provider_dir / "releases" / snapshot
	index_indicators: list[dict[str, Any]] = []
	for indicator, payload in payloads:
		filename = f"{indicator['slug']}.json"
		common.write_json(release_dir / filename, payload)
		index_indicators.append(index_entry(indicator, payload, snapshot))

	registry_source = {
		"url": args.registry if args.registry.startswith(("https://", "http://")) else None,
		"schema": registry_payload.get("schema"),
		"generatedAt": registry_payload.get("generatedAt"),
		"areaCount": len(area_by_iso3),
	}
	if registry_source["url"] is None:
		registry_source.pop("url")

	sampling_caution_entities = sorted(
		{
			name
			for starred in starred_by_id.values()
			for name in starred
			if name in mapping
		}
	)

	core_source_summary = {
		"id": "pisa-2025-volume-i",
		"publication": config["publication"],
		"publicationDate": config["publicationDate"],
		"url": config["sourceUrl"],
		"dataUrl": config["dataUrl"],
		"effectiveDataUrl": effective_url,
		"license": provider["license"],
		"licenseUrl": provider["licenseUrl"],
		"attribution": provider["attribution"],
		"sourceFileSha256": source_sha256,
		"sourceFileBytes": len(raw),
	}
	if headers.get("last-modified"):
		core_source_summary["lastModified"] = headers["last-modified"]
	if headers.get("etag"):
		core_source_summary["etag"] = headers["etag"]

	sources_by_id = {item["id"]: item for item in additional_sources}
	sources_by_id["pisa-2025-volume-i"] = core_source_summary
	for indicator, payload in additional_payloads:
		indicator_id = str(indicator["id"])
		source_observations_by_id[indicator_id] = int(payload["source"]["sourceObservations"])
		retained_observations_by_id[indicator_id] = int(payload["source"]["retainedObservations"])

	dataset = {
		"url": config["sourceUrl"],
		"dataUrl": config["dataUrl"],
		"effectiveDataUrl": effective_url,
		"publication": config["publication"],
		"publicationDate": config["publicationDate"],
		"sourceFileSha256": source_sha256,
		"sourceFileBytes": len(raw),
		"sourceEntities": len(all_source_names),
		"mappedEntities": len(mapping),
		"excludedPartialEntities": config["excludedPartialEntities"],
		"samplingCautionEntities2025": sampling_caution_entities,
		"scopeNote": "Top-level source/entity fields describe the PISA 2025 Volume I core workbook; dataset.sources describes every official source used by the 12 indicators.",
		"sources": [sources_by_id[source_id] for source_id in sorted(sources_by_id)],
		"indicatorRows": {
			indicator_id: {
				"sourceObservations": source_observations_by_id[indicator_id],
				"retainedObservations": retained_observations_by_id[indicator_id],
			}
			for indicator_id in sorted(source_observations_by_id)
		},
	}
	if headers.get("last-modified"):
		dataset["lastModified"] = headers["last-modified"]
	if headers.get("etag"):
		dataset["etag"] = headers["etag"]

	provider_metadata = dict(provider)
	provider_metadata["dataset"] = "PISA – offizielle aggregierte Ergebnistabellen"
	provider_metadata["license"] = additional_config["providerLicense"]["label"]
	provider_metadata["licenseUrl"] = additional_config["providerLicense"]["url"]
	provider_metadata["attribution"] = "OECD PISA; konkrete Publikation und Lizenz sind je Kennzahl in den Quellmetadaten angegeben."

	provider_index = {
		"schema": "kartensammlung.statistics-provider-index/v1",
		"provider": provider_metadata,
		"retrievedAt": now.isoformat().replace("+00:00", "Z"),
		"activeSnapshot": snapshot,
		"areaRegistry": registry_source,
		"dataset": dataset,
		"indicators": index_indicators,
		"notes": [
			"Official OECD PISA aggregated country/economy result tables are used directly; no student microdata are re-aggregated.",
			"Only actual assessment cycles are published. Missing calendar years are not interpolated or carried forward.",
			"Participation in optional and innovative PISA domains is much smaller than in mathematics, reading and science; sparse coverage is intentional and is not filled from other sources.",
			"Subnational samples without equivalent Kartensammlung geometries are deliberately excluded from whole-country choropleths. Hong Kong, Macao, Chinese Taipei, the Palestinian Authority and Kosovo remain separate where OECD publishes them separately and the registry has matching geometries.",
			"OECD sampling cautions are retained in source.samplingCautionAreasByYear for additional domains. PISA 2022 creative-thinking double-asterisk scale-linkage cautions are retained separately in source.scaleLinkageCautionAreasByYear.",
			"PISA 2025 written content is CC BY 4.0. Earlier OECD publications used by the historical/special-domain tables remain under the OECD Terms & Conditions; each indicator records its exact source licence metadata.",
		],
	}
	common.write_json(provider_dir / "index.json", provider_index)
	common.write_json(args.output_dir / "index.json", common.build_global_index(provider_catalog))
	print(f"Source SHA-256 (PISA 2025 core): {source_sha256}")
	print(f"Built extended OECD PISA statistics snapshot {snapshot}")
	print(f"Indicators: {len(index_indicators)} (core={len(core_payloads)}, additional={len(additional_payloads)})")


if __name__ == "__main__":
	main()
