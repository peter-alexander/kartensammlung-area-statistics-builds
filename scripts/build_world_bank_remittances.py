#!/usr/bin/env python3
from __future__ import annotations

import argparse
from collections import defaultdict
from datetime import date, datetime, timezone
import hashlib
import io
import math
from pathlib import Path
import time
from typing import Any
from urllib.error import HTTPError, URLError
from urllib.request import Request, urlopen

from openpyxl import load_workbook

import build_world_bank as common

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_CONFIG = ROOT / "config" / "world-bank-remittances-indicators.json"
DEFAULT_PROVIDERS = ROOT / "config" / "statistics-providers.json"
DEFAULT_OUTPUT_DIR = ROOT / "dist" / "statistics"
DEFAULT_REGISTRY = "https://tiles.radlobby.at/AreaStatistics/area-registry-countries.json"
USER_AGENT = "kartensammlung-area-statistics-builds/1"


def parse_args() -> argparse.Namespace:
	parser = argparse.ArgumentParser(
		description="Build country remittance statistics from World Bank / KNOMAD resources."
	)
	parser.add_argument("--config", type=Path, default=DEFAULT_CONFIG)
	parser.add_argument("--providers", type=Path, default=DEFAULT_PROVIDERS)
	parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
	parser.add_argument("--registry", default=DEFAULT_REGISTRY)
	parser.add_argument("--timeout", type=int, default=120)
	return parser.parse_args()


def parse_iso_date(value: str, label: str) -> date:
	try:
		return date.fromisoformat(value)
	except ValueError as error:
		raise ValueError(f"Invalid {label}: {value!r}") from error


def validate_config(payload: Any) -> dict[str, Any]:
	if not isinstance(payload, dict) or payload.get("schema") != "kartensammlung.world-bank-remittances-statistics/v1":
		raise ValueError("Invalid World Bank remittances statistics config.")
	for key in ("apiBase", "countryApiUrl", "sourcePage"):
		if not str(payload.get(key, "")).startswith("https://"):
			raise ValueError(f"World Bank remittances {key} must use HTTPS.")
	for key in ("datasetId", "expectedDatasetTitle", "expectedLicense", "expectedSecurityClassification"):
		if not str(payload.get(key, "")).strip():
			raise ValueError(f"World Bank remittances config is missing {key}.")

	expected_start = payload.get("expectedStartYear")
	minimum_latest = payload.get("minimumLatestYear")
	if not isinstance(expected_start, int) or not 1900 <= expected_start <= 2200:
		raise ValueError("World Bank remittances expectedStartYear is invalid.")
	if not isinstance(minimum_latest, int) or minimum_latest < expected_start:
		raise ValueError("World Bank remittances minimumLatestYear is invalid.")
	parse_iso_date(str(payload.get("minimumResourceUpdatedDate", "")), "minimumResourceUpdatedDate")

	provider = payload.get("provider")
	if not isinstance(provider, dict) or provider.get("id") != "world-bank-remittances":
		raise ValueError("World Bank remittances provider metadata is invalid.")
	for key in ("name", "dataset", "license", "licenseUrl", "attribution"):
		if not str(provider.get(key, "")).strip():
			raise ValueError(f"World Bank remittances provider metadata is missing {key}.")

	aliases = payload.get("nameAliases")
	if not isinstance(aliases, dict):
		raise ValueError("World Bank remittances nameAliases must be an object.")
	for name, iso3 in aliases.items():
		if not str(name).strip() or len(str(iso3).strip()) != 3 or not str(iso3).strip().isalpha():
			raise ValueError(f"Invalid World Bank remittances name alias: {name!r} -> {iso3!r}")
	ignored_names = payload.get("ignoredSourceNames")
	if not isinstance(ignored_names, list) or any(not str(name).strip() for name in ignored_names):
		raise ValueError("World Bank remittances ignoredSourceNames is invalid.")

	indicators = payload.get("indicators")
	if not isinstance(indicators, list) or not indicators:
		raise ValueError("World Bank remittances config requires indicators.")
	ids: set[str] = set()
	slugs: set[str] = set()
	resource_ids: set[str] = set()
	for indicator in indicators:
		if not isinstance(indicator, dict):
			raise ValueError("World Bank remittances indicator entries must be objects.")
		for key in ("id", "slug", "resourceId", "sourceTitlePrefix", "sheet", "header", "title", "description"):
			if not str(indicator.get(key, "")).strip():
				raise ValueError(f"World Bank remittances indicator entry is missing {key}.")
		indicator_id = str(indicator["id"])
		slug = str(indicator["slug"])
		resource_id = str(indicator["resourceId"])
		if indicator_id in ids or slug in slugs or resource_id in resource_ids:
			raise ValueError(f"Duplicate World Bank remittances indicator identifier: {indicator_id}")
		ids.add(indicator_id)
		slugs.add(slug)
		resource_ids.add(resource_id)
		if not resource_id.startswith("DR"):
			raise ValueError(f"Invalid World Bank Data Catalog resource id: {resource_id}")
		unit = indicator.get("unit")
		if not isinstance(unit, dict) or not str(unit.get("id", "")).strip() or not str(unit.get("label", "")).strip():
			raise ValueError(f"Indicator {indicator_id} has invalid unit metadata.")
		common.validate_classification(indicator_id, indicator.get("classification"))
		for key in ("minAreasWithAnyValue", "minAreasInDefaultYear"):
			value = indicator.get(key)
			if not isinstance(value, int) or value <= 0:
				raise ValueError(f"Indicator {indicator_id} has invalid {key}.")
		required_areas = indicator.get("requiredAreas")
		if not isinstance(required_areas, list) or not required_areas or any(not str(value).strip() for value in required_areas):
			raise ValueError(f"Indicator {indicator_id} requires requiredAreas.")
	return payload


def fetch_bytes(url: str, timeout: int, accept: str = "*/*") -> bytes:
	delays = (0, 5, 15, 30)
	last_error: Exception | None = None
	for attempt, delay in enumerate(delays, start=1):
		if delay:
			time.sleep(delay)
		request = Request(url, headers={"Accept": accept, "User-Agent": USER_AGENT})
		try:
			with urlopen(request, timeout=timeout) as response:
				return response.read()
		except (HTTPError, URLError, TimeoutError) as error:
			last_error = error
			print(f"Request failed ({attempt}/{len(delays)}): {url}: {error}")
	if last_error is None:
		raise RuntimeError(f"Request failed without an exception: {url}")
	raise RuntimeError(f"Request failed after {len(delays)} attempts: {url}") from last_error


def validate_dataset_metadata(config: dict[str, Any], timeout: int) -> dict[str, Any]:
	url = f"{str(config['apiBase']).rstrip('/')}/datasets/{config['datasetId']}"
	payload = common.fetch_json(url, timeout)
	if not isinstance(payload, dict):
		raise RuntimeError("World Bank KNOMAD dataset metadata is not an object.")
	if str(payload.get("dataset_unique_id", "")) != str(config["datasetId"]):
		raise RuntimeError("World Bank KNOMAD dataset id changed unexpectedly.")
	identification = payload.get("identification") or {}
	if str(identification.get("title", "")).strip() != str(config["expectedDatasetTitle"]):
		raise RuntimeError(f"World Bank KNOMAD dataset title changed: {identification.get('title')!r}")
	constraints = payload.get("constraints") or {}
	license_metadata = constraints.get("license") or {}
	license_id = str(license_metadata.get("license_id", "")).strip()
	if license_id != str(config["expectedLicense"]):
		raise RuntimeError(f"World Bank KNOMAD license changed: {license_id!r}")
	security = constraints.get("security") or {}
	classification = str(security.get("classification", "")).strip()
	if classification != str(config["expectedSecurityClassification"]):
		raise RuntimeError(f"World Bank KNOMAD security classification changed: {classification!r}")
	return payload


def load_world_bank_country_names(config: dict[str, Any], area_by_iso3: dict[str, str], timeout: int) -> dict[str, str]:
	payload = common.fetch_json(str(config["countryApiUrl"]), timeout)
	if not isinstance(payload, list) or len(payload) != 2 or not isinstance(payload[1], list):
		raise RuntimeError("Unexpected World Bank country-directory response shape.")
	result: dict[str, str] = {}
	for entry in payload[1]:
		if not isinstance(entry, dict):
			continue
		name = str(entry.get("name", "")).strip()
		iso3 = str(entry.get("id", "")).strip().upper()
		region_id = str((entry.get("region") or {}).get("id", "")).strip().upper()
		if not name or len(iso3) != 3 or region_id == "NA":
			continue
		if name in result and result[name] != iso3:
			raise RuntimeError(f"Duplicate World Bank country name: {name}")
		result[name] = iso3
	aliases = {str(name): str(iso3).upper() for name, iso3 in config["nameAliases"].items()}
	for name, iso3 in aliases.items():
		if iso3 not in area_by_iso3:
			raise RuntimeError(f"World Bank remittances alias target is missing from registry: {name} -> {iso3}")
		result[name] = iso3
	return result


def parse_year_label(value: Any) -> int | None:
	if isinstance(value, bool) or value is None:
		return None
	if isinstance(value, int):
		return value if 1900 <= value <= 2200 else None
	if isinstance(value, float) and value.is_integer():
		integer = int(value)
		return integer if 1900 <= integer <= 2200 else None
	text = str(value).strip()
	if len(text) == 4 and text.isdigit():
		year = int(text)
		return year if 1900 <= year <= 2200 else None
	return None


def maybe_number(value: Any) -> float | None:
	if value is None or value == "" or isinstance(value, bool):
		return None
	if isinstance(value, str) and not value.strip():
		return None
	try:
		number = float(value)
	except (TypeError, ValueError):
		return None
	return number if math.isfinite(number) else None


def parse_value(value: Any, context: str) -> int | float | None:
	if value is None or value == "" or (isinstance(value, str) and not value.strip()):
		return None
	if isinstance(value, bool):
		raise RuntimeError(f"Boolean remittance value in {context}.")
	try:
		number = float(value)
	except (TypeError, ValueError) as error:
		raise RuntimeError(f"Non-numeric remittance value in {context}: {value!r}") from error
	if not math.isfinite(number):
		raise RuntimeError(f"Non-finite remittance value in {context}.")
	if number < 0:
		raise RuntimeError(f"Negative remittance value in {context}: {number}")
	rounded = round(number, 6)
	if rounded.is_integer():
		return int(rounded)
	return rounded


def validate_resource_metadata(
	config: dict[str, Any],
	indicator: dict[str, Any],
	metadata: Any,
) -> dict[str, Any]:
	if not isinstance(metadata, dict):
		raise RuntimeError(f"World Bank resource metadata is not an object for {indicator['resourceId']}.")
	resource_id = str(indicator["resourceId"])
	if str(metadata.get("resource_unique_id", "")) != resource_id:
		raise RuntimeError(f"World Bank resource id changed for {resource_id}.")
	dataset = metadata.get("dataset") or {}
	if str(dataset.get("unique_id", "")) != str(config["datasetId"]):
		raise RuntimeError(f"World Bank resource {resource_id} moved to another dataset.")
	identification = metadata.get("identification") or {}
	title = str(identification.get("title", metadata.get("name", ""))).strip()
	if not title.startswith(str(indicator["sourceTitlePrefix"])):
		raise RuntimeError(f"World Bank resource title changed for {resource_id}: {title!r}")
	distribution = metadata.get("distribution") or {}
	if str(distribution.get("distribution_format", "")).strip().lower() != "xlsx":
		raise RuntimeError(f"World Bank resource format changed for {resource_id}.")
	security = (metadata.get("constraints") or {}).get("security") or {}
	if str(security.get("classification", "")).strip() != str(config["expectedSecurityClassification"]):
		raise RuntimeError(f"World Bank resource security classification changed for {resource_id}.")
	updated_text = str(metadata.get("last_updated_date", "")).strip()
	try:
		updated_date = datetime.fromisoformat(updated_text.replace("Z", "+00:00")).date()
	except ValueError as error:
		raise RuntimeError(f"World Bank resource {resource_id} has invalid last_updated_date: {updated_text!r}") from error
	minimum_updated = parse_iso_date(str(config["minimumResourceUpdatedDate"]), "minimumResourceUpdatedDate")
	if updated_date < minimum_updated:
		raise RuntimeError(
			f"World Bank resource {resource_id} is unexpectedly old: {updated_date.isoformat()} < {minimum_updated.isoformat()}."
		)
	return metadata


def load_resource(
	config: dict[str, Any],
	indicator: dict[str, Any],
	area_by_iso3: dict[str, str],
	country_by_name: dict[str, str],
	timeout: int,
) -> tuple[dict[int, dict[str, int | float]], dict[str, Any]]:
	api_base = str(config["apiBase"]).rstrip("/")
	resource_id = str(indicator["resourceId"])
	metadata_url = f"{api_base}/resources/{resource_id}"
	download_url = f"{api_base}/resources/{resource_id}/download"
	metadata = validate_resource_metadata(config, indicator, common.fetch_json(metadata_url, timeout))
	data = fetch_bytes(
		download_url,
		timeout,
		"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet,*/*",
	)
	distribution = metadata["distribution"]
	declared_size = str(distribution.get("distribution_size", "")).strip()
	if declared_size.isdigit() and len(data) != int(declared_size):
		raise RuntimeError(
			f"World Bank resource {resource_id} byte size changed during download: {len(data)} != {declared_size}."
		)

	workbook = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
	try:
		if str(indicator["sheet"]) not in workbook.sheetnames:
			raise RuntimeError(f"World Bank resource {resource_id} is missing sheet {indicator['sheet']!r}.")
		sheet = workbook[str(indicator["sheet"])]
		rows = list(sheet.iter_rows(values_only=True))
	finally:
		workbook.close()
	if not rows:
		raise RuntimeError(f"World Bank resource {resource_id} sheet is empty.")
	header = list(rows[0])
	if not header or str(header[0] or "").strip() != str(indicator["header"]):
		raise RuntimeError(f"World Bank resource {resource_id} header changed: {header[0] if header else None!r}")

	year_columns: dict[int, int] = {}
	for column_index, value in enumerate(header[1:], start=1):
		year = parse_year_label(value)
		if year is None:
			continue
		if year in year_columns:
			raise RuntimeError(f"Duplicate year column in World Bank resource {resource_id}: {year}")
		year_columns[year] = column_index
	if not year_columns:
		raise RuntimeError(f"World Bank resource {resource_id} contains no annual columns.")
	available_source_years = sorted(year_columns)
	expected_start = int(config["expectedStartYear"])
	if available_source_years[0] != expected_start:
		raise RuntimeError(
			f"World Bank resource {resource_id} starts in {available_source_years[0]}, expected {expected_start}."
		)
	if available_source_years != list(range(expected_start, available_source_years[-1] + 1)):
		raise RuntimeError(f"World Bank resource {resource_id} has a non-contiguous annual timeline.")
	if available_source_years[-1] < int(config["minimumLatestYear"]):
		raise RuntimeError(f"World Bank resource {resource_id} is unexpectedly old: {available_source_years[-1]}.")

	ignored_names = set(str(name) for name in config["ignoredSourceNames"])
	ignored_seen: list[str] = []
	footer_notes: list[str] = []
	mapped_source_rows = 0
	seen_iso3: set[str] = set()
	values_by_year: dict[int, dict[str, int | float]] = defaultdict(dict)
	for row_index, row in enumerate(rows[1:], start=2):
		name = str(row[0] or "").strip() if row else ""
		if not name:
			continue
		has_annual_number = any(
			maybe_number(row[column] if column < len(row) else None) is not None
			for column in year_columns.values()
		)
		if name in ignored_names:
			ignored_seen.append(name)
			continue
		iso3 = country_by_name.get(name)
		if iso3 is None:
			if has_annual_number:
				raise RuntimeError(
					f"Unmapped numeric World Bank remittance row {row_index} in {resource_id}: {name!r}"
				)
			footer_notes.append(name)
			continue
		if iso3 not in area_by_iso3:
			raise RuntimeError(
				f"World Bank remittance row {row_index} maps outside the country registry: {name!r} -> {iso3}"
			)
		if iso3 in seen_iso3:
			raise RuntimeError(f"Duplicate World Bank remittance country row in {resource_id}: {iso3}")
		seen_iso3.add(iso3)
		mapped_source_rows += 1
		area_id = area_by_iso3[iso3]
		for year, column in year_columns.items():
			raw_value = row[column] if column < len(row) else None
			value = parse_value(raw_value, f"{resource_id} row {row_index} year {year}")
			if value is None:
				continue
			values_by_year[year][area_id] = value

	missing_ignored = sorted(ignored_names - set(ignored_seen))
	if missing_ignored:
		raise RuntimeError(f"Expected ignored remittance source names disappeared from {resource_id}: {missing_ignored}")
	if mapped_source_rows < int(indicator["minAreasWithAnyValue"]):
		raise RuntimeError(f"World Bank resource {resource_id} has too few mapped country rows: {mapped_source_rows}.")

	resource_info = {
		"resourceId": resource_id,
		"title": str((metadata.get("identification") or {}).get("title", metadata.get("name", ""))).strip(),
		"metadataUrl": metadata_url,
		"downloadUrl": download_url,
		"fileName": str(distribution.get("file_name", "")).strip(),
		"bytes": len(data),
		"sha256": hashlib.sha256(data).hexdigest(),
		"lastUpdatedAt": str(metadata.get("last_updated_date", "")).strip(),
		"sheet": str(indicator["sheet"]),
		"header": str(indicator["header"]),
		"sourceStartYear": available_source_years[0],
		"sourceLatestYear": available_source_years[-1],
		"mappedSourceRows": mapped_source_rows,
		"ignoredSourceNames": sorted(set(ignored_seen)),
		"workbookNotes": footer_notes,
	}
	return values_by_year, resource_info


def choose_default_year(
	indicator: dict[str, Any],
	available_years: list[int],
	values_by_year: dict[int, dict[str, int | float]],
) -> int:
	minimum_coverage = int(indicator["minAreasInDefaultYear"])
	eligible = [year for year in available_years if len(values_by_year[year]) >= minimum_coverage]
	if not eligible:
		raise RuntimeError(
			f"World Bank remittances has no sufficiently complete default year for {indicator['id']}: "
			f"expected at least {minimum_coverage} areas."
		)
	return eligible[-1]


def build_payload(
	config: dict[str, Any],
	indicator: dict[str, Any],
	provider: dict[str, Any],
	area_by_iso3: dict[str, str],
	values_by_year: dict[int, dict[str, int | float]],
	resource_info: dict[str, Any],
) -> dict[str, Any]:
	available_years = sorted(year for year, values in values_by_year.items() if values)
	if not available_years:
		raise RuntimeError(f"World Bank remittances returned no mapped values for {indicator['id']}.")
	if available_years[0] != int(config["expectedStartYear"]):
		raise RuntimeError(f"World Bank remittances history starts unexpectedly for {indicator['id']}.")
	latest_year = available_years[-1]
	if latest_year < int(config["minimumLatestYear"]):
		raise RuntimeError(f"World Bank remittances latest year is too old for {indicator['id']}: {latest_year}.")
	mapped_areas = {area_id for values in values_by_year.values() for area_id in values}
	if len(mapped_areas) < int(indicator["minAreasWithAnyValue"]):
		raise RuntimeError(
			f"World Bank remittances coverage too small for {indicator['id']}: {len(mapped_areas)} areas."
		)
	default_year = choose_default_year(indicator, available_years, values_by_year)
	for required_area in indicator["requiredAreas"]:
		if required_area not in mapped_areas:
			raise RuntimeError(f"World Bank remittances required area is missing entirely: {required_area}")
		if required_area not in values_by_year[default_year]:
			raise RuntimeError(
				f"World Bank remittances required area is missing in default year {default_year}: {required_area}"
			)

	return {
		"schema": "kartensammlung.statistics-indicator/v1",
		"indicator": {
			"id": indicator["id"],
			"title": indicator["title"],
			"description": indicator["description"],
			"areaLevel": "country",
			"frequency": "annual",
			"unit": indicator["unit"],
			"classification": indicator["classification"],
		},
		"source": {
			"providerId": provider["id"],
			"providerName": provider["name"],
			"dataset": provider["dataset"],
			"datasetId": config["datasetId"],
			"license": provider["license"],
			"licenseUrl": provider["licenseUrl"],
			"attribution": provider["attribution"],
			"url": config["sourcePage"],
			"resourceId": resource_info["resourceId"],
			"resourceTitle": resource_info["title"],
			"resourceMetadataUrl": resource_info["metadataUrl"],
			"downloadUrl": resource_info["downloadUrl"],
			"fileName": resource_info["fileName"],
			"resourceUpdatedAt": resource_info["lastUpdatedAt"],
			"sha256": resource_info["sha256"],
			"bytes": resource_info["bytes"],
			"sheet": resource_info["sheet"],
			"header": resource_info["header"],
			"mappingAliases": config["nameAliases"],
			"ignoredSourceNames": resource_info["ignoredSourceNames"],
			"workbookNotes": resource_info["workbookNotes"],
		},
		"availableYears": available_years,
		"defaultYear": default_year,
		"coverage": {
			"registryAreas": len(area_by_iso3),
			"mappedSourceRows": resource_info["mappedSourceRows"],
			"areasWithAnyValue": len(mapped_areas),
			"latestYear": latest_year,
			"areasInLatestYear": len(values_by_year[latest_year]),
			"areasInDefaultYear": len(values_by_year[default_year]),
		},
		"values": {
			str(year): {
				area_id: values_by_year[year][area_id]
				for area_id in sorted(values_by_year[year])
			}
			for year in available_years
		},
	}


def main() -> None:
	args = parse_args()
	config = validate_config(common.read_json_path(args.config))
	provider_catalog = common.validate_provider_catalog(common.read_json_path(args.providers))
	provider = config["provider"]
	if not any(item["id"] == provider["id"] for item in provider_catalog["providers"]):
		raise RuntimeError("World Bank remittances is missing from statistics-providers.json.")
	area_by_iso3, registry_payload = common.load_registry(args.registry, args.timeout)
	dataset_metadata = validate_dataset_metadata(config, args.timeout)
	country_by_name = load_world_bank_country_names(config, area_by_iso3, args.timeout)

	indicator_payloads: list[tuple[dict[str, Any], dict[str, Any], dict[str, Any]]] = []
	for indicator in config["indicators"]:
		print(f"Fetching {indicator['id']} ({indicator['resourceId']})")
		values_by_year, resource_info = load_resource(
			config,
			indicator,
			area_by_iso3,
			country_by_name,
			args.timeout,
		)
		payload = build_payload(config, indicator, provider, area_by_iso3, values_by_year, resource_info)
		coverage = payload["coverage"]
		print(
			f"{indicator['id']}: mappedRows={coverage['mappedSourceRows']} "
			f"mapped={coverage['areasWithAnyValue']} years={payload['availableYears'][0]}-{coverage['latestYear']} "
			f"latestCoverage={coverage['areasInLatestYear']} defaultYear={payload['defaultYear']} "
			f"defaultCoverage={coverage['areasInDefaultYear']}"
		)
		indicator_payloads.append((indicator, payload, resource_info))

	hasher = hashlib.sha256()
	for indicator, payload, _ in sorted(indicator_payloads, key=lambda item: item[0]["id"]):
		hasher.update(str(indicator["id"]).encode("utf-8"))
		hasher.update(b"\0")
		hasher.update(common.canonical_bytes(payload))
		hasher.update(b"\0")
	snapshot = hasher.hexdigest()[:16]

	provider_dir = args.output_dir / provider["id"]
	release_dir = provider_dir / "releases" / snapshot
	index_indicators: list[dict[str, Any]] = []
	resources: list[dict[str, Any]] = []
	for indicator, payload, resource_info in indicator_payloads:
		filename = f"{indicator['slug']}.json"
		common.write_json(release_dir / filename, payload)
		index_indicators.append({
			"id": indicator["id"],
			"title": indicator["title"],
			"description": indicator["description"],
			"areaLevel": "country",
			"frequency": "annual",
			"unit": indicator["unit"],
			"classification": indicator["classification"],
			"sourceIndicator": indicator["resourceId"],
			"path": f"releases/{snapshot}/{filename}",
			"availableYears": payload["availableYears"],
			"defaultYear": payload["defaultYear"],
			"coverage": payload["coverage"],
		})
		resources.append({
			"resourceId": resource_info["resourceId"],
			"title": resource_info["title"],
			"fileName": resource_info["fileName"],
			"lastUpdatedAt": resource_info["lastUpdatedAt"],
			"bytes": resource_info["bytes"],
			"sha256": resource_info["sha256"],
			"sourceStartYear": resource_info["sourceStartYear"],
			"sourceLatestYear": resource_info["sourceLatestYear"],
			"mappedSourceRows": resource_info["mappedSourceRows"],
			"ignoredSourceNames": resource_info["ignoredSourceNames"],
		})

	registry_source = {
		"url": args.registry if args.registry.startswith(("https://", "http://")) else None,
		"schema": registry_payload.get("schema"),
		"generatedAt": registry_payload.get("generatedAt"),
		"areaCount": len(area_by_iso3),
	}
	if registry_source["url"] is None:
		registry_source.pop("url")

	dataset_constraints = dataset_metadata.get("constraints") or {}
	provider_index = {
		"schema": "kartensammlung.statistics-provider-index/v1",
		"provider": provider,
		"retrievedAt": datetime.now(timezone.utc).isoformat().replace("+00:00", "Z"),
		"activeSnapshot": snapshot,
		"areaRegistry": registry_source,
		"dataset": {
			"datasetId": config["datasetId"],
			"title": str((dataset_metadata.get("identification") or {}).get("title", "")).strip(),
			"url": config["sourcePage"],
			"license": str((dataset_constraints.get("license") or {}).get("license_id", "")).strip(),
			"securityClassification": str((dataset_constraints.get("security") or {}).get("classification", "")).strip(),
			"resources": sorted(resources, key=lambda item: item["resourceId"]),
		},
		"indicators": index_indicators,
		"notes": [
			"Annual inward and outward remittance flows are taken directly from the World Bank KNOMAD resources and reported in current nominal US dollars (millions).",
			"Zero values are retained as reported; missing cells remain missing. No interpolation, extrapolation or cross-source fallback is applied.",
			"Source country names are resolved through the World Bank country directory plus an explicit, validated alias table; fuzzy matching is not used.",
			"The source workbook notes are preserved in each indicator payload, including release-specific methodology and estimate notes.",
		],
	}
	common.write_json(provider_dir / "index.json", provider_index)
	common.write_json(args.output_dir / "index.json", common.build_global_index(provider_catalog))
	print(f"Built World Bank remittances snapshot {snapshot}")
	print(f"Indicators: {len(index_indicators)}")


if __name__ == "__main__":
	main()
