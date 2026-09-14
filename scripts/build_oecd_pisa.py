#!/usr/bin/env python3
from __future__ import annotations

import argparse
import hashlib
import io
import json
import math
import re
import time
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from urllib.error import HTTPError, URLError
from urllib.request import Request, urlopen

import pycountry
from openpyxl import load_workbook

import build_world_bank as common

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_CONFIG = ROOT / "config" / "oecd-pisa-indicators.json"
DEFAULT_PROVIDERS = ROOT / "config" / "statistics-providers.json"
DEFAULT_OUTPUT_DIR = ROOT / "dist" / "statistics"
DEFAULT_REGISTRY = "https://tiles.radlobby.at/AreaStatistics/area-registry-countries.json"
USER_AGENT = "kartensammlung-area-statistics-builds/1"
MISSING_MARKER = "m"


def parse_args() -> argparse.Namespace:
	parser = argparse.ArgumentParser(
		description="Build country PISA mean-score statistics from OECD PISA 2025 Annex B1A."
	)
	parser.add_argument("--config", type=Path, default=DEFAULT_CONFIG)
	parser.add_argument("--providers", type=Path, default=DEFAULT_PROVIDERS)
	parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
	parser.add_argument("--registry", default=DEFAULT_REGISTRY)
	parser.add_argument("--timeout", type=int, default=120)
	return parser.parse_args()


def validate_value_range(indicator_id: str, value_range: Any) -> None:
	if (
		not isinstance(value_range, list)
		or len(value_range) != 2
		or any(isinstance(value, bool) or not isinstance(value, (int, float)) for value in value_range)
		or not all(math.isfinite(float(value)) for value in value_range)
		or float(value_range[0]) >= float(value_range[1])
	):
		raise ValueError(f"OECD PISA indicator {indicator_id} has invalid valueRange.")


def validate_config(payload: Any) -> dict[str, Any]:
	if not isinstance(payload, dict) or payload.get("schema") != "kartensammlung.oecd-pisa-statistics/v1":
		raise ValueError("Invalid OECD PISA statistics config.")
	for key in ("sourceUrl", "dataUrl"):
		if not str(payload.get(key, "")).startswith("https://"):
			raise ValueError(f"OECD PISA {key} must use HTTPS.")
	publication_date = str(payload.get("publicationDate", ""))
	if publication_date != "2026-09-08":
		raise ValueError("Unexpected PISA 2025 publication date.")

	provider = payload.get("provider")
	if not isinstance(provider, dict) or provider.get("id") != "oecd-pisa":
		raise ValueError("OECD PISA provider metadata is invalid.")
	for key in ("name", "dataset", "license", "licenseUrl", "attribution"):
		if not str(provider.get(key, "")).strip():
			raise ValueError(f"OECD PISA provider metadata is missing {key}.")
	if provider.get("license") != "CC BY 4.0":
		raise ValueError("OECD PISA license must remain CC BY 4.0.")

	excluded = payload.get("excludedPartialEntities")
	expected_excluded = {
		"B-S-J-Z (China)",
		"Dushanbe (Tajikistan)",
		"Kurdistan Region (Iraq)",
		"Ukrainian regions (17 of 27)",
	}
	if not isinstance(excluded, list) or set(excluded) != expected_excluded or len(excluded) != len(expected_excluded):
		raise ValueError("OECD PISA excluded partial-country entities changed unexpectedly.")

	aliases = payload.get("countryAliases")
	if not isinstance(aliases, dict) or not aliases:
		raise ValueError("OECD PISA countryAliases are required.")
	for source_name, iso3 in aliases.items():
		if not str(source_name).strip() or not re.fullmatch(r"[A-Z0-9]{3}", str(iso3)):
			raise ValueError(f"Invalid OECD PISA country alias: {source_name!r} -> {iso3!r}")

	expected_indicators = {
		"education.pisa-science-mean-score": (
			"Table I.B1.2a.36",
			"Table I.B1.2a.1",
			[2006, 2009, 2012, 2015, 2018, 2022, 2025],
		),
		"education.pisa-reading-mean-score": (
			"Table I.B1.2a.37",
			"Table I.B1.2a.2",
			[2000, 2003, 2006, 2009, 2012, 2015, 2018, 2022, 2025],
		),
		"education.pisa-mathematics-mean-score": (
			"Table I.B1.2a.38",
			"Table I.B1.2a.3",
			[2003, 2006, 2009, 2012, 2015, 2018, 2022, 2025],
		),
	}
	indicators = payload.get("indicators")
	if not isinstance(indicators, list) or len(indicators) != len(expected_indicators):
		raise ValueError("OECD PISA config requires exactly three indicators.")

	ids: set[str] = set()
	slugs: set[str] = set()
	for indicator in indicators:
		if not isinstance(indicator, dict):
			raise ValueError("OECD PISA indicator must be an object.")
		for key in ("id", "slug", "sourceIndicator", "currentTable", "title", "description"):
			if not str(indicator.get(key, "")).strip():
				raise ValueError(f"OECD PISA indicator is missing {key}.")
		indicator_id = str(indicator["id"])
		slug = str(indicator["slug"])
		if indicator_id in ids:
			raise ValueError(f"Duplicate OECD PISA indicator id: {indicator_id}")
		if slug in slugs:
			raise ValueError(f"Duplicate OECD PISA indicator slug: {slug}")
		ids.add(indicator_id)
		slugs.add(slug)
		expected = expected_indicators.get(indicator_id)
		actual = (
			str(indicator["sourceIndicator"]),
			str(indicator["currentTable"]),
			indicator.get("expectedYears"),
		)
		if expected is None or actual != expected:
			raise ValueError(f"Unexpected OECD PISA table/year mapping for {indicator_id}: {actual}")
		unit = indicator.get("unit")
		if not isinstance(unit, dict) or unit.get("id") != "pisa-points":
			raise ValueError(f"OECD PISA target unit must be pisa-points for {indicator_id}.")
		common.validate_classification(indicator_id, indicator.get("classification"))
		validate_value_range(indicator_id, indicator.get("valueRange"))
		for key in ("minAreasWithAnyValue", "minAreasInDefaultYear"):
			value = indicator.get(key)
			if not isinstance(value, int) or value <= 0:
				raise ValueError(f"OECD PISA indicator {indicator_id} has invalid {key}.")
		required_areas = indicator.get("requiredAreas")
		if not isinstance(required_areas, list) or not required_areas:
			raise ValueError(f"OECD PISA indicator {indicator_id} requires requiredAreas.")
	return payload


def fetch_xlsx(url: str, timeout: int) -> tuple[bytes, dict[str, str], str]:
	last_error: Exception | None = None
	for attempt, delay in enumerate((0, 3, 10, 20, 40), start=1):
		if delay:
			time.sleep(delay)
		request = Request(
			url,
			headers={
				"Accept": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet,application/octet-stream;q=0.9,*/*;q=0.1",
				"User-Agent": USER_AGENT,
			},
		)
		try:
			with urlopen(request, timeout=timeout) as response:
				raw = response.read()
				headers = {key.lower(): value for key, value in response.headers.items()}
				effective_url = response.geturl()
			if not raw:
				raise RuntimeError("OECD PISA Statlink returned an empty response.")
			if not raw.startswith(b"PK"):
				raise RuntimeError("OECD PISA Statlink did not return an XLSX/ZIP payload.")
			return raw, headers, effective_url
		except (HTTPError, URLError, TimeoutError, OSError, RuntimeError) as error:
			last_error = error
			print(f"OECD PISA request failed ({attempt}/5): {url}: {error}")
	if last_error is None:
		raise RuntimeError("OECD PISA request failed without an exception.")
	raise RuntimeError("OECD PISA request failed after 5 attempts.") from last_error


def clean_source_name(value: str) -> tuple[str, bool]:
	raw = value.strip()
	starred = raw.endswith("*")
	return raw.removesuffix("*").strip(), starred


def extract_year_columns(sheet: Any, expected_years: list[int]) -> dict[int, int]:
	header = list(next(sheet.iter_rows(min_row=8, max_row=8, values_only=True)))
	year_columns: dict[int, int] = {}
	for column, value in enumerate(header):
		if not isinstance(value, str):
			continue
		match = re.fullmatch(r"PISA (\d{4})", value.strip())
		if match:
			year = int(match.group(1))
			if year in year_columns:
				raise RuntimeError(f"Duplicate PISA year column {year} in {sheet.title}.")
			year_columns[year] = column
	if sorted(year_columns) != expected_years:
		raise RuntimeError(
			f"Unexpected PISA cycle columns in {sheet.title}: {sorted(year_columns)}; expected {expected_years}."
		)
	return year_columns


def extract_trend_rows(
	sheet: Any,
	expected_years: list[int],
	value_range: list[int | float],
) -> tuple[dict[str, dict[int, float]], set[str]]:
	year_columns = extract_year_columns(sheet, expected_years)
	rows: dict[str, dict[int, float]] = {}
	starred_entities: set[str] = set()
	value_min, value_max = (float(item) for item in value_range)
	for row in sheet.iter_rows(values_only=True):
		name_raw = row[0] if row else None
		if not isinstance(name_raw, str):
			continue
		name, starred = clean_source_name(name_raw)
		if not name or name.startswith("OECD average"):
			continue
		values: dict[int, float] = {}
		for year, column in year_columns.items():
			value = row[column]
			if value is None:
				continue
			if isinstance(value, str):
				marker = value.strip()
				if not marker or marker == MISSING_MARKER:
					continue
				raise RuntimeError(f"Unexpected PISA marker {marker!r} for {name} {year} in {sheet.title}.")
			if isinstance(value, bool) or not isinstance(value, (int, float)):
				raise RuntimeError(f"Unexpected PISA value type for {name} {year} in {sheet.title}: {value!r}")
			number = float(value)
			if not math.isfinite(number) or number < value_min or number > value_max:
				raise RuntimeError(f"PISA value outside configured range: {name} {year} {number} in {sheet.title}.")
			values[year] = number
		if not values:
			continue
		if name in rows:
			raise RuntimeError(f"Duplicate PISA entity row {name!r} in {sheet.title}.")
		rows[name] = values
		if starred:
			starred_entities.add(name)
	return rows, starred_entities


def extract_current_2025(sheet: Any, value_range: list[int | float]) -> dict[str, float]:
	rows: dict[str, float] = {}
	value_min, value_max = (float(item) for item in value_range)
	for row in sheet.iter_rows(values_only=True):
		if not row or not isinstance(row[0], str):
			continue
		name, _ = clean_source_name(row[0])
		if not name or name.startswith("OECD "):
			continue
		value = row[1] if len(row) > 1 else None
		if value is None or isinstance(value, str):
			continue
		if isinstance(value, bool) or not isinstance(value, (int, float)):
			raise RuntimeError(f"Unexpected PISA 2025 current-table value for {name}: {value!r}")
		number = float(value)
		if not math.isfinite(number) or number < value_min or number > value_max:
			raise RuntimeError(f"PISA 2025 current-table value outside configured range: {name} {number}.")
		if name in rows:
			raise RuntimeError(f"Duplicate PISA 2025 current-table entity: {name}")
		rows[name] = number
	return rows


def iso3_for_source_name(name: str, aliases: dict[str, str]) -> str | None:
	if name in aliases:
		return aliases[name]
	try:
		return pycountry.countries.lookup(name).alpha_3
	except LookupError:
		return None


def build_country_mapping(
	all_source_names: set[str],
	config: dict[str, Any],
	area_by_iso3: dict[str, str],
) -> dict[str, tuple[str, str]]:
	excluded = set(config["excludedPartialEntities"])
	missing_exclusions = sorted(excluded - all_source_names)
	if missing_exclusions:
		raise RuntimeError(f"Configured PISA partial-country exclusions disappeared: {missing_exclusions}")

	aliases = {str(key): str(value) for key, value in config["countryAliases"].items()}
	mapping: dict[str, tuple[str, str]] = {}
	unresolved: list[str] = []
	missing_registry: list[tuple[str, str]] = []
	used_iso3: dict[str, str] = {}
	for name in sorted(all_source_names):
		if name in excluded:
			continue
		iso3 = iso3_for_source_name(name, aliases)
		if iso3 is None:
			unresolved.append(name)
			continue
		area_id = area_by_iso3.get(iso3)
		if area_id is None:
			missing_registry.append((name, iso3))
			continue
		previous = used_iso3.get(iso3)
		if previous is not None and previous != name:
			raise RuntimeError(f"Multiple PISA entities map to {iso3}: {previous!r}, {name!r}")
		used_iso3[iso3] = name
		mapping[name] = (iso3, area_id)
	if unresolved:
		raise RuntimeError(f"Unresolved OECD PISA country/economy names: {unresolved}")
	if missing_registry:
		raise RuntimeError(f"OECD PISA mappings missing from area registry: {missing_registry}")
	return mapping


def validate_current_matches_trend(
	indicator: dict[str, Any],
	trend_rows: dict[str, dict[int, float]],
	current_rows: dict[str, float],
) -> None:
	trend_2025 = {
		name: values[2025]
		for name, values in trend_rows.items()
		if 2025 in values
	}
	if set(current_rows) != set(trend_2025):
		missing_current = sorted(set(trend_2025) - set(current_rows))
		missing_trend = sorted(set(current_rows) - set(trend_2025))
		raise RuntimeError(
			f"PISA 2025 current/trend entity mismatch for {indicator['id']}: "
			f"missingCurrent={missing_current}, missingTrend={missing_trend}."
		)
	for name, current_value in current_rows.items():
		trend_value = trend_2025[name]
		if abs(current_value - trend_value) > 1e-10:
			raise RuntimeError(
				f"PISA 2025 current/trend value mismatch for {indicator['id']} {name}: "
				f"{current_value} != {trend_value}."
			)


def normalize_values(
	indicator: dict[str, Any],
	trend_rows: dict[str, dict[int, float]],
	mapping: dict[str, tuple[str, str]],
	excluded: set[str],
) -> tuple[dict[int, dict[str, int | float]], int, int]:
	values: dict[int, dict[str, int | float]] = defaultdict(dict)
	source_observations = 0
	retained_observations = 0
	for name, source_values in trend_rows.items():
		for year, value in source_values.items():
			source_observations += 1
			if name in excluded:
				continue
			mapped = mapping.get(name)
			if mapped is None:
				raise RuntimeError(f"Missing PISA country mapping for {name!r}.")
			_, area_id = mapped
			if area_id in values[year]:
				raise RuntimeError(f"Duplicate PISA value for {indicator['id']} {area_id} {year}.")
			values[year][area_id] = common.normalize_number(value)
			retained_observations += 1
	return values, source_observations, retained_observations


def choose_default_year(
	indicator: dict[str, Any],
	available_years: list[int],
	values: dict[int, dict[str, int | float]],
) -> int:
	minimum = int(indicator["minAreasInDefaultYear"])
	eligible = [year for year in available_years if len(values[year]) >= minimum]
	if not eligible:
		raise RuntimeError(
			f"OECD PISA {indicator['id']} has no cycle with at least {minimum} mapped countries."
		)
	return eligible[-1]


def build_payload(
	config: dict[str, Any],
	indicator: dict[str, Any],
	area_by_iso3: dict[str, str],
	values: dict[int, dict[str, int | float]],
	mapping: dict[str, tuple[str, str]],
	starred_entities: set[str],
	source_sha256: str,
	source_observations: int,
	retained_observations: int,
) -> dict[str, Any]:
	provider = config["provider"]
	available_years = sorted(year for year, mapped in values.items() if mapped)
	if available_years != indicator["expectedYears"]:
		raise RuntimeError(
			f"OECD PISA mapped years changed for {indicator['id']}: {available_years}; "
			f"expected {indicator['expectedYears']}."
		)
	mapped_areas = {area_id for mapped in values.values() for area_id in mapped}
	if len(mapped_areas) < int(indicator["minAreasWithAnyValue"]):
		raise RuntimeError(
			f"OECD PISA coverage too small for {indicator['id']}: {len(mapped_areas)} mapped countries/economies, "
			f"expected at least {indicator['minAreasWithAnyValue']}."
		)
	default_year = choose_default_year(indicator, available_years, values)
	if default_year != 2025:
		raise RuntimeError(f"OECD PISA default year must currently be 2025 for {indicator['id']}, got {default_year}.")
	for area_id in indicator["requiredAreas"]:
		if area_id not in values[default_year]:
			raise RuntimeError(
				f"OECD PISA {indicator['id']} default cycle {default_year} is missing required area {area_id}."
			)

	sampling_caution_areas = sorted(
		mapping[name][1]
		for name in starred_entities
		if name in mapping and 2025 in next(
			(source_values for source_name, source_values in [] if source_name == name),
			{},
		)
	)
	# The asterisk is attached to country/economy names in the OECD 2025 trend tables.
	# Store every mapped starred entity as a 2025 sampling caution; no historical values are altered.
	sampling_caution_areas = sorted(mapping[name][1] for name in starred_entities if name in mapping)

	latest_year = available_years[-1]
	payload = {
		"schema": "kartensammlung.statistics-indicator/v1",
		"indicator": {
			"id": indicator["id"],
			"title": indicator["title"],
			"description": indicator["description"],
			"areaLevel": "country",
			"frequency": "pisa-cycle",
			"unit": indicator["unit"],
			"classification": indicator["classification"],
		},
		"source": {
			"providerId": provider["id"],
			"providerName": provider["name"],
			"dataset": provider["dataset"],
			"publication": config["publication"],
			"publicationDate": config["publicationDate"],
			"license": provider["license"],
			"licenseUrl": provider["licenseUrl"],
			"attribution": provider["attribution"],
			"url": config["sourceUrl"],
			"dataUrl": config["dataUrl"],
			"indicator": indicator["sourceIndicator"],
			"currentTable": indicator["currentTable"],
			"sourceFileSha256": source_sha256,
			"sourceObservations": source_observations,
			"retainedObservations": retained_observations,
			"excludedPartialEntities": config["excludedPartialEntities"],
			"samplingCautionAreas2025": sampling_caution_areas,
		},
		"availableYears": available_years,
		"defaultYear": default_year,
		"coverage": {
			"registryAreas": len(area_by_iso3),
			"areasWithAnyValue": len(mapped_areas),
			"latestYear": latest_year,
			"areasInLatestYear": len(values[latest_year]),
			"areasInDefaultYear": len(values[default_year]),
			"observations": retained_observations,
		},
		"values": {
			str(year): {
				area_id: values[year][area_id]
				for area_id in sorted(values[year])
			}
			for year in available_years
		},
	}
	print(
		f"{indicator['id']}: observations={retained_observations}/{source_observations} "
		f"mappedAreas={len(mapped_areas)} cycles={available_years} "
		f"latestCoverage={len(values[latest_year])} samplingCautions={sampling_caution_areas}"
	)
	return payload


def main() -> None:
	args = parse_args()
	config = validate_config(common.read_json_path(args.config))
	provider_catalog = common.validate_provider_catalog(common.read_json_path(args.providers))
	provider = config["provider"]
	if not any(item["id"] == provider["id"] for item in provider_catalog["providers"]):
		raise RuntimeError("OECD PISA is missing from statistics-providers.json.")

	area_by_iso3, registry_payload = common.load_registry(args.registry, args.timeout)
	print(f"Fetching OECD PISA Annex B1A from {config['dataUrl']}")
	raw, headers, effective_url = fetch_xlsx(str(config["dataUrl"]), args.timeout)
	source_sha256 = hashlib.sha256(raw).hexdigest()
	workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)

	expected_sheets = {
		str(indicator["sourceIndicator"])
		for indicator in config["indicators"]
	} | {
		str(indicator["currentTable"])
		for indicator in config["indicators"]
	}
	missing_sheets = sorted(expected_sheets - set(workbook.sheetnames))
	if missing_sheets:
		workbook.close()
		raise RuntimeError(f"OECD PISA workbook is missing target sheets: {missing_sheets}")

	trend_rows_by_id: dict[str, dict[str, dict[int, float]]] = {}
	starred_by_id: dict[str, set[str]] = {}
	all_source_names: set[str] = set()
	for indicator in config["indicators"]:
		indicator_id = str(indicator["id"])
		trend_rows, starred = extract_trend_rows(
			workbook[str(indicator["sourceIndicator"])],
			list(indicator["expectedYears"]),
			list(indicator["valueRange"]),
		)
		current_rows = extract_current_2025(
			workbook[str(indicator["currentTable"])],
			list(indicator["valueRange"]),
		)
		validate_current_matches_trend(indicator, trend_rows, current_rows)
		trend_rows_by_id[indicator_id] = trend_rows
		starred_by_id[indicator_id] = starred
		all_source_names.update(trend_rows)
		print(
			f"{indicator_id}: sourceEntities={len(trend_rows)} current2025={len(current_rows)} "
			f"exactCurrentTrendMatches={len(current_rows)}"
		)
	workbook.close()

	mapping = build_country_mapping(all_source_names, config, area_by_iso3)
	print(
		f"PISA entity mapping: source={len(all_source_names)} mapped={len(mapping)} "
		f"excludedPartial={len(config['excludedPartialEntities'])} registry={len(area_by_iso3)}"
	)

	excluded = set(config["excludedPartialEntities"])
	payloads: list[tuple[dict[str, Any], dict[str, Any]]] = []
	source_observations_by_id: dict[str, int] = {}
	retained_observations_by_id: dict[str, int] = {}
	for indicator in config["indicators"]:
		indicator_id = str(indicator["id"])
		values, source_observations, retained_observations = normalize_values(
			indicator,
			trend_rows_by_id[indicator_id],
			mapping,
			excluded,
		)
		payload = build_payload(
			config,
			indicator,
			area_by_iso3,
			values,
			mapping,
			starred_by_id[indicator_id],
			source_sha256,
			source_observations,
			retained_observations,
		)
		payloads.append((indicator, payload))
		source_observations_by_id[indicator_id] = source_observations
		retained_observations_by_id[indicator_id] = retained_observations

	hasher = hashlib.sha256()
	for indicator, payload in payloads:
		hasher.update(str(indicator["id"]).encode("utf-8"))
		hasher.update(b"\0")
		hasher.update(common.canonical_bytes(payload))
		hasher.update(b"\0")
	snapshot = hasher.hexdigest()[:16]

	now = datetime.now(timezone.utc)
	provider_dir = args.output_dir / provider["id"]
	release_dir = provider_dir / "releases" / snapshot
	index_indicators: list[dict[str, Any]] = []
	for indicator, payload in payloads:
		filename = f"{indicator['slug']}.json"
		common.write_json(release_dir / filename, payload)
		index_indicators.append(
			{
				"id": indicator["id"],
				"title": indicator["title"],
				"description": indicator["description"],
				"areaLevel": "country",
				"frequency": "pisa-cycle",
				"unit": indicator["unit"],
				"classification": indicator["classification"],
				"sourceIndicator": indicator["sourceIndicator"],
				"path": f"releases/{snapshot}/{filename}",
				"availableYears": payload["availableYears"],
				"defaultYear": payload["defaultYear"],
				"coverage": payload["coverage"],
			}
		)

	registry_source = {
		"url": args.registry if args.registry.startswith(("https://", "http://")) else None,
		"schema": registry_payload.get("schema"),
		"generatedAt": registry_payload.get("generatedAt"),
		"areaCount": len(area_by_iso3),
	}
	if registry_source["url"] is None:
		registry_source.pop("url")

	sampling_caution_entities = sorted(
		{
			name
			for starred in starred_by_id.values()
			for name in starred
			if name in mapping
		}
	)
	dataset = {
		"url": config["sourceUrl"],
		"dataUrl": config["dataUrl"],
		"effectiveDataUrl": effective_url,
		"publication": config["publication"],
		"publicationDate": config["publicationDate"],
		"sourceFileSha256": source_sha256,
		"sourceFileBytes": len(raw),
		"sourceEntities": len(all_source_names),
		"mappedEntities": len(mapping),
		"excludedPartialEntities": config["excludedPartialEntities"],
		"samplingCautionEntities2025": sampling_caution_entities,
		"indicatorRows": {
			indicator_id: {
				"sourceObservations": source_observations_by_id[indicator_id],
				"retainedObservations": retained_observations_by_id[indicator_id],
			}
			for indicator_id in sorted(source_observations_by_id)
		},
	}
	if headers.get("last-modified"):
		dataset["lastModified"] = headers["last-modified"]
	if headers.get("etag"):
		dataset["etag"] = headers["etag"]

	provider_index = {
		"schema": "kartensammlung.statistics-provider-index/v1",
		"provider": provider,
		"retrievedAt": now.isoformat().replace("+00:00", "Z"),
		"activeSnapshot": snapshot,
		"areaRegistry": registry_source,
		"dataset": dataset,
		"indicators": index_indicators,
		"notes": [
			"Official OECD PISA 2025 Annex B1A country/economy mean scores are used directly; no student microdata are re-aggregated.",
			"Only actual PISA assessment cycles are published. Missing calendar years are not interpolated or carried forward.",
			"Hong Kong (China), Macao (China), Chinese Taipei, the Palestinian Authority and Kosovo are mapped to their own country/economy geometries in the Kartensammlung registry and are not folded into another state.",
			"Four partial-country samples are deliberately excluded from whole-country choropleths: B-S-J-Z (China), Dushanbe (Tajikistan), Kurdistan Region (Iraq), and Ukrainian regions (17 of 27).",
			"OECD cycle labels are retained as published. Some historical PISA 2000/2009+ and PISA for Development assessments were administered in a following calendar year but remain assigned to the OECD cycle year.",
			"Countries/economies marked with an asterisk by OECD for PISA 2025 sampling standards remain in the dataset and are listed in dataset.samplingCautionEntities2025; their values are not altered.",
		],
	}
	common.write_json(provider_dir / "index.json", provider_index)
	common.write_json(args.output_dir / "index.json", common.build_global_index(provider_catalog))
	print(f"Source SHA-256: {source_sha256}")
	print(f"Built OECD PISA statistics snapshot {snapshot}")
	print(f"Indicators: {len(index_indicators)}")


if __name__ == "__main__":
	main()
