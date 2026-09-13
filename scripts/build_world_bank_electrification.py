#!/usr/bin/env python3
from __future__ import annotations

import argparse
import hashlib
import math
import tempfile
import time
from collections import Counter
from pathlib import Path
from typing import Any
from urllib.error import HTTPError, URLError
from urllib.request import Request, urlopen

from openpyxl import load_workbook

import build_world_bank as common

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_CONFIG = ROOT / "config" / "world-bank-electrification-indicators.json"
DEFAULT_PROVIDERS = ROOT / "config" / "statistics-providers.json"
DEFAULT_OUTPUT_DIR = ROOT / "dist" / "statistics"
DEFAULT_REGISTRY = "https://tiles.radlobby.at/AreaStatistics/area-registry-countries.json"
USER_AGENT = "kartensammlung-area-statistics-builds/1"
SOURCE_SHEET = "UN reporting"


def parse_args() -> argparse.Namespace:
	parser = argparse.ArgumentParser(
		description="Build country electricity-access statistics from the World Bank Global Electrification Database."
	)
	parser.add_argument("--config", type=Path, default=DEFAULT_CONFIG)
	parser.add_argument("--providers", type=Path, default=DEFAULT_PROVIDERS)
	parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
	parser.add_argument("--registry", default=DEFAULT_REGISTRY)
	parser.add_argument("--timeout", type=int, default=120)
	return parser.parse_args()


def validate_config(payload: Any) -> dict[str, Any]:
	if not isinstance(payload, dict) or payload.get("schema") != "kartensammlung.world-bank-electrification-statistics/v1":
		raise ValueError("Invalid World Bank electrification statistics config.")
	for key in ("downloadUrl", "sourceUrl", "catalogUrl"):
		if not str(payload.get(key, "")).startswith("https://"):
			raise ValueError(f"World Bank electrification {key} must use HTTPS.")
	provider = payload.get("provider")
	if not isinstance(provider, dict) or provider.get("id") != "world-bank-electrification":
		raise ValueError("World Bank electrification provider metadata is invalid.")
	for key in ("name", "dataset", "license", "licenseUrl", "attribution"):
		if not str(provider.get(key, "")).strip():
			raise ValueError(f"World Bank electrification provider metadata is missing {key}.")
	distribution = payload.get("distribution")
	if not isinstance(distribution, dict) or not str(distribution.get("provider", "")).strip():
		raise ValueError("World Bank electrification distribution metadata is invalid.")
	fallback = payload.get("wdiFallback")
	if not isinstance(fallback, dict) or not str(fallback.get("apiBase", "")).startswith("https://"):
		raise ValueError("World Bank electrification WDI fallback metadata is invalid.")
	if not isinstance(fallback.get("sourceId"), int) or int(fallback["sourceId"]) <= 0:
		raise ValueError("World Bank electrification WDI fallback sourceId is invalid.")
	if not isinstance(fallback.get("maximumYear"), int):
		raise ValueError("World Bank electrification WDI fallback maximumYear is invalid.")
	for key in ("providerId", "providerName"):
		if not str(fallback.get(key, "")).strip():
			raise ValueError(f"World Bank electrification WDI fallback metadata is missing {key}.")
	invariants = payload.get("sourceInvariants")
	if not isinstance(invariants, dict) or not invariants:
		raise ValueError("World Bank electrification source invariants are missing.")
	nature_labels = payload.get("natureLabels")
	if not isinstance(nature_labels, dict) or not nature_labels:
		raise ValueError("World Bank electrification nature labels are missing.")
	expected_ignored = payload.get("expectedIgnoredIso3Codes")
	if not isinstance(expected_ignored, list):
		raise ValueError("World Bank electrification expectedIgnoredIso3Codes is invalid.")
	indicators = payload.get("indicators")
	if not isinstance(indicators, list) or len(indicators) != 1:
		raise ValueError("World Bank electrification config requires exactly one indicator.")
	indicator = indicators[0]
	for key in (
		"id", "slug", "sourceIndicator", "sourceIndicatorName", "fallbackWdiIndicator",
		"title", "description",
	):
		if not str(indicator.get(key, "")).strip():
			raise ValueError(f"World Bank electrification indicator is missing {key}.")
	unit = indicator.get("unit")
	if not isinstance(unit, dict) or not str(unit.get("id", "")).strip() or not str(unit.get("label", "")).strip():
		raise ValueError("World Bank electrification indicator has invalid unit metadata.")
	common.validate_classification(str(indicator["id"]), indicator.get("classification"))
	for key in (
		"minAreasWithAnyValue", "minAreasInDefaultYear", "minimumDirectObservations",
		"minimumLatestYear", "minimumDefaultYear", "maximumDirectStartYear",
	):
		value = indicator.get(key)
		if not isinstance(value, int) or value <= 0:
			raise ValueError(f"World Bank electrification indicator has invalid {key}.")
	max_difference = indicator.get("maximumOverlapDifference")
	if not isinstance(max_difference, (int, float)) or float(max_difference) < 0:
		raise ValueError("World Bank electrification maximumOverlapDifference is invalid.")
	required_areas = indicator.get("requiredAreas")
	if not isinstance(required_areas, list) or not required_areas:
		raise ValueError("World Bank electrification indicator requires requiredAreas.")
	return payload


def download_workbook(url: str, timeout: int) -> tuple[Path, dict[str, Any]]:
	last_error: Exception | None = None
	for attempt, delay in enumerate((0, 3, 10, 20, 40), start=1):
		if delay:
			time.sleep(delay)
		handle = tempfile.NamedTemporaryFile(prefix="world-bank-electrification-", suffix=".xlsx", delete=False)
		path = Path(handle.name)
		handle.close()
		try:
			request = Request(url, headers={"Accept": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet,*/*", "User-Agent": USER_AGENT})
			with urlopen(request, timeout=timeout) as response, path.open("wb") as output:
				content_type = str(response.headers.get("Content-Type") or "")
				last_modified = str(response.headers.get("Last-Modified") or "").strip() or None
				while True:
					chunk = response.read(1024 * 1024)
					if not chunk:
						break
					output.write(chunk)
			bytes_count = path.stat().st_size
			if bytes_count < 2_000_000:
				raise RuntimeError(f"Electrification workbook unexpectedly small: {bytes_count} bytes")
			return path, {
				"bytes": bytes_count,
				"contentType": content_type,
				"lastModified": last_modified,
			}
		except (HTTPError, URLError, TimeoutError, OSError, RuntimeError) as error:
			last_error = error
			path.unlink(missing_ok=True)
			print(f"Electrification workbook download failed ({attempt}/5): {url}: {error}")
	if last_error is None:
		raise RuntimeError("Electrification workbook download failed without an exception.")
	raise RuntimeError("Electrification workbook download failed after 5 attempts.") from last_error


def numeric(value: Any) -> int | float | None:
	if value is None:
		return None
	if isinstance(value, str) and value.strip().upper() in ("", "NA", "N/A", "..", "..."):
		return None
	return common.normalize_number(value)


def validate_percentage(indicator_id: str, area_id: str, year: int, value: int | float, source: str) -> None:
	if not math.isfinite(float(value)) or float(value) < 0 or float(value) > 100:
		raise RuntimeError(f"Invalid {source} percentage for {indicator_id} {area_id} {year}: {value}")


def parse_direct_values(
	path: Path,
	config: dict[str, Any],
	indicator: dict[str, Any],
	area_by_iso3: dict[str, str],
	now_year: int,
) -> tuple[dict[int, dict[str, int | float]], dict[int, dict[str, dict[str, Any]]], dict[str, Any]]:
	workbook = load_workbook(path, read_only=True, data_only=True)
	if SOURCE_SHEET not in workbook.sheetnames:
		raise RuntimeError(f"Electrification workbook is missing sheet {SOURCE_SHEET!r}.")
	worksheet = workbook[SOURCE_SHEET]
	rows = worksheet.iter_rows(values_only=True)
	header = next(rows, None)
	if header is None:
		raise RuntimeError("Electrification UN reporting sheet is empty.")
	column_index = {str(value).strip(): index for index, value in enumerate(header) if value is not None}
	required_columns = {
		"SeriesCode", "Indicator", "GeoAreaName/Reference Area Name", "TimePeriod", "Value",
		"Units", "Nature", "Location", "Reporting Type", "Source", "ISOalpha3", "Type",
	}
	missing_columns = sorted(required_columns - set(column_index))
	if missing_columns:
		raise RuntimeError(f"Electrification workbook is missing columns: {', '.join(missing_columns)}")

	invariants = config["sourceInvariants"]
	anchors = ("SeriesCode", "Indicator", "Location", "Type")
	source_rows: list[tuple[Any, ...]] = []
	for row in rows:
		if all(str(row[column_index[field]] or "").strip() == str(invariants[field]) for field in anchors):
			source_rows.append(row)
	if not source_rows:
		raise RuntimeError("Electrification workbook returned no direct source rows.")
	for field, expected in invariants.items():
		observed = {str(row[column_index[field]] or "").strip() for row in source_rows}
		if observed != {str(expected)}:
			raise RuntimeError(f"Unexpected electrification {field} values: {sorted(observed)}")

	values_by_year: dict[int, dict[str, int | float]] = {}
	metadata_by_year: dict[int, dict[str, dict[str, Any]]] = {}
	ignored_codes: set[str] = set()
	nature_counts: Counter[str] = Counter()
	nature_labels = config["natureLabels"]
	direct_observations = 0
	for row in source_rows:
		iso3 = str(row[column_index["ISOalpha3"]] or "").strip().upper()
		if iso3 not in area_by_iso3:
			if iso3:
				ignored_codes.add(iso3)
			continue
		year_raw = row[column_index["TimePeriod"]]
		if isinstance(year_raw, bool) or not isinstance(year_raw, (int, float)):
			continue
		year = int(year_raw)
		if year > now_year:
			continue
		value = numeric(row[column_index["Value"]])
		if value is None:
			continue
		area_id = area_by_iso3[iso3]
		validate_percentage(str(indicator["id"]), area_id, year, value, "direct electrification")
		year_values = values_by_year.setdefault(year, {})
		if area_id in year_values:
			raise RuntimeError(f"Duplicate direct electrification value for {area_id} {year}.")
		year_values[area_id] = value
		nature_code = str(row[column_index["Nature"]] or "").strip()
		if nature_code not in nature_labels:
			raise RuntimeError(f"Unknown electrification Nature code: {nature_code!r}")
		metadata_by_year.setdefault(year, {})[area_id] = {
			"natureCode": nature_code,
			"natureLabel": str(nature_labels[nature_code]),
		}
		nature_counts[nature_code] += 1
		direct_observations += 1

	if not values_by_year:
		raise RuntimeError("Electrification workbook returned no mapped direct values.")
	expected_ignored = sorted(str(code).strip().upper() for code in config["expectedIgnoredIso3Codes"])
	if sorted(ignored_codes) != expected_ignored:
		raise RuntimeError(
			f"Unexpected ignored electrification ISO3 codes: {sorted(ignored_codes)}; expected {expected_ignored}."
		)
	if direct_observations < int(indicator["minimumDirectObservations"]):
		raise RuntimeError(f"Direct electrification observation count unexpectedly small: {direct_observations}.")
	if min(values_by_year) > int(indicator["maximumDirectStartYear"]):
		raise RuntimeError(f"Direct electrification history starts unexpectedly late: {min(values_by_year)}.")
	return values_by_year, metadata_by_year, {
		"worksheetRows": worksheet.max_row,
		"worksheetColumns": worksheet.max_column,
		"sourceRows": len(source_rows),
		"directObservations": direct_observations,
		"directStartYear": min(values_by_year),
		"directEndYear": max(values_by_year),
		"ignoredCodes": sorted(ignored_codes),
		"natureCounts": dict(sorted(nature_counts.items())),
	}


def add_wdi_fallback(
	config: dict[str, Any],
	indicator: dict[str, Any],
	values_by_year: dict[int, dict[str, int | float]],
	metadata_by_year: dict[int, dict[str, dict[str, Any]]],
	area_by_iso3: dict[str, str],
	now_year: int,
	timeout: int,
) -> dict[str, Any]:
	fallback = config["wdiFallback"]
	fallback_code = str(indicator["fallbackWdiIndicator"])
	records = common.fetch_world_bank_records(
		str(fallback["apiBase"]).rstrip("/"),
		f"country/all/indicator/{fallback_code}",
		int(fallback["sourceId"]),
		timeout,
	)
	fallback_count = 0
	fallback_areas: set[str] = set()
	fallback_years: set[int] = set()
	overlap_count = 0
	overlap_differences = 0
	max_overlap_difference = 0.0
	unexpected_fallbacks: list[tuple[str, int]] = []
	maximum_fallback_year = int(fallback["maximumYear"])

	for record in records:
		if record.get("value") is None or str(record.get("obs_status", "")).strip().upper() == "F":
			continue
		iso3 = str(record.get("countryiso3code", "")).strip().upper()
		if iso3 not in area_by_iso3:
			continue
		year_text = str(record.get("date", "")).strip()
		if len(year_text) != 4 or not year_text.isdigit():
			continue
		year = int(year_text)
		if year > now_year:
			continue
		area_id = area_by_iso3[iso3]
		value = common.normalize_number(record["value"])
		validate_percentage(str(indicator["id"]), area_id, year, value, "WDI fallback")
		year_values = values_by_year.setdefault(year, {})
		if area_id in year_values:
			overlap_count += 1
			difference = abs(float(year_values[area_id]) - float(value))
			max_overlap_difference = max(max_overlap_difference, difference)
			if difference > 1e-9:
				overlap_differences += 1
			continue
		if year > maximum_fallback_year:
			unexpected_fallbacks.append((area_id, year))
			continue
		year_values[area_id] = value
		metadata_by_year.setdefault(year, {})[area_id] = {
			"fallback": True,
			"sourceProviderId": str(fallback["providerId"]),
			"sourceProviderName": str(fallback["providerName"]),
			"sourceIndicator": fallback_code,
			"provenance": "Older World Bank WDI distribution of the same Global Electrification Database series; used only for historical country-year observations unavailable in the current Tracking SDG 7 workbook",
		}
		fallback_count += 1
		fallback_areas.add(area_id)
		fallback_years.add(year)

	if unexpected_fallbacks:
		sample = ", ".join(f"{area_id} {year}" for area_id, year in unexpected_fallbacks[:10])
		raise RuntimeError(
			f"WDI would be required after configured historical fallback period ({maximum_fallback_year}): {sample}"
		)
	if max_overlap_difference > float(indicator["maximumOverlapDifference"]):
		raise RuntimeError(
			f"Direct/WDI electrification overlap differs too much: max {max_overlap_difference}."
		)
	return {
		"fallbackObservations": fallback_count,
		"fallbackAreas": len(fallback_areas),
		"fallbackYears": sorted(fallback_years),
		"overlapObservations": overlap_count,
		"overlapDifferences": overlap_differences,
		"maxOverlapDifference": max_overlap_difference,
	}


def build_indicator_payload(
	config: dict[str, Any],
	indicator: dict[str, Any],
	path: Path,
	area_by_iso3: dict[str, str],
	now_year: int,
	timeout: int,
) -> tuple[dict[str, Any], dict[str, Any]]:
	values_by_year, metadata_by_year, direct_stats = parse_direct_values(
		path, config, indicator, area_by_iso3, now_year
	)
	fallback_stats = add_wdi_fallback(
		config, indicator, values_by_year, metadata_by_year, area_by_iso3, now_year, timeout
	)
	available_years = sorted(year for year, values in values_by_year.items() if values)
	areas_with_any_value = set().union(*(set(values_by_year[year]) for year in available_years))
	if len(areas_with_any_value) < int(indicator["minAreasWithAnyValue"]):
		raise RuntimeError(f"Electrification coverage too small: {len(areas_with_any_value)} areas.")
	default_candidates = [
		year for year in available_years
		if len(values_by_year[year]) >= int(indicator["minAreasInDefaultYear"])
	]
	if not default_candidates:
		raise RuntimeError("Electrification has no sufficiently covered default year.")
	default_year = default_candidates[-1]
	if available_years[-1] < int(indicator["minimumLatestYear"]):
		raise RuntimeError(f"Electrification latest year is too old: {available_years[-1]}.")
	if default_year < int(indicator["minimumDefaultYear"]):
		raise RuntimeError(f"Electrification default year is too old: {default_year}.")
	for area_id in indicator["requiredAreas"]:
		if area_id not in values_by_year[default_year]:
			raise RuntimeError(f"Required area {area_id} has no electrification value in {default_year}.")

	provider = config["provider"]
	fallback = config["wdiFallback"]
	payload: dict[str, Any] = {
		"schema": "kartensammlung.statistics-indicator/v1",
		"indicator": {
			"id": indicator["id"],
			"title": indicator["title"],
			"description": indicator["description"],
			"areaLevel": "country",
			"frequency": "annual",
			"unit": indicator["unit"],
			"classification": indicator["classification"],
		},
		"source": {
			"providerId": provider["id"],
			"providerName": provider["name"],
			"dataset": provider["dataset"],
			"indicator": indicator["sourceIndicator"],
			"indicatorName": indicator["sourceIndicatorName"],
			"provenance": "World Bank Global Electrification Database distributed through the Tracking SDG 7 electrification workbook",
			"license": provider["license"],
			"licenseUrl": provider["licenseUrl"],
			"attribution": provider["attribution"],
			"url": config["sourceUrl"],
			"downloadUrl": config["downloadUrl"],
			"catalogUrl": config["catalogUrl"],
			"distribution": config["distribution"],
			"sheet": SOURCE_SHEET,
			"filters": config["sourceInvariants"],
			"natureLabels": config["natureLabels"],
			"timeCoverage": {
				"startYear": direct_stats["directStartYear"],
				"endYear": direct_stats["directEndYear"],
			},
			"fallback": {
				"providerId": fallback["providerId"],
				"providerName": fallback["providerName"],
				"indicator": indicator["fallbackWdiIndicator"],
				"maximumYear": fallback["maximumYear"],
				"usage": "Only historical country-year observations missing from the current direct workbook",
				"provenance": "Older World Bank WDI distribution of the same Global Electrification Database series",
			},
		},
		"availableYears": available_years,
		"defaultYear": default_year,
		"coverage": {
			"registryAreas": len(area_by_iso3),
			"areasWithAnyValue": len(areas_with_any_value),
			"latestYear": available_years[-1],
			"areasInLatestYear": len(values_by_year[available_years[-1]]),
			"areasInDefaultYear": len(values_by_year[default_year]),
			"observations": sum(len(values) for values in values_by_year.values()),
			"directObservations": direct_stats["directObservations"],
			"fallbackObservations": fallback_stats["fallbackObservations"],
			"fallbackAreas": fallback_stats["fallbackAreas"],
			"directNatureCounts": direct_stats["natureCounts"],
		},
		"values": {
			str(year): {area_id: values_by_year[year][area_id] for area_id in sorted(values_by_year[year])}
			for year in available_years
		},
		"observationMetadata": {
			str(year): {area_id: metadata_by_year[year][area_id] for area_id in sorted(metadata_by_year[year])}
			for year in sorted(metadata_by_year)
		},
	}
	stats = {**direct_stats, **fallback_stats}
	print(
		f"{indicator['id']}: direct={stats['directObservations']} fallback={stats['fallbackObservations']} "
		f"areas={len(areas_with_any_value)} years={available_years[0]}-{available_years[-1]} "
		f"defaultYear={default_year} defaultCoverage={len(values_by_year[default_year])} "
		f"overlap={stats['overlapObservations']} overlapDiffs={stats['overlapDifferences']} "
		f"maxOverlapDiff={stats['maxOverlapDifference']} ignored={stats['ignoredCodes']}"
	)
	return payload, stats


def main() -> None:
	args = parse_args()
	config = validate_config(common.read_json_path(args.config))
	provider_catalog = common.validate_provider_catalog(common.read_json_path(args.providers))
	provider = config["provider"]
	if not any(item["id"] == provider["id"] for item in provider_catalog["providers"]):
		raise RuntimeError("World Bank Global Electrification Database is missing from statistics-providers.json.")
	area_by_iso3, registry_payload = common.load_registry(args.registry, args.timeout)
	now = common.utc_now()
	path, download_metadata = download_workbook(str(config["downloadUrl"]), args.timeout)
	try:
		payloads: list[tuple[dict[str, Any], dict[str, Any], dict[str, Any]]] = []
		for indicator in config["indicators"]:
			payload, stats = build_indicator_payload(
				config, indicator, path, area_by_iso3, now.year, args.timeout
			)
			payloads.append((indicator, payload, stats))
	finally:
		path.unlink(missing_ok=True)

	hasher = hashlib.sha256()
	for indicator, payload, _ in sorted(payloads, key=lambda item: item[0]["id"]):
		hasher.update(str(indicator["id"]).encode("utf-8") + b"\0" + common.canonical_bytes(payload) + b"\0")
	snapshot = hasher.hexdigest()[:16]
	provider_dir = args.output_dir / provider["id"]
	release_dir = provider_dir / "releases" / snapshot
	index_indicators: list[dict[str, Any]] = []
	for indicator, payload, _ in payloads:
		filename = f"{indicator['slug']}.json"
		common.write_json(release_dir / filename, payload)
		index_indicators.append({
			"id": indicator["id"],
			"title": indicator["title"],
			"description": indicator["description"],
			"areaLevel": "country",
			"frequency": "annual",
			"unit": indicator["unit"],
			"classification": indicator["classification"],
			"sourceIndicator": indicator["sourceIndicator"],
			"path": f"releases/{snapshot}/{filename}",
			"availableYears": payload["availableYears"],
			"defaultYear": payload["defaultYear"],
			"coverage": payload["coverage"],
		})

	registry_source = {
		"url": args.registry if args.registry.startswith(("https://", "http://")) else None,
		"schema": registry_payload.get("schema"),
		"generatedAt": registry_payload.get("generatedAt"),
		"areaCount": len(area_by_iso3),
	}
	if registry_source["url"] is None:
		registry_source.pop("url")
	first_stats = payloads[0][2]
	provider_index = {
		"schema": "kartensammlung.statistics-provider-index/v1",
		"provider": provider,
		"retrievedAt": now.isoformat().replace("+00:00", "Z"),
		"activeSnapshot": snapshot,
		"areaRegistry": registry_source,
		"dataset": {
			"sourceUrl": config["sourceUrl"],
			"catalogUrl": config["catalogUrl"],
			"downloadUrl": config["downloadUrl"],
			"distribution": config["distribution"],
			"sheet": SOURCE_SHEET,
			"bytes": download_metadata["bytes"],
			"contentType": download_metadata["contentType"],
			"lastModified": download_metadata["lastModified"],
			"worksheetRows": first_stats["worksheetRows"],
			"worksheetColumns": first_stats["worksheetColumns"],
			"sourceRows": first_stats["sourceRows"],
			"ignoredIso3Codes": first_stats["ignoredCodes"],
		},
		"indicators": index_indicators,
		"notes": [
			"World Bank Global Electrification Database values from the current Tracking SDG 7 workbook are canonical.",
			"World Bank WDI is used only for historical observations through 1999 that are absent from the current direct workbook.",
			"Direct observation Nature codes are retained in observationMetadata; WDI fallback observations are explicitly marked there as fallback.",
			"The World Bank aggregate CHI (Channel Islands) is intentionally not mapped to Jersey or Guernsey.",
		],
	}
	common.write_json(provider_dir / "index.json", provider_index)
	common.write_json(args.output_dir / "index.json", common.build_global_index(provider_catalog))
	print(f"Built World Bank electrification snapshot {snapshot}")
	print(f"Provider index: {provider_dir / 'index.json'}")
	print(f"Indicators: {len(index_indicators)}")


if __name__ == "__main__":
	main()
