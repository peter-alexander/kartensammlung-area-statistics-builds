#!/usr/bin/env python3
from __future__ import annotations

import argparse
import hashlib
import json
import math
import tempfile
import time
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any
from urllib.error import HTTPError, URLError
from urllib.parse import urlencode
from urllib.request import Request, urlopen

from openpyxl import load_workbook

import build_world_bank as common

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_CONFIG = ROOT / "config" / "edgar-indicators.json"
DEFAULT_PROVIDERS = ROOT / "config" / "statistics-providers.json"
DEFAULT_OUTPUT_DIR = ROOT / "dist" / "statistics"
DEFAULT_REGISTRY = "https://tiles.radlobby.at/AreaStatistics/area-registry-countries.json"
USER_AGENT = "kartensammlung-area-statistics-builds/1"


def parse_args() -> argparse.Namespace:
	parser = argparse.ArgumentParser(
		description="Build country methane and nitrous-oxide statistics from the EDGAR Community GHG Database."
	)
	parser.add_argument("--config", type=Path, default=DEFAULT_CONFIG)
	parser.add_argument("--providers", type=Path, default=DEFAULT_PROVIDERS)
	parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
	parser.add_argument("--registry", default=DEFAULT_REGISTRY)
	parser.add_argument("--timeout", type=int, default=120)
	return parser.parse_args()


def validate_config(payload: Any) -> dict[str, Any]:
	if not isinstance(payload, dict) or payload.get("schema") != "kartensammlung.edgar-statistics/v1":
		raise ValueError("Invalid EDGAR statistics config.")
	for key in ("downloadUrl", "sourceUrl", "datasetUrl"):
		if not str(payload.get(key, "")).startswith("https://"):
			raise ValueError(f"EDGAR {key} must use HTTPS.")
	provider = payload.get("provider")
	if not isinstance(provider, dict) or provider.get("id") != "edgar":
		raise ValueError("EDGAR provider metadata is invalid.")
	for key in ("name", "dataset", "license", "licenseUrl", "attribution"):
		if not str(provider.get(key, "")).strip():
			raise ValueError(f"EDGAR provider metadata is missing {key}.")
	if payload.get("sheet") != "GHG_by_sector_and_country":
		raise ValueError("EDGAR source sheet must be GHG_by_sector_and_country.")

	invariants = payload.get("sourceInvariants")
	if not isinstance(invariants, dict):
		raise ValueError("EDGAR source invariants are missing.")
	if invariants.get("startYear") != 1970 or invariants.get("endYear") != 2025:
		raise ValueError("EDGAR source year invariants are invalid.")
	sectors = invariants.get("sectors")
	if not isinstance(sectors, list) or len(sectors) != 8 or len(set(sectors)) != 8:
		raise ValueError("EDGAR sector invariants are invalid.")
	allowed_substances = invariants.get("allowedSubstances")
	if allowed_substances != ["GWP_100_AR5_CH4", "GWP_100_AR5_N2O"]:
		raise ValueError("EDGAR provider must be restricted to CH4 and N2O AR5 substances.")

	excluded = payload.get("excludedCombinedAreas")
	if not isinstance(excluded, dict) or set(excluded) != {"CHE", "ESP", "FRA", "ISR", "ITA", "SDN"}:
		raise ValueError("EDGAR excludedCombinedAreas must contain the six audited grouped country codes.")
	if any(len(str(code)) != 3 or not str(name).strip() for code, name in excluded.items()):
		raise ValueError("EDGAR excludedCombinedAreas contains invalid entries.")

	fallback = payload.get("wdiFallback")
	if not isinstance(fallback, dict) or not str(fallback.get("apiBase", "")).startswith("https://"):
		raise ValueError("EDGAR WDI fallback metadata is invalid.")
	if not isinstance(fallback.get("sourceId"), int) or int(fallback["sourceId"]) <= 0:
		raise ValueError("EDGAR WDI fallback sourceId is invalid.")
	if fallback.get("maximumYear") != 2024:
		raise ValueError("EDGAR WDI fallback maximumYear must be 2024 for the audited release.")
	for key in ("providerId", "providerName"):
		if not str(fallback.get(key, "")).strip():
			raise ValueError(f"EDGAR WDI fallback metadata is missing {key}.")

	indicators = payload.get("indicators")
	if not isinstance(indicators, list) or len(indicators) != 2:
		raise ValueError("EDGAR config requires exactly CH4 and N2O indicators.")
	expected_sources = {"GWP_100_AR5_CH4", "GWP_100_AR5_N2O"}
	seen_sources: set[str] = set()
	ids: set[str] = set()
	slugs: set[str] = set()
	for indicator in indicators:
		if not isinstance(indicator, dict):
			raise ValueError("EDGAR indicator entries must be objects.")
		for key in ("id", "slug", "sourceIndicator", "fallbackWdiIndicator", "title", "description"):
			if not str(indicator.get(key, "")).strip():
				raise ValueError(f"EDGAR indicator entry is missing {key}.")
		indicator_id = str(indicator["id"])
		slug = str(indicator["slug"])
		source_indicator = str(indicator["sourceIndicator"])
		if indicator_id in ids or slug in slugs or source_indicator in seen_sources:
			raise ValueError(f"Duplicate EDGAR indicator metadata for {indicator_id}.")
		ids.add(indicator_id)
		slugs.add(slug)
		seen_sources.add(source_indicator)
		unit = indicator.get("unit")
		if not isinstance(unit, dict) or unit.get("id") != "mt-co2e":
			raise ValueError(f"EDGAR indicator {indicator_id} must use mt-co2e.")
		common.validate_classification(indicator_id, indicator.get("classification"))
		for key in (
			"minAreasWithAnyValue", "minAreasInDefaultYear", "minimumDirectObservations",
			"minimumLatestYear", "minimumDefaultYear",
		):
			value = indicator.get(key)
			if not isinstance(value, int) or value <= 0:
				raise ValueError(f"EDGAR indicator {indicator_id} has invalid {key}.")
		required_areas = indicator.get("requiredAreas")
		if not isinstance(required_areas, list) or not required_areas:
			raise ValueError(f"EDGAR indicator {indicator_id} requires requiredAreas.")
	if seen_sources != expected_sources:
		raise ValueError(f"Unexpected EDGAR source substances: {sorted(seen_sources)}")
	return payload


def download_workbook(url: str, timeout: int) -> tuple[Path, dict[str, Any]]:
	last_error: Exception | None = None
	for attempt, delay in enumerate((0, 3, 10, 20, 40), start=1):
		if delay:
			time.sleep(delay)
		handle = tempfile.NamedTemporaryFile(prefix="edgar-2026-", suffix=".xlsx", delete=False)
		path = Path(handle.name)
		handle.close()
		try:
			request = Request(
				url,
				headers={
					"Accept": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet,*/*",
					"User-Agent": USER_AGENT,
				},
			)
			with urlopen(request, timeout=timeout) as response, path.open("wb") as output:
				content_type = str(response.headers.get("Content-Type") or "")
				last_modified = str(response.headers.get("Last-Modified") or "").strip() or None
				while True:
					chunk = response.read(1024 * 1024)
					if not chunk:
						break
					output.write(chunk)
			bytes_count = path.stat().st_size
			if bytes_count < 4_000_000:
				raise RuntimeError(f"EDGAR workbook unexpectedly small: {bytes_count} bytes")
			return path, {
				"bytes": bytes_count,
				"contentType": content_type,
				"lastModified": last_modified,
			}
		except (HTTPError, URLError, TimeoutError, OSError, RuntimeError) as error:
			last_error = error
			path.unlink(missing_ok=True)
			print(f"EDGAR workbook download failed ({attempt}/5): {url}: {error}")
	if last_error is None:
		raise RuntimeError("EDGAR workbook download failed without an exception.")
	raise RuntimeError("EDGAR workbook download failed after 5 attempts.") from last_error


def validate_workbook_metadata(workbook: Any, config: dict[str, Any]) -> dict[str, Any]:
	for sheet in ("info", "Citations and references", str(config["sheet"])):
		if sheet not in workbook.sheetnames:
			raise RuntimeError(f"EDGAR workbook is missing sheet {sheet!r}.")

	info_text = "\n".join(
		str(value)
		for row in workbook["info"].iter_rows(values_only=True)
		for value in row
		if value is not None
	)
	if "GHG emissions of all world countries, 2026 Report" not in info_text:
		raise RuntimeError("Unexpected EDGAR report edition.")
	if "IPCC AR5 (GWP-100 AR5)" not in info_text:
		raise RuntimeError("EDGAR workbook no longer declares IPCC AR5 GWP-100.")
	if "GHG_by_sector_and_country sheet are expressed in Mt CO2eq/yr" not in info_text:
		raise RuntimeError("EDGAR source unit declaration changed.")

	citation_text = "\n".join(
		str(value)
		for row in workbook["Citations and references"].iter_rows(values_only=True)
		for value in row
		if value is not None
	)
	if "Creative Commons Attribution 4.0 International" not in citation_text:
		raise RuntimeError("EDGAR workbook no longer declares the EU CC BY 4.0 reuse terms.")
	if "IEA-EDGAR CO2" not in citation_text or "CC BY-NC-ND 4.0" not in citation_text:
		raise RuntimeError("EDGAR workbook no longer exposes the separate restricted fossil-CO2 licence exception.")
	return {
		"report": "GHG emissions of all world countries, 2026 Report",
		"gwp": str(config["sourceInvariants"]["gwp"]),
		"unit": str(config["sourceInvariants"]["unit"]),
	}


def numeric(value: Any) -> int | float | None:
	if value is None or (isinstance(value, str) and value.strip().upper() in ("", "NA", "N/A", "..", "...")):
		return None
	return common.normalize_number(value)


def parse_direct_values(
	path: Path,
	config: dict[str, Any],
	area_by_iso3: dict[str, str],
	now_year: int,
) -> tuple[dict[str, dict[int, dict[str, int | float]]], dict[str, dict[str, Any]], dict[str, Any]]:
	workbook = load_workbook(path, read_only=True, data_only=True)
	workbook_metadata = validate_workbook_metadata(workbook, config)
	worksheet = workbook[str(config["sheet"])]
	rows = worksheet.iter_rows(values_only=True)
	header = next(rows, None)
	if header is None:
		raise RuntimeError("EDGAR source sheet is empty.")
	column_index = {str(value).strip(): index for index, value in enumerate(header) if value is not None}
	required_columns = {"Substance", "Sector", "EDGAR Country Code", "Country"}
	missing = sorted(required_columns - set(column_index))
	if missing:
		raise RuntimeError(f"EDGAR source sheet is missing columns: {', '.join(missing)}")

	start_year = int(config["sourceInvariants"]["startYear"])
	end_year = int(config["sourceInvariants"]["endYear"])
	year_columns = {
		int(value): index
		for index, value in enumerate(header)
		if isinstance(value, int) and start_year <= int(value) <= end_year
	}
	if set(year_columns) != set(range(start_year, end_year + 1)):
		raise RuntimeError(
			f"Unexpected EDGAR source years: {min(year_columns, default=None)}-"
			f"{max(year_columns, default=None)}, count={len(year_columns)}."
		)

	allowed_substances = set(str(value) for value in config["sourceInvariants"]["allowedSubstances"])
	expected_sectors = set(str(value) for value in config["sourceInvariants"]["sectors"])
	excluded = {str(code): str(name) for code, name in config["excludedCombinedAreas"].items()}
	values_by_source: dict[str, dict[int, dict[str, int | float]]] = {
		source: {} for source in allowed_substances
	}
	sector_values: dict[str, dict[str, set[str]]] = {
		source: defaultdict(set) for source in allowed_substances
	}
	totals: dict[str, dict[tuple[str, int], float]] = {
		source: defaultdict(float) for source in allowed_substances
	}
	excluded_rows: Counter[tuple[str, str]] = Counter()
	ignored_codes: Counter[str] = Counter()
	source_rows: Counter[str] = Counter()

	for row in rows:
		substance = str(row[column_index["Substance"]] or "").strip()
		if substance not in allowed_substances:
			continue
		source_rows[substance] += 1
		sector = str(row[column_index["Sector"]] or "").strip()
		if sector not in expected_sectors:
			raise RuntimeError(f"Unexpected EDGAR sector for {substance}: {sector!r}")
		iso3 = str(row[column_index["EDGAR Country Code"]] or "").strip().upper()
		name = str(row[column_index["Country"]] or "").strip()
		if iso3 in excluded:
			if name != excluded[iso3]:
				raise RuntimeError(
					f"EDGAR grouped area {iso3} changed from {excluded[iso3]!r} to {name!r}; audit required."
				)
			excluded_rows[(substance, iso3)] += 1
			continue
		if iso3 not in area_by_iso3:
			if iso3:
				ignored_codes[iso3] += 1
			continue
		area_id = area_by_iso3[iso3]
		sector_values[substance][area_id].add(sector)
		for year, column in year_columns.items():
			if year > now_year:
				continue
			value = numeric(row[column])
			if value is None:
				continue
			if not math.isfinite(float(value)) or float(value) < 0:
				raise RuntimeError(f"Invalid EDGAR value for {substance} {area_id} {year}: {value}")
			totals[substance][(area_id, year)] += float(value)

	expected_ignored = {"AIR", "ANT", "EU27", "GLOBAL TOTAL", "SCG", "SEA"}
	if set(ignored_codes) != expected_ignored:
		raise RuntimeError(
			f"Unexpected ignored EDGAR country codes: {sorted(ignored_codes)}; expected {sorted(expected_ignored)}."
		)
	for substance in sorted(allowed_substances):
		if source_rows[substance] < 1500:
			raise RuntimeError(f"EDGAR source row count unexpectedly small for {substance}: {source_rows[substance]}")
		if set(sector_values[substance]) != set(sector_values[next(iter(allowed_substances))]):
			# Both gases should expose the same safe country geography in this report edition.
			pass
		for area_id, sectors in sector_values[substance].items():
			if sectors != expected_sectors:
				raise RuntimeError(
					f"EDGAR {substance} sectors incomplete for {area_id}: {sorted(sectors)}"
				)
		if len(sector_values[substance]) != 200:
			raise RuntimeError(
				f"EDGAR safe direct geography changed for {substance}: {len(sector_values[substance])} areas; expected 200."
			)
		if {iso3 for source, iso3 in excluded_rows if source == substance} != set(excluded):
			raise RuntimeError(f"EDGAR grouped-area exclusions incomplete for {substance}.")
		for (area_id, year), value in totals[substance].items():
			values_by_source[substance].setdefault(year, {})[area_id] = common.normalize_number(value)
		if set(values_by_source[substance]) != set(range(start_year, min(end_year, now_year) + 1)):
			raise RuntimeError(f"EDGAR year coverage is incomplete for {substance}.")

	stats = {
		"worksheetRows": worksheet.max_row,
		"worksheetColumns": worksheet.max_column,
		"sourceRows": dict(sorted(source_rows.items())),
		"ignoredCodes": sorted(ignored_codes),
		"ignoredCodeRows": dict(sorted(ignored_codes.items())),
		"excludedCombinedAreas": dict(sorted(excluded.items())),
		"excludedRows": {
			f"{source}:{iso3}": count for (source, iso3), count in sorted(excluded_rows.items())
		},
		"workbookMetadata": workbook_metadata,
	}
	per_source_stats: dict[str, dict[str, Any]] = {}
	for substance, values_by_year in values_by_source.items():
		areas = set().union(*(set(values) for values in values_by_year.values()))
		observations = sum(len(values) for values in values_by_year.values())
		per_source_stats[substance] = {
			"directAreas": len(areas),
			"directObservations": observations,
			"directStartYear": min(values_by_year),
			"directEndYear": max(values_by_year),
		}
	stats["perSource"] = per_source_stats
	return values_by_source, per_source_stats, stats


def fetch_wdi_records(
	api_base: str,
	source_id: int,
	indicator: str,
	start_year: int,
	end_year: int,
	timeout: int,
) -> list[dict[str, Any]]:
	query = {
		"format": "json",
		"source": str(source_id),
		"per_page": "20000",
		"page": "1",
		"date": f"{start_year}:{end_year}",
	}
	path = f"country/all/indicator/{indicator}"
	url = f"{api_base.rstrip('/')}/{path}?{urlencode(query)}"
	metadata, records = common.world_bank_payload(url, timeout)
	pages = int(metadata.get("pages", 1))
	all_records = list(records)
	for page in range(2, pages + 1):
		query["page"] = str(page)
		page_url = f"{api_base.rstrip('/')}/{path}?{urlencode(query)}"
		page_metadata, page_records = common.world_bank_payload(page_url, timeout)
		if int(page_metadata.get("page", page)) != page:
			raise RuntimeError(f"World Bank pagination mismatch for {indicator}: expected page {page}.")
		all_records.extend(page_records)
	return all_records


def add_wdi_fallback(
	config: dict[str, Any],
	indicator: dict[str, Any],
	values_by_year: dict[int, dict[str, int | float]],
	area_by_iso3: dict[str, str],
	now_year: int,
	timeout: int,
) -> tuple[dict[int, dict[str, dict[str, Any]]], dict[str, Any]]:
	fallback = config["wdiFallback"]
	maximum_year = min(int(fallback["maximumYear"]), now_year)
	minimum_year = int(config["sourceInvariants"]["startYear"])
	fallback_code = str(indicator["fallbackWdiIndicator"])
	records = fetch_wdi_records(
		str(fallback["apiBase"]),
		int(fallback["sourceId"]),
		fallback_code,
		minimum_year,
		maximum_year,
		timeout,
	)
	metadata_by_year: dict[int, dict[str, dict[str, Any]]] = {}
	fallback_count = 0
	fallback_areas: set[str] = set()
	fallback_years: set[int] = set()
	overlap_count = 0
	overlap_differences = 0
	max_overlap_difference = 0.0

	for record in records:
		if record.get("value") is None or str(record.get("obs_status", "")).strip().upper() == "F":
			continue
		iso3 = str(record.get("countryiso3code", "")).strip().upper()
		if iso3 not in area_by_iso3:
			continue
		year_text = str(record.get("date", "")).strip()
		if len(year_text) != 4 or not year_text.isdigit():
			continue
		year = int(year_text)
		if year < minimum_year or year > maximum_year:
			continue
		area_id = area_by_iso3[iso3]
		value = common.normalize_number(record["value"])
		year_values = values_by_year.setdefault(year, {})
		if area_id in year_values:
			overlap_count += 1
			difference = abs(float(year_values[area_id]) - float(value))
			max_overlap_difference = max(max_overlap_difference, difference)
			if difference > 1e-9:
				overlap_differences += 1
			continue
		year_values[area_id] = value
		metadata_by_year.setdefault(year, {})[area_id] = {
			"fallback": True,
			"sourceProviderId": str(fallback["providerId"]),
			"sourceProviderName": str(fallback["providerName"]),
			"sourceIndicator": fallback_code,
			"provenance": (
				"Older World Bank WDI distribution of the same EDGAR greenhouse-gas series; "
				"used only where the EDGAR 2026 report has no safe one-to-one country observation"
			),
		}
		fallback_count += 1
		fallback_areas.add(area_id)
		fallback_years.add(year)

	return metadata_by_year, {
		"fallbackObservations": fallback_count,
		"fallbackAreas": len(fallback_areas),
		"fallbackYears": sorted(fallback_years),
		"overlapObservations": overlap_count,
		"overlapDifferences": overlap_differences,
		"maxOverlapDifference": max_overlap_difference,
	}


def build_indicator_payload(
	config: dict[str, Any],
	indicator: dict[str, Any],
	direct_values_by_source: dict[str, dict[int, dict[str, int | float]]],
	direct_stats_by_source: dict[str, dict[str, Any]],
	area_by_iso3: dict[str, str],
	now_year: int,
	timeout: int,
) -> tuple[dict[str, Any], dict[str, Any]]:
	source_indicator = str(indicator["sourceIndicator"])
	values_by_year = {
		year: dict(values)
		for year, values in direct_values_by_source[source_indicator].items()
	}
	direct_stats = direct_stats_by_source[source_indicator]
	metadata_by_year, fallback_stats = add_wdi_fallback(
		config, indicator, values_by_year, area_by_iso3, now_year, timeout
	)
	available_years = sorted(year for year, values in values_by_year.items() if values)
	areas_with_any_value = set().union(*(set(values_by_year[year]) for year in available_years))
	if len(areas_with_any_value) < int(indicator["minAreasWithAnyValue"]):
		raise RuntimeError(
			f"EDGAR coverage too small for {indicator['id']}: {len(areas_with_any_value)} areas."
		)
	default_candidates = [
		year for year in available_years
		if len(values_by_year[year]) >= int(indicator["minAreasInDefaultYear"])
	]
	if not default_candidates:
		raise RuntimeError(f"EDGAR has no sufficiently covered default year for {indicator['id']}.")
	default_year = default_candidates[-1]
	if available_years[-1] < int(indicator["minimumLatestYear"]):
		raise RuntimeError(f"EDGAR latest year is too old for {indicator['id']}: {available_years[-1]}.")
	if default_year < int(indicator["minimumDefaultYear"]):
		raise RuntimeError(f"EDGAR default year is too old for {indicator['id']}: {default_year}.")
	for area_id in indicator["requiredAreas"]:
		if area_id not in values_by_year[default_year]:
			raise RuntimeError(f"Required area {area_id} has no direct EDGAR value in {default_year}.")
	if int(direct_stats["directObservations"]) < int(indicator["minimumDirectObservations"]):
		raise RuntimeError(
			f"EDGAR direct observation count too small for {indicator['id']}: "
			f"{direct_stats['directObservations']}."
		)

	provider = config["provider"]
	fallback = config["wdiFallback"]
	payload: dict[str, Any] = {
		"schema": "kartensammlung.statistics-indicator/v1",
		"indicator": {
			"id": indicator["id"],
			"title": indicator["title"],
			"description": indicator["description"],
			"areaLevel": "country",
			"frequency": "annual",
			"unit": indicator["unit"],
			"classification": indicator["classification"],
		},
		"source": {
			"providerId": provider["id"],
			"providerName": provider["name"],
			"dataset": provider["dataset"],
			"indicator": source_indicator,
			"provenance": (
				"European Commission JRC EDGAR 2026 country-sector greenhouse-gas data, "
				"already expressed as IPCC AR5 GWP-100 CO2-equivalents; only CH4 and N2O are redistributed"
			),
			"license": provider["license"],
			"licenseUrl": provider["licenseUrl"],
			"attribution": provider["attribution"],
			"url": config["sourceUrl"],
			"downloadUrl": config["downloadUrl"],
			"datasetUrl": config["datasetUrl"],
			"sheet": config["sheet"],
			"filters": {
				"substance": source_indicator,
				"sectors": config["sourceInvariants"]["sectors"],
				"excludedCombinedAreas": config["excludedCombinedAreas"],
			},
			"timeCoverage": {
				"startYear": direct_stats["directStartYear"],
				"endYear": direct_stats["directEndYear"],
			},
			"fallback": {
				"providerId": fallback["providerId"],
				"providerName": fallback["providerName"],
				"indicator": indicator["fallbackWdiIndicator"],
				"maximumYear": fallback["maximumYear"],
				"usage": "Only country-year observations without a safe one-to-one EDGAR 2026 report value",
				"provenance": "Older World Bank WDI distribution of the same EDGAR greenhouse-gas series",
			},
		},
		"availableYears": available_years,
		"defaultYear": default_year,
		"coverage": {
			"registryAreas": len(area_by_iso3),
			"areasWithAnyValue": len(areas_with_any_value),
			"latestYear": available_years[-1],
			"areasInLatestYear": len(values_by_year[available_years[-1]]),
			"areasInDefaultYear": len(values_by_year[default_year]),
			"observations": sum(len(values) for values in values_by_year.values()),
			"directObservations": direct_stats["directObservations"],
			"directAreas": direct_stats["directAreas"],
			"fallbackObservations": fallback_stats["fallbackObservations"],
			"fallbackAreas": fallback_stats["fallbackAreas"],
		},
		"values": {
			str(year): {area_id: values_by_year[year][area_id] for area_id in sorted(values_by_year[year])}
			for year in available_years
		},
		"observationMetadata": {
			str(year): {area_id: metadata_by_year[year][area_id] for area_id in sorted(metadata_by_year[year])}
			for year in sorted(metadata_by_year)
		},
	}
	stats = {**direct_stats, **fallback_stats}
	print(
		f"{indicator['id']}: direct={stats['directObservations']} fallback={stats['fallbackObservations']} "
		f"areas={len(areas_with_any_value)} years={available_years[0]}-{available_years[-1]} "
		f"defaultYear={default_year} defaultCoverage={len(values_by_year[default_year])} "
		f"overlap={stats['overlapObservations']} overlapDiffs={stats['overlapDifferences']} "
		f"maxOverlapDiff={stats['maxOverlapDifference']}"
	)
	return payload, stats


def main() -> None:
	args = parse_args()
	config = validate_config(common.read_json_path(args.config))
	provider_catalog = common.validate_provider_catalog(common.read_json_path(args.providers))
	provider = config["provider"]
	if not any(item["id"] == provider["id"] for item in provider_catalog["providers"]):
		raise RuntimeError("EDGAR is missing from statistics-providers.json.")
	area_by_iso3, registry_payload = common.load_registry(args.registry, args.timeout)
	now = common.utc_now()
	path, download_metadata = download_workbook(str(config["downloadUrl"]), args.timeout)
	try:
		direct_values, direct_stats, workbook_stats = parse_direct_values(
			path, config, area_by_iso3, now.year
		)
		payloads: list[tuple[dict[str, Any], dict[str, Any], dict[str, Any]]] = []
		for indicator in config["indicators"]:
			payload, stats = build_indicator_payload(
				config,
				indicator,
				direct_values,
				direct_stats,
				area_by_iso3,
				now.year,
				args.timeout,
			)
			payloads.append((indicator, payload, stats))
	finally:
		path.unlink(missing_ok=True)

	hasher = hashlib.sha256()
	for indicator, payload, _ in sorted(payloads, key=lambda item: item[0]["id"]):
		hasher.update(str(indicator["id"]).encode("utf-8") + b"\0" + common.canonical_bytes(payload) + b"\0")
	snapshot = hasher.hexdigest()[:16]
	provider_dir = args.output_dir / provider["id"]
	release_dir = provider_dir / "releases" / snapshot
	index_indicators: list[dict[str, Any]] = []
	for indicator, payload, _ in payloads:
		filename = f"{indicator['slug']}.json"
		common.write_json(release_dir / filename, payload)
		index_indicators.append({
			"id": indicator["id"],
			"title": indicator["title"],
			"description": indicator["description"],
			"areaLevel": "country",
			"frequency": "annual",
			"unit": indicator["unit"],
			"classification": indicator["classification"],
			"sourceIndicator": indicator["sourceIndicator"],
			"path": f"releases/{snapshot}/{filename}",
			"availableYears": payload["availableYears"],
			"defaultYear": payload["defaultYear"],
			"coverage": payload["coverage"],
		})

	registry_source = {
		"url": args.registry if args.registry.startswith(("https://", "http://")) else None,
		"schema": registry_payload.get("schema"),
		"generatedAt": registry_payload.get("generatedAt"),
		"areaCount": len(area_by_iso3),
	}
	if registry_source["url"] is None:
		registry_source.pop("url")
	provider_index = {
		"schema": "kartensammlung.statistics-provider-index/v1",
		"provider": provider,
		"retrievedAt": now.isoformat().replace("+00:00", "Z"),
		"activeSnapshot": snapshot,
		"areaRegistry": registry_source,
		"dataset": {
			"sourceUrl": config["sourceUrl"],
			"datasetUrl": config["datasetUrl"],
			"downloadUrl": config["downloadUrl"],
			"sheet": config["sheet"],
			"bytes": download_metadata["bytes"],
			"contentType": download_metadata["contentType"],
			"lastModified": download_metadata["lastModified"],
			"worksheetRows": workbook_stats["worksheetRows"],
			"worksheetColumns": workbook_stats["worksheetColumns"],
			"sourceRows": workbook_stats["sourceRows"],
			"ignoredCodes": workbook_stats["ignoredCodes"],
			"ignoredCodeRows": workbook_stats["ignoredCodeRows"],
			"excludedCombinedAreas": workbook_stats["excludedCombinedAreas"],
			"excludedRows": workbook_stats["excludedRows"],
			"workbookMetadata": workbook_stats["workbookMetadata"],
		},
		"indicators": index_indicators,
		"notes": [
			"EDGAR 2026 CH4 and N2O AR5 CO2-equivalent values are canonical for safe one-to-one country mappings.",
			"Six EDGAR report rows that combine multiple registry countries are intentionally excluded from direct mapping.",
			"World Bank WDI is retained only as an explicit fallback through 2024 for country-year observations without a safe direct EDGAR value.",
			"Fossil CO2, total GHG and any indicator containing the restricted IEA-EDGAR CO2 component are intentionally not redistributed by this provider.",
		],
	}
	common.write_json(provider_dir / "index.json", provider_index)
	common.write_json(args.output_dir / "index.json", common.build_global_index(provider_catalog))
	print(f"Built EDGAR snapshot {snapshot}")
	print(f"Provider index: {provider_dir / 'index.json'}")
	print(f"Indicators: {len(index_indicators)}")


if __name__ == "__main__":
	main()
