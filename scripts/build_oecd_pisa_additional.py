#!/usr/bin/env python3
from __future__ import annotations

import hashlib
import io
import math
import re
import time
from collections import defaultdict
from dataclasses import dataclass
from typing import Any
from urllib.error import HTTPError, URLError
from urllib.request import Request, urlopen

import pycountry
import xlrd
from openpyxl import load_workbook

import build_world_bank as common

USER_AGENT = "kartensammlung-area-statistics-builds/1"
MISSING_MARKER = "m"


@dataclass(frozen=True)
class SourceFile:
	raw: bytes
	headers: dict[str, str]
	effective_url: str


def validate_config(payload: Any) -> dict[str, Any]:
	if not isinstance(payload, dict) or payload.get("schema") != "kartensammlung.oecd-pisa-additional/v1":
		raise ValueError("Invalid OECD PISA additional statistics config.")

	sources = payload.get("sources")
	if not isinstance(sources, dict) or not sources:
		raise ValueError("OECD PISA additional config requires sources.")
	for source_id, source in sources.items():
		if not isinstance(source, dict):
			raise ValueError(f"Invalid OECD PISA source {source_id}.")
		for key in ("publication", "sourceUrl", "dataUrl", "license", "licenseUrl", "attribution"):
			if not str(source.get(key, "")).strip():
				raise ValueError(f"OECD PISA source {source_id} is missing {key}.")
		for key in ("sourceUrl", "dataUrl", "licenseUrl"):
			if not str(source[key]).startswith("https://"):
				raise ValueError(f"OECD PISA source {source_id} {key} must use HTTPS.")

	aliases = payload.get("countryAliases")
	if not isinstance(aliases, dict):
		raise ValueError("OECD PISA additional countryAliases must be an object.")
	for source_name, iso3 in aliases.items():
		if not str(source_name).strip() or not re.fullmatch(r"[A-Z0-9]{3}", str(iso3)):
			raise ValueError(f"Invalid OECD PISA additional alias: {source_name!r} -> {iso3!r}")

	indicators = payload.get("indicators")
	if not isinstance(indicators, list) or len(indicators) != 9:
		raise ValueError("OECD PISA additional config requires exactly nine indicators.")
	ids: set[str] = set()
	slugs: set[str] = set()
	for indicator in indicators:
		if not isinstance(indicator, dict):
			raise ValueError("OECD PISA additional indicator must be an object.")
		for key in ("id", "slug", "title", "description"):
			if not str(indicator.get(key, "")).strip():
				raise ValueError(f"OECD PISA additional indicator is missing {key}.")
		indicator_id = str(indicator["id"])
		slug = str(indicator["slug"])
		if indicator_id in ids or slug in slugs:
			raise ValueError(f"Duplicate OECD PISA additional id/slug: {indicator_id} / {slug}")
		ids.add(indicator_id)
		slugs.add(slug)
		unit = indicator.get("unit")
		if not isinstance(unit, dict) or not str(unit.get("id", "")).strip():
			raise ValueError(f"OECD PISA additional unit missing for {indicator_id}.")
		common.validate_classification(indicator_id, indicator.get("classification"))
		value_range = indicator.get("valueRange")
		if (
			not isinstance(value_range, list)
			or len(value_range) != 2
			or any(isinstance(value, bool) or not isinstance(value, (int, float)) for value in value_range)
			or not all(math.isfinite(float(value)) for value in value_range)
			or float(value_range[0]) >= float(value_range[1])
		):
			raise ValueError(f"Invalid OECD PISA additional valueRange for {indicator_id}.")
		expected_years = indicator.get("expectedYears")
		if not isinstance(expected_years, list) or not expected_years or any(not isinstance(year, int) for year in expected_years):
			raise ValueError(f"Invalid OECD PISA additional expectedYears for {indicator_id}.")
		if sorted(set(expected_years)) != expected_years:
			raise ValueError(f"OECD PISA additional expectedYears must be sorted/unique for {indicator_id}.")
		expected_coverage = indicator.get("expectedAreasByYear")
		if not isinstance(expected_coverage, dict) or set(expected_coverage) != {str(year) for year in expected_years}:
			raise ValueError(f"Invalid OECD PISA additional expectedAreasByYear for {indicator_id}.")
		if any(not isinstance(value, int) or value <= 0 for value in expected_coverage.values()):
			raise ValueError(f"Invalid OECD PISA additional coverage count for {indicator_id}.")
		expected_any = indicator.get("expectedAreasWithAnyValue")
		if not isinstance(expected_any, int) or expected_any <= 0:
			raise ValueError(f"Invalid OECD PISA additional expectedAreasWithAnyValue for {indicator_id}.")
		inputs = indicator.get("inputs")
		if not isinstance(inputs, list) or not inputs:
			raise ValueError(f"OECD PISA additional inputs missing for {indicator_id}.")
		for input_spec in inputs:
			if not isinstance(input_spec, dict):
				raise ValueError(f"Invalid input for {indicator_id}.")
			if input_spec.get("source") not in sources:
				raise ValueError(f"Unknown source for {indicator_id}: {input_spec.get('source')}")
			if input_spec.get("parser") not in {"single-mean", "financial-trend"}:
				raise ValueError(f"Unsupported parser for {indicator_id}: {input_spec.get('parser')}")
			if not str(input_spec.get("sheet", "")).strip():
				raise ValueError(f"Missing sheet for {indicator_id}.")
			if input_spec["parser"] == "single-mean" and not isinstance(input_spec.get("year"), int):
				raise ValueError(f"single-mean input requires year for {indicator_id}.")
			if input_spec["parser"] == "financial-trend":
				years = input_spec.get("years")
				if years != [2012, 2015, 2018, 2022]:
					raise ValueError(f"Unexpected financial-literacy years for {indicator_id}: {years}")
	return payload


def fetch_source(url: str, timeout: int) -> SourceFile:
	last_error: Exception | None = None
	for attempt, delay in enumerate((0, 3, 10, 20, 40), start=1):
		if delay:
			time.sleep(delay)
		request = Request(url, headers={"Accept": "*/*", "User-Agent": USER_AGENT})
		try:
			with urlopen(request, timeout=timeout) as response:
				raw = response.read()
				headers = {key.lower(): value for key, value in response.headers.items()}
				effective_url = response.geturl()
			if not raw:
				raise RuntimeError("OECD PISA source returned an empty response.")
			if not (raw.startswith(b"PK") or raw.startswith(b"\xd0\xcf\x11\xe0")):
				raise RuntimeError("OECD PISA source did not return XLSX/XLS data.")
			return SourceFile(raw=raw, headers=headers, effective_url=effective_url)
		except (HTTPError, URLError, TimeoutError, OSError, RuntimeError) as error:
			last_error = error
			print(f"OECD PISA additional request failed ({attempt}/5): {url}: {error}")
	if last_error is None:
		raise RuntimeError("OECD PISA additional request failed without an exception.")
	raise RuntimeError(f"OECD PISA additional request failed after 5 attempts: {url}") from last_error


def workbook_rows(raw: bytes) -> dict[str, list[tuple[Any, ...]]]:
	if raw.startswith(b"PK"):
		workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
		try:
			return {
				sheet_name: [tuple(row) for row in workbook[sheet_name].iter_rows(values_only=True)]
				for sheet_name in workbook.sheetnames
			}
		finally:
			workbook.close()
	if raw.startswith(b"\xd0\xcf\x11\xe0"):
		workbook = xlrd.open_workbook(file_contents=raw, on_demand=True)
		try:
			result: dict[str, list[tuple[Any, ...]]] = {}
			for sheet_name in workbook.sheet_names():
				sheet = workbook.sheet_by_name(sheet_name)
				result[sheet_name] = [tuple(sheet.row_values(row_index)) for row_index in range(sheet.nrows)]
			return result
		finally:
			workbook.release_resources()
	raise RuntimeError("Unsupported OECD PISA workbook format.")


def compact_row(row: tuple[Any, ...]) -> list[Any]:
	return [value for value in row if value is not None and value != ""]


def clean_name(raw: str) -> tuple[str, int]:
	name = " ".join(raw.strip().split())
	star_count = len(name) - len(name.rstrip("*"))
	name = name.rstrip("*").strip()
	name = re.sub(r"^Cyprus1,\s*2$", "Cyprus", name)
	return name, star_count


def is_aggregate(name: str) -> bool:
	normalized = name.strip().lower()
	return (
		normalized.startswith("oecd average")
		or normalized.startswith("oecd total")
		or normalized.startswith("overall average")
	)


def first_number(cells: list[Any]) -> float | None:
	for value in cells[1:]:
		if isinstance(value, bool):
			continue
		if isinstance(value, (int, float)):
			return float(value)
	return None


def parse_single_mean(
	rows: list[tuple[Any, ...]],
	year: int,
	value_range: list[int | float],
	section_start: str | None = None,
	section_end: str | None = None,
) -> tuple[dict[str, dict[int, float]], dict[int, set[str]], dict[int, set[str]]]:
	values: dict[str, dict[int, float]] = {}
	sampling_cautions: dict[int, set[str]] = defaultdict(set)
	scale_cautions: dict[int, set[str]] = defaultdict(set)
	minimum, maximum = (float(value) for value in value_range)
	active = section_start is None
	start_norm = section_start.lower() if section_start else None
	end_norm = section_end.lower() if section_end else None
	for row in rows:
		cells = compact_row(row)
		if not cells:
			continue
		first_text = str(cells[0]).strip() if isinstance(cells[0], str) else ""
		normalized = " ".join(first_text.lower().split())
		if not active:
			if normalized == start_norm:
				active = True
			continue
		if end_norm is not None and normalized == end_norm:
			break
		if not isinstance(cells[0], str):
			continue
		name, star_count = clean_name(cells[0])
		if not name or is_aggregate(name):
			continue
		value = first_number(cells)
		if value is None:
			continue
		if not math.isfinite(value) or value < minimum or value > maximum:
			continue
		if name in values:
			raise RuntimeError(f"Duplicate OECD PISA additional row {name!r} in year {year}.")
		values[name] = {year: value}
		if star_count == 1:
			sampling_cautions[year].add(name)
		elif star_count >= 2:
			scale_cautions[year].add(name)
	return values, sampling_cautions, scale_cautions


def parse_financial_trend(
	rows: list[tuple[Any, ...]],
	years: list[int],
	value_range: list[int | float],
	caution_year: int,
) -> tuple[dict[str, dict[int, float]], dict[int, set[str]], dict[int, set[str]]]:
	values: dict[str, dict[int, float]] = {}
	sampling_cautions: dict[int, set[str]] = defaultdict(set)
	scale_cautions: dict[int, set[str]] = defaultdict(set)
	minimum, maximum = (float(value) for value in value_range)
	mean_positions = [1, 3, 5, 7]
	for row in rows:
		cells = compact_row(row)
		if not cells or not isinstance(cells[0], str):
			continue
		name, star_count = clean_name(cells[0])
		if not name or is_aggregate(name):
			continue
		row_values: dict[int, float] = {}
		for year, position in zip(years, mean_positions):
			if position >= len(cells):
				continue
			value = cells[position]
			if isinstance(value, str):
				if value.strip() in {"", MISSING_MARKER}:
					continue
				continue
			if isinstance(value, bool) or not isinstance(value, (int, float)):
				continue
			number = float(value)
			if not math.isfinite(number) or number < minimum or number > maximum:
				continue
			row_values[year] = number
		if not row_values:
			continue
		if name in values:
			raise RuntimeError(f"Duplicate OECD PISA financial-literacy row {name!r}.")
		values[name] = row_values
		if star_count == 1:
			sampling_cautions[caution_year].add(name)
		elif star_count >= 2:
			scale_cautions[caution_year].add(name)
	return values, sampling_cautions, scale_cautions


def merge_source_rows(
	target: dict[str, dict[int, float]],
	incoming: dict[str, dict[int, float]],
	indicator_id: str,
) -> None:
	for name, source_values in incoming.items():
		destination = target.setdefault(name, {})
		for year, value in source_values.items():
			if year in destination:
				raise RuntimeError(f"Duplicate OECD PISA additional value for {indicator_id}: {name} {year}.")
			destination[year] = value


def merge_cautions(target: dict[int, set[str]], incoming: dict[int, set[str]]) -> None:
	for year, names in incoming.items():
		target[year].update(names)


def iso3_for_source_name(name: str, aliases: dict[str, str]) -> str | None:
	if name in aliases:
		return aliases[name]
	try:
		return pycountry.countries.lookup(name).alpha_3
	except LookupError:
		return None


def build_mapping(
	names: set[str],
	excluded: set[str],
	aliases: dict[str, str],
	area_by_iso3: dict[str, str],
	indicator_id: str,
) -> dict[str, tuple[str, str]]:
	missing_exclusions = sorted(excluded - names)
	if missing_exclusions:
		raise RuntimeError(f"Configured OECD PISA exclusions disappeared for {indicator_id}: {missing_exclusions}")
	mapping: dict[str, tuple[str, str]] = {}
	unresolved: list[str] = []
	missing_registry: list[tuple[str, str]] = []
	used_iso3: dict[str, str] = {}
	for name in sorted(names):
		if name in excluded:
			continue
		iso3 = iso3_for_source_name(name, aliases)
		if iso3 is None:
			unresolved.append(name)
			continue
		area_id = area_by_iso3.get(iso3)
		if area_id is None:
			missing_registry.append((name, iso3))
			continue
		previous = used_iso3.get(iso3)
		if previous is not None and previous != name:
			raise RuntimeError(f"Multiple OECD PISA names map to {iso3} for {indicator_id}: {previous!r}, {name!r}")
		used_iso3[iso3] = name
		mapping[name] = (iso3, area_id)
	if unresolved:
		raise RuntimeError(f"Unresolved OECD PISA country/economy names for {indicator_id}: {unresolved}")
	if missing_registry:
		raise RuntimeError(f"OECD PISA mappings missing from area registry for {indicator_id}: {missing_registry}")
	return mapping


def mapped_cautions(
	cautions: dict[int, set[str]],
	mapping: dict[str, tuple[str, str]],
	excluded: set[str],
) -> dict[str, list[str]]:
	result: dict[str, list[str]] = {}
	for year in sorted(cautions):
		areas = sorted(mapping[name][1] for name in cautions[year] if name not in excluded and name in mapping)
		if areas:
			result[str(year)] = areas
	return result


def source_summary(source_id: str, source: dict[str, Any], source_file: SourceFile) -> dict[str, Any]:
	result = {
		"id": source_id,
		"publication": source["publication"],
		"publicationDate": source.get("publicationDate"),
		"url": source["sourceUrl"],
		"dataUrl": source["dataUrl"],
		"effectiveDataUrl": source_file.effective_url,
		"license": source["license"],
		"licenseUrl": source["licenseUrl"],
		"attribution": source["attribution"],
		"sourceFileSha256": hashlib.sha256(source_file.raw).hexdigest(),
		"sourceFileBytes": len(source_file.raw),
	}
	if result["publicationDate"] is None:
		result.pop("publicationDate")
	if source_file.headers.get("last-modified"):
		result["lastModified"] = source_file.headers["last-modified"]
	if source_file.headers.get("etag"):
		result["etag"] = source_file.headers["etag"]
	return result


def build_additional_payloads(
	config: dict[str, Any],
	provider: dict[str, Any],
	area_by_iso3: dict[str, str],
	timeout: int,
	cached_sources: dict[str, SourceFile] | None = None,
) -> tuple[list[tuple[dict[str, Any], dict[str, Any]]], list[dict[str, Any]]]:
	config = validate_config(config)
	aliases = {str(key): str(value) for key, value in config["countryAliases"].items()}
	source_cache: dict[str, SourceFile] = dict(cached_sources or {})
	workbook_cache: dict[str, dict[str, list[tuple[Any, ...]]]] = {}
	source_summaries: dict[str, dict[str, Any]] = {}

	def get_source(source_id: str) -> tuple[dict[str, Any], SourceFile, dict[str, list[tuple[Any, ...]]]]:
		source = config["sources"][source_id]
		data_url = str(source["dataUrl"])
		if data_url not in source_cache:
			print(f"Fetching OECD PISA additional source {source_id} from {data_url}")
			source_cache[data_url] = fetch_source(data_url, timeout)
		source_file = source_cache[data_url]
		if data_url not in workbook_cache:
			workbook_cache[data_url] = workbook_rows(source_file.raw)
		source_summaries[source_id] = source_summary(source_id, source, source_file)
		return source, source_file, workbook_cache[data_url]

	payloads: list[tuple[dict[str, Any], dict[str, Any]]] = []
	for indicator in config["indicators"]:
		indicator_id = str(indicator["id"])
		source_rows: dict[str, dict[int, float]] = {}
		sampling_cautions: dict[int, set[str]] = defaultdict(set)
		scale_cautions: dict[int, set[str]] = defaultdict(set)
		input_source_ids: list[str] = []
		input_tables: list[str] = []
		input_source_observations: dict[str, int] = {}

		for input_index, input_spec in enumerate(indicator["inputs"]):
			source_id = str(input_spec["source"])
			source, _, sheets = get_source(source_id)
			sheet_name = str(input_spec["sheet"])
			if sheet_name not in sheets:
				raise RuntimeError(f"OECD PISA source {source_id} is missing sheet {sheet_name} for {indicator_id}.")
			parser = str(input_spec["parser"])
			if parser == "single-mean":
				parsed, sample_warnings, scale_warnings = parse_single_mean(
					sheets[sheet_name],
					int(input_spec["year"]),
					list(indicator["valueRange"]),
					str(input_spec["sectionStart"]) if input_spec.get("sectionStart") else None,
					str(input_spec["sectionEnd"]) if input_spec.get("sectionEnd") else None,
				)
			elif parser == "financial-trend":
				parsed, sample_warnings, scale_warnings = parse_financial_trend(
					sheets[sheet_name],
					list(input_spec["years"]),
					list(indicator["valueRange"]),
					int(input_spec.get("cautionYear", max(input_spec["years"]))),
				)
			else:
				raise RuntimeError(f"Unsupported OECD PISA additional parser: {parser}")
			merge_source_rows(source_rows, parsed, indicator_id)
			merge_cautions(sampling_cautions, sample_warnings)
			merge_cautions(scale_cautions, scale_warnings)
			input_source_ids.append(source_id)
			input_tables.append(sheet_name)
			input_source_observations[f"{input_index}:{source_id}:{sheet_name}"] = sum(len(item) for item in parsed.values())
			print(
				f"{indicator_id}: source={source_id} table={sheet_name} "
				f"entities={len(parsed)} observations={sum(len(item) for item in parsed.values())}"
			)

		excluded = {str(item) for item in indicator.get("excludedEntities", [])}
		mapping = build_mapping(set(source_rows), excluded, aliases, area_by_iso3, indicator_id)
		values: dict[int, dict[str, int | float]] = defaultdict(dict)
		source_observations = 0
		retained_observations = 0
		for name, yearly in source_rows.items():
			for year, value in yearly.items():
				source_observations += 1
				if name in excluded:
					continue
				if name not in mapping:
					raise RuntimeError(f"Missing OECD PISA additional mapping for {indicator_id}: {name}")
				area_id = mapping[name][1]
				if area_id in values[year]:
					raise RuntimeError(f"Duplicate OECD PISA additional area/year for {indicator_id}: {area_id} {year}")
				values[year][area_id] = common.normalize_number(f"{value:.5f}")
				retained_observations += 1

		available_years = sorted(year for year, mapped in values.items() if mapped)
		if available_years != indicator["expectedYears"]:
			raise RuntimeError(f"Unexpected OECD PISA additional years for {indicator_id}: {available_years}")
		for year in available_years:
			expected = int(indicator["expectedAreasByYear"][str(year)])
			actual = len(values[year])
			if actual != expected:
				raise RuntimeError(f"Unexpected OECD PISA additional coverage for {indicator_id} {year}: {actual}, expected {expected}")
		mapped_areas = {area_id for mapped in values.values() for area_id in mapped}
		if len(mapped_areas) != int(indicator["expectedAreasWithAnyValue"]):
			raise RuntimeError(
				f"Unexpected OECD PISA additional area union for {indicator_id}: {len(mapped_areas)}, "
				f"expected {indicator['expectedAreasWithAnyValue']}"
			)
		default_year = available_years[-1]
		for area_id in indicator.get("requiredAreas", []):
			if area_id not in mapped_areas:
				raise RuntimeError(f"OECD PISA additional {indicator_id} is missing required area {area_id}.")

		unique_source_ids = list(dict.fromkeys(input_source_ids))
		source_details = [source_summaries[source_id] for source_id in unique_source_ids]
		licenses = {(item["license"], item["licenseUrl"]) for item in source_details}
		if len(licenses) != 1:
			raise RuntimeError(f"Mixed source licenses within OECD PISA indicator {indicator_id} are not supported.")
		license_name, license_url = next(iter(licenses))
		primary_source = source_details[-1]
		source_metadata: dict[str, Any] = {
			"providerId": provider["id"],
			"providerName": provider["name"],
			"dataset": "OECD PISA official aggregated result tables",
			"publication": primary_source["publication"],
			"license": license_name,
			"licenseUrl": license_url,
			"attribution": primary_source["attribution"],
			"url": primary_source["url"],
			"dataUrl": primary_source["dataUrl"],
			"indicator": ", ".join(input_tables),
			"sourceTables": input_tables,
			"sources": source_details,
			"sourceObservations": source_observations,
			"retainedObservations": retained_observations,
			"sourceObservationsByInput": input_source_observations,
			"excludedPartialEntities": sorted(excluded),
		}
		if primary_source.get("publicationDate"):
			source_metadata["publicationDate"] = primary_source["publicationDate"]
		mapped_sampling = mapped_cautions(sampling_cautions, mapping, excluded)
		mapped_scale = mapped_cautions(scale_cautions, mapping, excluded)
		if mapped_sampling:
			source_metadata["samplingCautionAreasByYear"] = mapped_sampling
		if mapped_scale:
			source_metadata["scaleLinkageCautionAreasByYear"] = mapped_scale

		payload = {
			"schema": "kartensammlung.statistics-indicator/v1",
			"indicator": {
				"id": indicator_id,
				"title": indicator["title"],
				"description": indicator["description"],
				"areaLevel": "country",
				"frequency": "pisa-cycle",
				"unit": indicator["unit"],
				"classification": indicator["classification"],
			},
			"source": source_metadata,
			"availableYears": available_years,
			"defaultYear": default_year,
			"coverage": {
				"registryAreas": len(area_by_iso3),
				"areasWithAnyValue": len(mapped_areas),
				"latestYear": default_year,
				"areasInLatestYear": len(values[default_year]),
				"areasInDefaultYear": len(values[default_year]),
				"observations": retained_observations,
			},
			"values": {
				str(year): {area_id: values[year][area_id] for area_id in sorted(values[year])}
				for year in available_years
			},
		}
		payloads.append((indicator, payload))
		print(
			f"{indicator_id}: mappedAreas={len(mapped_areas)} years={available_years} "
			f"coverage={[len(values[year]) for year in available_years]} "
			f"observations={retained_observations}/{source_observations}"
		)

	return payloads, [source_summaries[source_id] for source_id in sorted(source_summaries)]
