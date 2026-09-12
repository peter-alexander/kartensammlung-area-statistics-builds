#!/usr/bin/env python3
from __future__ import annotations

import argparse
from datetime import datetime, timezone
import hashlib
import io
import json
import math
from pathlib import Path
import time
from typing import Any
from urllib.error import HTTPError, URLError
from urllib.request import Request, urlopen

from openpyxl import load_workbook

import build_world_bank as common

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_CONFIG = ROOT / "config" / "who-ghed-indicators.json"
DEFAULT_PROVIDERS = ROOT / "config" / "statistics-providers.json"
DEFAULT_OUTPUT_DIR = ROOT / "dist" / "statistics"
DEFAULT_REGISTRY = "https://tiles.radlobby.at/AreaStatistics/area-registry-countries.json"
USER_AGENT = "kartensammlung-area-statistics-builds/1"


def parse_args() -> argparse.Namespace:
	parser = argparse.ArgumentParser(description="Build normalized country statistics from WHO GHED.")
	parser.add_argument("--config", type=Path, default=DEFAULT_CONFIG)
	parser.add_argument("--providers", type=Path, default=DEFAULT_PROVIDERS)
	parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
	parser.add_argument("--registry", default=DEFAULT_REGISTRY)
	parser.add_argument("--timeout", type=int, default=120)
	return parser.parse_args()


def validate_config(payload: Any) -> dict[str, Any]:
	if not isinstance(payload, dict) or payload.get("schema") != "kartensammlung.who-ghed-statistics/v1":
		raise ValueError("Invalid WHO GHED statistics config.")
	for key in ("sourceUrl", "documentationUrl", "downloadUrl"):
		if not str(payload.get(key, "")).startswith("https://"):
			raise ValueError(f"WHO GHED {key} must use HTTPS.")
	provider = payload.get("provider")
	if not isinstance(provider, dict) or provider.get("id") != "who-ghed":
		raise ValueError("WHO GHED provider metadata is invalid.")
	for key in ("name", "dataset", "license", "licenseUrl", "attribution"):
		if not str(provider.get(key, "")).strip():
			raise ValueError(f"WHO GHED provider metadata is missing {key}.")
	preliminary_years = payload.get("preliminaryYears", [])
	if not isinstance(preliminary_years, list) or any(not isinstance(year, int) or not 1900 <= year <= 2200 for year in preliminary_years):
		raise ValueError("WHO GHED preliminaryYears is invalid.")
	indicators = payload.get("indicators")
	if not isinstance(indicators, list) or not indicators:
		raise ValueError("WHO GHED config requires indicators.")
	ids: set[str] = set()
	slugs: set[str] = set()
	columns: set[str] = set()
	for indicator in indicators:
		if not isinstance(indicator, dict):
			raise ValueError("WHO GHED indicator entries must be objects.")
		for key in ("id", "slug", "sourceColumn", "title", "description"):
			if not str(indicator.get(key, "")).strip():
				raise ValueError(f"WHO GHED indicator entry is missing {key}.")
		indicator_id = str(indicator["id"])
		slug = str(indicator["slug"])
		column = str(indicator["sourceColumn"])
		if indicator_id in ids or slug in slugs or column in columns:
			raise ValueError(f"Duplicate WHO GHED indicator metadata: {indicator_id}.")
		ids.add(indicator_id)
		slugs.add(slug)
		columns.add(column)
		unit = indicator.get("unit")
		if not isinstance(unit, dict) or not str(unit.get("id", "")).strip() or not str(unit.get("label", "")).strip():
			raise ValueError(f"WHO GHED indicator {indicator_id} has invalid unit metadata.")
		common.validate_classification(indicator_id, indicator.get("classification"))
		for key in ("minAreasWithAnyValue", "minAreasInDefaultYear"):
			if not isinstance(indicator.get(key), int) or int(indicator[key]) <= 0:
				raise ValueError(f"WHO GHED indicator {indicator_id} has invalid {key}.")
		required = indicator.get("requiredAreas")
		if not isinstance(required, list) or not required:
			raise ValueError(f"WHO GHED indicator {indicator_id} requires requiredAreas.")
		value_range = indicator.get("valueRange")
		if (
			not isinstance(value_range, list)
			or len(value_range) != 2
			or any(isinstance(value, bool) or not isinstance(value, (int, float)) for value in value_range)
			or not math.isfinite(float(value_range[0]))
			or not math.isfinite(float(value_range[1]))
			or float(value_range[0]) >= float(value_range[1])
		):
			raise ValueError(f"WHO GHED indicator {indicator_id} has invalid valueRange.")
	return payload


def fetch_bytes(url: str, timeout: int) -> bytes:
	delays = (0, 5, 15, 30)
	last_error: Exception | None = None
	for attempt, delay in enumerate(delays, start=1):
		if delay:
			time.sleep(delay)
		request = Request(
			url,
			headers={
				"Accept": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet,*/*",
				"User-Agent": USER_AGENT,
			},
		)
		try:
			with urlopen(request, timeout=timeout) as response:
				data = response.read()
			if len(data) < 1_000_000 or not data.startswith(b"PK"):
				raise RuntimeError(f"WHO GHED download is not a plausible XLSX file: {len(data)} bytes.")
			return data
		except (HTTPError, URLError, TimeoutError, RuntimeError) as error:
			last_error = error
			print(f"WHO GHED request failed ({attempt}/{len(delays)}): {error}")
	if last_error is None:
		raise RuntimeError("WHO GHED request failed without exception.")
	raise RuntimeError("WHO GHED request failed after retries.") from last_error


def normalize_number(value: Any) -> int | float:
	if isinstance(value, bool) or value is None:
		raise ValueError("Missing or invalid numeric value.")
	number = float(value)
	if not math.isfinite(number):
		raise ValueError("Numeric value is not finite.")
	if number.is_integer():
		return int(number)
	return number


def workbook_version(workbook: Any) -> list[str]:
	if "Version" not in workbook.sheetnames:
		raise RuntimeError("WHO GHED workbook is missing Version sheet.")
	lines = [str(row[0]).strip() for row in workbook["Version"].iter_rows(values_only=True) if row and row[0]]
	if not lines or lines[0] != "WHO Global Health Expenditure Database (GHED)":
		raise RuntimeError(f"Unexpected WHO GHED Version sheet: {lines}")
	if not any(line.startswith("Last updated:") for line in lines):
		raise RuntimeError(f"WHO GHED Version sheet has no update date: {lines}")
	return lines


def load_codebook(workbook: Any) -> dict[str, dict[str, Any]]:
	if "Codebook" not in workbook.sheetnames:
		raise RuntimeError("WHO GHED workbook is missing Codebook sheet.")
	rows = workbook["Codebook"].iter_rows(values_only=True)
	header = [str(value).strip() if value is not None else "" for value in next(rows)]
	required = {"variable code", "variable name", "long code (GHED data explorer)", "unit", "currency"}
	if not required.issubset(set(header)):
		raise RuntimeError(f"WHO GHED Codebook schema changed: {header}")
	index = {name: i for i, name in enumerate(header)}
	result: dict[str, dict[str, Any]] = {}
	for row in rows:
		code = str(row[index["variable code"]] or "").strip()
		if not code:
			continue
		if code in result:
			raise RuntimeError(f"Duplicate WHO GHED Codebook variable: {code}")
		result[code] = {
			"variableName": str(row[index["variable name"]] or "").strip(),
			"longCode": str(row[index["long code (GHED data explorer)"]] or "").strip(),
			"unit": str(row[index["unit"]] or "").strip(),
			"currency": str(row[index["currency"]] or "").strip(),
		}
	return result


def choose_default_year(indicator: dict[str, Any], values_by_year: dict[int, dict[str, int | float]]) -> int:
	minimum = int(indicator["minAreasInDefaultYear"])
	eligible = [year for year, values in values_by_year.items() if len(values) >= minimum]
	if not eligible:
		raise RuntimeError(
			f"WHO GHED has no sufficiently complete default year for {indicator['id']}; "
			f"expected at least {minimum} mapped areas."
		)
	return max(eligible)


def build_payload(
	config: dict[str, Any],
	indicator: dict[str, Any],
	values_by_year: dict[int, dict[str, int | float]],
	area_count: int,
	codebook: dict[str, dict[str, Any]],
	version_lines: list[str],
) -> dict[str, Any]:
	available_years = sorted(year for year, values in values_by_year.items() if values)
	if not available_years:
		raise RuntimeError(f"WHO GHED returned no values for {indicator['id']}.")
	areas_with_any_value = set().union(*(set(values_by_year[year]) for year in available_years))
	if len(areas_with_any_value) < int(indicator["minAreasWithAnyValue"]):
		raise RuntimeError(
			f"WHO GHED indicator {indicator['id']} maps only {len(areas_with_any_value)} areas; "
			f"expected at least {indicator['minAreasWithAnyValue']}."
		)
	default_year = choose_default_year(indicator, values_by_year)
	for area_id in indicator["requiredAreas"]:
		if area_id not in values_by_year[default_year]:
			raise RuntimeError(
				f"Required area {area_id} has no WHO GHED value for {indicator['id']} in {default_year}."
			)

	column = str(indicator["sourceColumn"])
	metadata = codebook[column]
	provider = config["provider"]
	preliminary_years = set(int(year) for year in config.get("preliminaryYears", []))
	observation_metadata = {
		str(year): {
			area_id: {"preliminary": True}
			for area_id in sorted(values_by_year[year])
		}
		for year in available_years
		if year in preliminary_years
	}
	payload: dict[str, Any] = {
		"schema": "kartensammlung.statistics-indicator/v1",
		"indicator": {
			"id": indicator["id"],
			"title": indicator["title"],
			"description": indicator["description"],
			"areaLevel": "country",
			"frequency": "annual",
			"unit": indicator["unit"],
			"classification": indicator["classification"],
		},
		"source": {
			"providerId": provider["id"],
			"providerName": provider["name"],
			"dataset": provider["dataset"],
			"indicator": column,
			"indicatorName": metadata["variableName"],
			"indicatorLongCode": metadata["longCode"],
			"sourceUnit": metadata["unit"],
			"sourceCurrency": metadata["currency"],
			"license": provider["license"],
			"licenseUrl": provider["licenseUrl"],
			"attribution": provider["attribution"],
			"url": config["sourceUrl"],
			"documentationUrl": config["documentationUrl"],
			"downloadUrl": config["downloadUrl"],
			"workbookVersion": version_lines,
			"timeCoverage": {"startYear": available_years[0], "endYear": available_years[-1]},
			"preliminaryYears": sorted(year for year in preliminary_years if year in available_years),
		},
		"availableYears": available_years,
		"defaultYear": default_year,
		"coverage": {
			"registryAreas": area_count,
			"areasWithAnyValue": len(areas_with_any_value),
			"latestYear": available_years[-1],
			"areasInLatestYear": len(values_by_year[available_years[-1]]),
			"areasInDefaultYear": len(values_by_year[default_year]),
			"observations": sum(len(values_by_year[year]) for year in available_years),
		},
		"values": {
			str(year): {area_id: values_by_year[year][area_id] for area_id in sorted(values_by_year[year])}
			for year in available_years
		},
	}
	if observation_metadata:
		payload["observationMetadata"] = observation_metadata
	return payload


def main() -> None:
	args = parse_args()
	config = validate_config(common.read_json_path(args.config))
	provider_catalog = common.validate_provider_catalog(common.read_json_path(args.providers))
	provider = config["provider"]
	if not any(item["id"] == provider["id"] for item in provider_catalog["providers"]):
		raise RuntimeError("WHO GHED is missing from statistics-providers.json.")
	area_by_iso3, registry_payload = common.load_registry(args.registry, args.timeout)
	data = fetch_bytes(str(config["downloadUrl"]), max(args.timeout, 180))
	print(f"WHO GHED downloaded={len(data)} bytes")
	workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
	for required_sheet in ("Data", "Codebook", "Version"):
		if required_sheet not in workbook.sheetnames:
			raise RuntimeError(f"WHO GHED workbook is missing {required_sheet} sheet.")
	version_lines = workbook_version(workbook)
	print("WHO GHED version: " + " | ".join(version_lines))
	codebook = load_codebook(workbook)
	for indicator in config["indicators"]:
		column = str(indicator["sourceColumn"])
		if column not in codebook:
			raise RuntimeError(f"WHO GHED Codebook is missing configured column {column}.")

	data_sheet = workbook["Data"]
	rows = data_sheet.iter_rows(values_only=True)
	header = [str(value).strip() if value is not None else "" for value in next(rows)]
	header_index = {name: i for i, name in enumerate(header) if name}
	for required in ("location", "code", "year"):
		if required not in header_index:
			raise RuntimeError(f"WHO GHED Data sheet is missing {required}.")
	for indicator in config["indicators"]:
		if indicator["sourceColumn"] not in header_index:
			raise RuntimeError(f"WHO GHED Data sheet is missing {indicator['sourceColumn']}.")

	values: dict[str, dict[int, dict[str, int | float]]] = {
		str(indicator["id"]): {} for indicator in config["indicators"]
	}
	unmapped_codes: set[str] = set()
	seen_country_years: set[tuple[str, int]] = set()
	for row in rows:
		iso3 = str(row[header_index["code"]] or "").strip().upper()
		year_raw = row[header_index["year"]]
		if not iso3 or year_raw is None:
			continue
		try:
			year = int(year_raw)
		except (TypeError, ValueError):
			continue
		if (iso3, year) in seen_country_years:
			raise RuntimeError(f"Duplicate WHO GHED country-year row: {iso3} {year}")
		seen_country_years.add((iso3, year))
		area_id = area_by_iso3.get(iso3)
		if not area_id:
			unmapped_codes.add(iso3)
			continue
		for indicator in config["indicators"]:
			raw = row[header_index[str(indicator["sourceColumn"])]]
			if raw is None or raw == "":
				continue
			value = normalize_number(raw)
			low, high = (float(item) for item in indicator["valueRange"])
			if not low <= float(value) <= high:
				raise RuntimeError(
					f"WHO GHED value outside configured range for {indicator['id']} {iso3} {year}: "
					f"{value} not in {indicator['valueRange']}"
				)
			year_values = values[str(indicator["id"])].setdefault(year, {})
			if area_id in year_values:
				raise RuntimeError(f"Duplicate WHO GHED value for {indicator['id']} {area_id} {year}.")
			year_values[area_id] = value

	if unmapped_codes:
		raise RuntimeError(f"Unexpected unmapped WHO GHED country codes: {sorted(unmapped_codes)}")

	indicator_payloads: list[tuple[dict[str, Any], dict[str, Any]]] = []
	for indicator in config["indicators"]:
		payload = build_payload(
			config,
			indicator,
			values[str(indicator["id"])],
			len(area_by_iso3),
			codebook,
			version_lines,
		)
		indicator_payloads.append((indicator, payload))
		coverage = payload["coverage"]
		print(
			f"{indicator['id']}: observations={coverage['observations']} "
			f"areas={coverage['areasWithAnyValue']} years={payload['availableYears'][0]}-{payload['availableYears'][-1]} "
			f"defaultYear={payload['defaultYear']} defaultCoverage={coverage['areasInDefaultYear']} "
			f"latestCoverage={coverage['areasInLatestYear']}"
		)

	hasher = hashlib.sha256()
	for indicator, payload in sorted(indicator_payloads, key=lambda item: item[0]["id"]):
		hasher.update(str(indicator["id"]).encode("utf-8") + b"\0" + common.canonical_bytes(payload) + b"\0")
	snapshot = hasher.hexdigest()[:16]
	provider_dir = args.output_dir / provider["id"]
	release_dir = provider_dir / "releases" / snapshot
	index_indicators: list[dict[str, Any]] = []
	for indicator, payload in indicator_payloads:
		filename = f"{indicator['slug']}.json"
		common.write_json(release_dir / filename, payload)
		index_indicators.append({
			"id": indicator["id"],
			"title": indicator["title"],
			"description": indicator["description"],
			"areaLevel": "country",
			"frequency": "annual",
			"unit": indicator["unit"],
			"classification": indicator["classification"],
			"sourceIndicator": indicator["sourceColumn"],
			"path": f"releases/{snapshot}/{filename}",
			"availableYears": payload["availableYears"],
			"defaultYear": payload["defaultYear"],
			"coverage": payload["coverage"],
		})

	registry_source: dict[str, Any] = {
		"schema": registry_payload.get("schema"),
		"generatedAt": registry_payload.get("generatedAt"),
		"areaCount": len(area_by_iso3),
	}
	if args.registry.startswith(("https://", "http://")):
		registry_source["url"] = args.registry
	provider_index = {
		"schema": "kartensammlung.statistics-provider-index/v1",
		"provider": provider,
		"retrievedAt": datetime.now(timezone.utc).isoformat().replace("+00:00", "Z"),
		"activeSnapshot": snapshot,
		"areaRegistry": registry_source,
		"indicators": index_indicators,
		"notes": [
			"WHO Global Health Expenditure Database (GHED) is the canonical source for these health-financing indicators.",
			"The workbook's preliminary 2024 observations are retained but do not become the default year unless broad country coverage is reached.",
			"WHO GHED is the original source of the corresponding World Bank WDI health-expenditure series.",
		],
	}
	common.write_json(provider_dir / "index.json", provider_index)
	common.write_json(args.output_dir / "index.json", common.build_global_index(provider_catalog))
	print(f"Built WHO GHED snapshot {snapshot}")
	print(f"Indicators: {len(index_indicators)}")


if __name__ == "__main__":
	main()
