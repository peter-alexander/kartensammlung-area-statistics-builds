#!/usr/bin/env python3
from __future__ import annotations

import argparse
import hashlib
import math
import time
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from typing import Any
from urllib.error import HTTPError, URLError
from urllib.request import Request, urlopen

import openpyxl

import build_world_bank as common

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_CONFIG = ROOT / "config" / "who-ghe-indicators.json"
DEFAULT_PROVIDERS = ROOT / "config" / "statistics-providers.json"
DEFAULT_OUTPUT_DIR = ROOT / "dist" / "statistics"
DEFAULT_REGISTRY = "https://tiles.radlobby.at/AreaStatistics/area-registry-countries.json"
USER_AGENT = "kartensammlung-area-statistics-builds/1"


def parse_args() -> argparse.Namespace:
	parser = argparse.ArgumentParser(
		description="Build normalized country mortality statistics from WHO Global Health Estimates."
	)
	parser.add_argument("--config", type=Path, default=DEFAULT_CONFIG)
	parser.add_argument("--providers", type=Path, default=DEFAULT_PROVIDERS)
	parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
	parser.add_argument("--registry", default=DEFAULT_REGISTRY)
	parser.add_argument("--timeout", type=int, default=90)
	return parser.parse_args()


def validate_config(payload: Any) -> dict[str, Any]:
	if not isinstance(payload, dict) or payload.get("schema") != "kartensammlung.who-ghe-statistics/v1":
		raise ValueError("Invalid WHO GHE statistics config.")

	for key in ("workbookUrl", "sourceUrl"):
		if not str(payload.get(key, "")).startswith("https://"):
			raise ValueError(f"WHO GHE {key} must use HTTPS.")

	provider = payload.get("provider")
	if not isinstance(provider, dict):
		raise ValueError("WHO GHE provider metadata is required.")
	for key in ("id", "name", "dataset", "license", "licenseUrl", "attribution"):
		if not str(provider.get(key, "")).strip():
			raise ValueError(f"WHO GHE provider metadata is missing {key}.")
	if provider.get("id") != "who-ghe":
		raise ValueError("WHO GHE provider id must be 'who-ghe'.")

	measure = payload.get("measure")
	if not isinstance(measure, dict) or not str(measure.get("name", "")).strip():
		raise ValueError("WHO GHE measure metadata is invalid.")
	if str(measure.get("sex", "")).strip() != "Persons":
		raise ValueError("WHO GHE builder currently requires sex='Persons'.")

	sheets = payload.get("sheets")
	if not isinstance(sheets, list) or not sheets:
		raise ValueError("WHO GHE sheets are required.")
	years: set[int] = set()
	names: set[str] = set()
	for sheet in sheets:
		if not isinstance(sheet, dict):
			raise ValueError("WHO GHE sheet entries must be objects.")
		year = sheet.get("year")
		name = str(sheet.get("name", "")).strip()
		if not isinstance(year, int) or not 1900 <= year <= 2200 or not name:
			raise ValueError(f"Invalid WHO GHE sheet entry: {sheet!r}")
		if year in years or name in names:
			raise ValueError(f"Duplicate WHO GHE sheet entry: {sheet!r}")
		years.add(year)
		names.add(name)

	minimum = payload.get("minMappedAreas")
	if not isinstance(minimum, int) or minimum < 100:
		raise ValueError("WHO GHE minMappedAreas is invalid.")
	required_areas = payload.get("requiredAreas")
	if not isinstance(required_areas, list) or not required_areas:
		raise ValueError("WHO GHE requiredAreas are required.")

	indicators = payload.get("indicators")
	if not isinstance(indicators, list) or not indicators:
		raise ValueError("WHO GHE indicators are required.")
	ids: set[str] = set()
	slugs: set[str] = set()
	codes: set[int] = set()
	for indicator in indicators:
		if not isinstance(indicator, dict):
			raise ValueError("WHO GHE indicator entries must be objects.")
		for key in ("id", "slug", "sourceCause", "title", "description"):
			if not str(indicator.get(key, "")).strip():
				raise ValueError(f"WHO GHE indicator is missing {key}: {indicator!r}")
		code = indicator.get("gheCode")
		if not isinstance(code, int) or code < 0:
			raise ValueError(f"WHO GHE indicator has invalid gheCode: {indicator['id']}")
		indicator_id = str(indicator["id"])
		slug = str(indicator["slug"])
		if indicator_id in ids or slug in slugs or code in codes:
			raise ValueError(f"Duplicate WHO GHE indicator metadata: {indicator_id}")
		ids.add(indicator_id)
		slugs.add(slug)
		codes.add(code)
		unit = indicator.get("unit")
		if not isinstance(unit, dict) or not str(unit.get("id", "")).strip() or not str(unit.get("label", "")).strip():
			raise ValueError(f"WHO GHE indicator {indicator_id} has invalid unit metadata.")
		common.validate_classification(indicator_id, indicator.get("classification"))
		min_year = indicator.get("minYear")
		if min_year is not None and (not isinstance(min_year, int) or min_year not in years):
			raise ValueError(f"WHO GHE indicator {indicator_id} has invalid minYear.")

	return payload


def fetch_workbook(url: str, timeout: int) -> tuple[bytes, int]:
	delays = (0, 5, 15, 30)
	last_error: Exception | None = None
	for attempt, delay in enumerate(delays, start=1):
		if delay:
			time.sleep(delay)
		request = Request(
			url,
			headers={
				"Accept": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet,*/*",
				"User-Agent": USER_AGENT,
			},
		)
		try:
			with urlopen(request, timeout=timeout) as response:
				data = response.read()
			if len(data) < 100_000:
				raise RuntimeError(f"WHO GHE workbook unexpectedly small: {len(data)} bytes.")
			return data, len(data)
		except (HTTPError, URLError, TimeoutError, RuntimeError) as error:
			last_error = error
			print(f"WHO GHE request failed ({attempt}/{len(delays)}): {error}")
	if last_error is None:
		raise RuntimeError("WHO GHE request failed without exception.")
	raise RuntimeError(f"WHO GHE request failed after {len(delays)} attempts.") from last_error


def normalize_iso3(value: Any) -> str:
	text = str(value if value is not None else "").strip().upper()
	return text if len(text) == 3 and text.isalpha() else ""


def parse_rate(value: Any, indicator_id: str, year: int, iso3: str) -> int | float:
	if value is None or str(value).strip() == "":
		raise RuntimeError(f"Missing WHO GHE value for {indicator_id}, {year}, {iso3}.")
	try:
		number = float(value)
	except (TypeError, ValueError) as error:
		raise RuntimeError(f"Invalid WHO GHE value for {indicator_id}, {year}, {iso3}: {value!r}") from error
	if not math.isfinite(number) or number < 0:
		raise RuntimeError(f"Invalid WHO GHE rate for {indicator_id}, {year}, {iso3}: {number}")
	if number.is_integer():
		return int(number)
	return number


def sheet_country_columns(
	sheet: Any,
	area_by_iso3: dict[str, str],
	minimum: int,
) -> tuple[list[tuple[int, str, str]], set[str]]:
	row = next(sheet.iter_rows(min_row=8, max_row=8, values_only=True))
	columns: list[tuple[int, str, str]] = []
	ignored: set[str] = set()
	seen_area_ids: set[str] = set()
	for index, value in enumerate(row):
		iso3 = normalize_iso3(value)
		if not iso3:
			continue
		area_id = area_by_iso3.get(iso3)
		if area_id is None:
			ignored.add(iso3)
			continue
		if area_id in seen_area_ids:
			raise RuntimeError(f"Duplicate WHO GHE country column in {sheet.title}: {iso3} / {area_id}")
		seen_area_ids.add(area_id)
		columns.append((index, iso3, area_id))

	if len(columns) < minimum:
		raise RuntimeError(
			f"WHO GHE mapped country coverage too small in {sheet.title}: {len(columns)}, expected at least {minimum}."
		)
	if ignored:
		raise RuntimeError(f"Unmapped WHO GHE ISO3 codes in {sheet.title}: {sorted(ignored)}")
	return columns, ignored


def cause_rows(sheet: Any, indicators: list[dict[str, Any]], sex: str) -> dict[int, tuple[Any, ...]]:
	expected = {int(indicator["gheCode"]): indicator for indicator in indicators}
	found: dict[int, tuple[Any, ...]] = {}
	for row in sheet.iter_rows(min_row=10, values_only=True):
		if str(row[0] or "").strip() != sex:
			continue
		code_value = row[1]
		try:
			code = int(code_value)
		except (TypeError, ValueError):
			continue
		if code not in expected:
			continue
		if code in found:
			raise RuntimeError(f"Duplicate WHO GHE cause code {code} in {sheet.title} for sex={sex}.")
		cause = str(row[5] or row[4] or row[3] or "").strip()
		expected_cause = str(expected[code]["sourceCause"]).strip()
		if cause != expected_cause:
			raise RuntimeError(
				f"WHO GHE cause name changed for code {code} in {sheet.title}: {cause!r}, expected {expected_cause!r}."
			)
		found[code] = row

	missing = sorted(set(expected) - set(found))
	if missing:
		raise RuntimeError(f"WHO GHE cause codes missing from {sheet.title}: {missing}")
	return found


def build_indicator_payload(
	config: dict[str, Any],
	indicator: dict[str, Any],
	values_by_year: dict[int, dict[str, int | float]],
	registry_area_count: int,
) -> dict[str, Any]:
	available_years = sorted(values_by_year)
	if not available_years:
		raise RuntimeError(f"WHO GHE indicator has no years: {indicator['id']}")
	default_year = available_years[-1]
	areas_with_any = len({area_id for values in values_by_year.values() for area_id in values})
	minimum = int(config["minMappedAreas"])
	if areas_with_any < minimum:
		raise RuntimeError(
			f"WHO GHE coverage too small for {indicator['id']}: {areas_with_any}, expected at least {minimum}."
		)
	for area_id in config["requiredAreas"]:
		if area_id not in values_by_year[default_year]:
			raise RuntimeError(f"Required area {area_id} has no WHO GHE value for {indicator['id']} in {default_year}.")

	provider = config["provider"]
	measure = config["measure"]
	return {
		"schema": "kartensammlung.statistics-indicator/v1",
		"indicator": {
			"id": indicator["id"],
			"title": indicator["title"],
			"description": indicator["description"],
			"areaLevel": "country",
			"frequency": "selected-years",
			"unit": indicator["unit"],
			"classification": indicator["classification"],
		},
		"source": {
			"providerId": provider["id"],
			"providerName": provider["name"],
			"dataset": provider["dataset"],
			"indicator": f"GHE code {indicator['gheCode']}",
			"indicatorName": indicator["sourceCause"],
			"metric": measure["name"],
			"sex": measure["sex"],
			"provenance": "WHO Global Health Estimates modeled estimate",
			"license": provider["license"],
			"licenseUrl": provider["licenseUrl"],
			"attribution": provider["attribution"],
			"url": config["sourceUrl"],
			"downloadUrl": config["workbookUrl"],
		},
		"availableYears": available_years,
		"defaultYear": default_year,
		"coverage": {
			"registryAreas": registry_area_count,
			"areasWithAnyValue": areas_with_any,
			"latestYear": default_year,
			"areasInLatestYear": len(values_by_year[default_year]),
			"areasInDefaultYear": len(values_by_year[default_year]),
		},
		"values": {
			str(year): {area_id: values_by_year[year][area_id] for area_id in sorted(values_by_year[year])}
			for year in available_years
		},
	}


def main() -> None:
	args = parse_args()
	config = validate_config(common.read_json_path(args.config))
	provider_catalog = common.validate_provider_catalog(common.read_json_path(args.providers))
	provider = config["provider"]
	if not any(item["id"] == provider["id"] for item in provider_catalog["providers"]):
		raise RuntimeError("WHO GHE is missing from statistics-providers.json.")

	area_by_iso3, registry_payload = common.load_registry(args.registry, args.timeout)
	registry_area_count = len(area_by_iso3)
	if registry_area_count < 245:
		raise RuntimeError(f"Country registry unexpectedly small: {registry_area_count} areas.")

	data, byte_count = fetch_workbook(str(config["workbookUrl"]), args.timeout)
	print(f"WHO GHE workbook downloaded={byte_count} bytes")
	try:
		workbook = openpyxl.load_workbook(BytesIO(data), read_only=True, data_only=True)
	except Exception as error:
		raise RuntimeError("WHO GHE workbook could not be opened.") from error

	expected_sheet_names = {str(sheet["name"]) for sheet in config["sheets"]}
	missing_sheets = sorted(expected_sheet_names - set(workbook.sheetnames))
	if missing_sheets:
		raise RuntimeError(f"WHO GHE workbook is missing expected sheets: {missing_sheets}")

	all_values: dict[str, dict[int, dict[str, int | float]]] = {
		str(indicator["id"]): {} for indicator in config["indicators"]
	}
	for sheet_config in sorted(config["sheets"], key=lambda item: int(item["year"])):
		year = int(sheet_config["year"])
		sheet = workbook[str(sheet_config["name"])]
		columns, _ignored = sheet_country_columns(sheet, area_by_iso3, int(config["minMappedAreas"]))
		rows = cause_rows(sheet, config["indicators"], str(config["measure"]["sex"]))
		print(f"WHO GHE {year}: mappedCountries={len(columns)} causes={len(rows)}")

		for indicator in config["indicators"]:
			min_year = int(indicator.get("minYear", year))
			if year < min_year:
				continue
			row = rows[int(indicator["gheCode"])]
			values: dict[str, int | float] = {}
			for column_index, iso3, area_id in columns:
				if column_index >= len(row):
					raise RuntimeError(
						f"WHO GHE row too short for {indicator['id']}, {year}, {iso3}: column {column_index}."
					)
				values[area_id] = parse_rate(row[column_index], str(indicator["id"]), year, iso3)
			if len(values) < int(config["minMappedAreas"]):
				raise RuntimeError(
					f"WHO GHE year coverage too small for {indicator['id']} {year}: {len(values)}."
				)
			all_values[str(indicator["id"])][year] = values

	now = datetime.now(timezone.utc)
	indicator_payloads: list[tuple[dict[str, Any], dict[str, Any]]] = []
	for indicator in config["indicators"]:
		payload = build_indicator_payload(
			config,
			indicator,
			all_values[str(indicator["id"])],
			registry_area_count,
		)
		indicator_payloads.append((indicator, payload))
		coverage = payload["coverage"]
		print(
			f"{indicator['id']}: years={payload['availableYears']} "
			f"defaultYear={payload['defaultYear']} areas={coverage['areasInDefaultYear']}/{coverage['registryAreas']}"
		)

	hasher = hashlib.sha256()
	for indicator, payload in sorted(indicator_payloads, key=lambda item: item[0]["id"]):
		hasher.update(str(indicator["id"]).encode("utf-8") + b"\0" + common.canonical_bytes(payload) + b"\0")
	snapshot = hasher.hexdigest()[:16]
	provider_dir = args.output_dir / provider["id"]
	release_dir = provider_dir / "releases" / snapshot
	index_indicators: list[dict[str, Any]] = []
	for indicator, payload in indicator_payloads:
		filename = f"{indicator['slug']}.json"
		common.write_json(release_dir / filename, payload)
		index_indicators.append({
			"id": indicator["id"],
			"title": indicator["title"],
			"description": indicator["description"],
			"areaLevel": "country",
			"frequency": "selected-years",
			"unit": indicator["unit"],
			"classification": indicator["classification"],
			"sourceIndicator": f"GHE code {indicator['gheCode']}",
			"path": f"releases/{snapshot}/{filename}",
			"availableYears": payload["availableYears"],
			"defaultYear": payload["defaultYear"],
			"coverage": payload["coverage"],
		})

	registry_source = {
		"url": args.registry if args.registry.startswith(("https://", "http://")) else None,
		"schema": registry_payload.get("schema"),
		"generatedAt": registry_payload.get("generatedAt"),
		"areaCount": registry_area_count,
	}
	if registry_source["url"] is None:
		registry_source.pop("url")

	provider_index = {
		"schema": "kartensammlung.statistics-provider-index/v1",
		"provider": provider,
		"retrievedAt": now.isoformat().replace("+00:00", "Z"),
		"activeSnapshot": snapshot,
		"areaRegistry": registry_source,
		"dataset": {
			"release": "Global Health Estimates 2021",
			"releaseDate": "2024-06",
			"measure": config["measure"]["name"],
			"sex": config["measure"]["sex"],
			"availableYears": sorted(int(sheet["year"]) for sheet in config["sheets"]),
		},
		"indicators": index_indicators,
		"notes": [
			"Die Werte sind modellierte WHO Global Health Estimates und nicht rohe registrierte Todesfallzahlen.",
			"Dargestellt werden altersstandardisierte Sterberaten je 100.000 Einwohner für beide Geschlechter zusammen.",
			"Die WHO weist darauf hin, dass die GHE-2021-Reihe wegen methodischer Änderungen nicht mit früheren GHE-Veröffentlichungen verglichen werden soll.",
			"Für COVID-19 werden nur 2020 und 2021 veröffentlicht; vorpandemische Nullwerte im Arbeitsbuch werden bewusst ausgelassen.",
		],
	}
	common.write_json(provider_dir / "index.json", provider_index)
	common.write_json(args.output_dir / "index.json", common.build_global_index(provider_catalog))
	print(f"Built WHO GHE snapshot {snapshot}")
	print(f"Provider index: {provider_dir / 'index.json'}")
	print(f"Indicators: {len(index_indicators)}")


if __name__ == "__main__":
	main()
